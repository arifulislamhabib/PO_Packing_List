# # """
# # PO PDF -> Packing List XLSX
# # Run: streamlit run app.py
# # """
# # import io
# # import re
# # from collections import defaultdict
# # from dataclasses import dataclass

# # import streamlit as st
# # import pdfplumber
# # from openpyxl import Workbook
# # from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


# # # ===============================================================
# # # CONFIG
# # # ===============================================================
# # SIZES = ["XS", "S", "M", "L", "XL", "XXL", "XXL+"]

# # NW  = {"XS": 0.510, "S": 0.520, "M": 0.550, "L": 0.560,
# #        "XL": 0.610, "XXL": 0.650, "XXL+": 0.690}
# # NNW = {"XS": 0.49,  "S": 0.50,  "M": 0.53,  "L": 0.54,
# #        "XL": 0.59,  "XXL": 0.63,  "XXL+": 0.67}

# # CARTON_ADDON = {"MULTI": 0.71, "SINGLE": 0.71, "BULK": 0.42}
# # NET_DEDUCT   = {"MULTI": 0.42, "SINGLE": 0.24, "BULK": 0.24}

# # ANCHOR_RE = re.compile(
# #     r"(\d{3})(\d{6})-(\d{3})\s+"
# #     r"(XXL\+|XXL|XL|XS|S|M|L)\s+"
# #     r"([\d.]+)\s+([\d,]+\.\d{2})"
# # )

# # COLOR_STOPWORDS = {"SIZE", "TOTAL", "COLOR", "PCS", "CTNS"}


# # # ===============================================================
# # # MODELS
# # # ===============================================================
# # @dataclass
# # class Item:
# #     pack: str
# #     color_code: str
# #     color_name: str
# #     size: str
# #     qty: int


# # @dataclass
# # class Carton:
# #     """A carton GROUP (one row in the packing list)."""
# #     ctn_no: str
# #     ctn_mes: str
# #     ctn_qty: int              # how many cartons in this group
# #     color: str
# #     size_qty: dict            # per ONE carton: {size: qty or packs}
# #     units_per_ctn: int        # MULTI: 7 (pcs) ; SINGLE: 2 (multiplier) ; BULK: 1
# #     qty_per_ctn: int          # total pcs per carton
# #     total_qty: int            # ctn_qty * qty_per_ctn
# #     gross_wt: float = 0.0     # per carton
# #     net_wt: float = 0.0       # per carton


# # # ===============================================================
# # # PDF PARSER
# # # ===============================================================
# # def _search(text, pattern):
# #     m = re.search(pattern, text, re.I)
# #     return m.group(1).strip() if m else ""


# # def _looks_like_color(line: str) -> bool:
# #     if not line:
# #         return False
# #     s = line.strip()
# #     if not (3 <= len(s) <= 40):
# #         return False
# #     if not s.isupper():
# #         return False
# #     if s in COLOR_STOPWORDS:
# #         return False
# #     if not all(ch.isalpha() or ch == " " for ch in s):
# #         return False
# #     if sum(ch.isalpha() for ch in s) < 3:
# #         return False
# #     if any(w in COLOR_STOPWORDS for w in s.split()):
# #         return False
# #     return True


# # def parse_po(pdf_bytes: bytes):
# #     header, items, color_map = {}, [], {}

# #     with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
# #         t0 = pdf.pages[0].extract_text() or ""
# #         header["po"]          = _search(t0, r"Destination\s*Purchase\s*Order\s*#?\s*(\d{6,})")
# #         header["ship_before"] = _search(t0, r"Do Not Ship Before Date\s*(\S+)")
# #         header["ship_cancel"] = _search(t0, r"Ship Cancel Date\s*(\S+)")
# #         header["in_dc"]       = _search(t0, r"In DC Date\s*(\S+)")
# #         header["stock"]       = _search(t0, r"Planned Stock Date\s*(\S+)")

# #         for page in pdf.pages[1:]:
# #             text = page.extract_text() or ""
# #             lines = [l.rstrip() for l in text.splitlines()]

# #             for i, line in enumerate(lines):
# #                 ls = line.strip()
# #                 if not ls.startswith("3340"):
# #                     continue
# #                 m = ANCHOR_RE.search(ls)
# #                 if not m:
# #                     continue
# #                 color_code = m.group(3)
# #                 size       = m.group(4)
# #                 cost       = float(m.group(5))
# #                 total      = float(m.group(6).replace(",", ""))
# #                 if cost <= 0:
# #                     continue
# #                 qty = round(total / cost)

# #                 if "Single" in ls:
# #                     pt = "SINGLE"
# #                 elif "Multi" in ls:
# #                     pt = "MULTI"
# #                 elif "Bulk" in ls:
# #                     pt = "BULK"
# #                 else:
# #                     continue

# #                 # dynamic color from next line
# #                 color_name = color_map.get(color_code, "")
# #                 if i + 1 < len(lines):
# #                     nxt = lines[i + 1].strip()
# #                     if _looks_like_color(nxt):
# #                         color_name = nxt
# #                         color_map[color_code] = nxt
# #                 if not color_name:
# #                     color_name = f"COLOR-{color_code}"

# #                 items.append(Item(pt, color_code, color_name, size, qty))

# #     seen, clean = set(), []
# #     for it in items:
# #         k = (it.pack, it.color_code, it.size, it.qty)
# #         if k in seen:
# #             continue
# #         seen.add(k)
# #         clean.append(it)

# #     for it in clean:
# #         if it.color_name.startswith("COLOR-") and it.color_code in color_map:
# #             it.color_name = color_map[it.color_code]

# #     return header, clean, color_map


# # # ===============================================================
# # # CARTON PLANNING
# # # ===============================================================
# # def plan_cartons(items):
# #     grouped = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
# #     for it in items:
# #         grouped[it.pack][it.color_name][it.size] += it.qty

# #     result = {}
# #     for pt, by_color in grouped.items():
# #         result[pt] = {}
# #         for color, sizes in by_color.items():
# #             if pt == "MULTI":
# #                 result[pt][color] = _plan_multi(color, sizes)
# #             elif pt == "SINGLE":
# #                 result[pt][color] = _plan_single(color, sizes)
# #             else:
# #                 result[pt][color] = _plan_bulk(color, sizes)
# #     return result


# # def _plan_multi(color, sizes):
# #     """All cartons identical -> ONE row with CTN QTY=n."""
# #     total = sum(sizes.values())
# #     n = total // 7
# #     if n == 0:
# #         return []
# #     per = {s: sizes.get(s, 0) // n for s in SIZES}
# #     return [Carton(
# #         ctn_no="1",
# #         ctn_mes="G82",
# #         ctn_qty=n,
# #         color=color,
# #         size_qty=dict(per),
# #         units_per_ctn=7,
# #         qty_per_ctn=7,
# #         total_qty=n * 7,
# #     )]


# # def _plan_single(color, sizes):
# #     """sizes: {size: pcs}. Split into cartons of up to 20 pcs (10 packs)."""
# #     raw = []
# #     for sz in SIZES:
# #         pcs = sizes.get(sz, 0)
# #         while pcs > 0:
# #             take = min(20, pcs)
# #             pcs -= take
# #             raw.append(Carton(
# #                 ctn_no="",
# #                 ctn_mes="G81" if take == 20 else "G82",
# #                 ctn_qty=1,
# #                 color=color,
# #                 size_qty={sz: take // 2},   # packs
# #                 units_per_ctn=2,
# #                 qty_per_ctn=take,
# #                 total_qty=take,
# #             ))
# #     return _group_cartons(raw)


# # def _plan_bulk(color, sizes):
# #     """sizes: {size: pcs}. Split into cartons of up to 20 pcs."""
# #     raw = []
# #     for sz in SIZES:
# #         qty = sizes.get(sz, 0)
# #         while qty > 0:
# #             take = min(20, qty)
# #             qty -= take
# #             raw.append(Carton(
# #                 ctn_no="",
# #                 ctn_mes="G81" if take == 20 else "G82",
# #                 ctn_qty=1,
# #                 color=color,
# #                 size_qty={sz: take},
# #                 units_per_ctn=1,
# #                 qty_per_ctn=take,
# #                 total_qty=take,
# #             ))
# #     return _group_cartons(raw)


# # def _group_cartons(raw):
# #     """Collapse consecutive cartons with identical size distribution + carton code."""
# #     if not raw:
# #         return []
# #     groups = []
# #     cur = raw[0]
# #     for c in raw[1:]:
# #         if c.size_qty == cur.size_qty and c.ctn_mes == cur.ctn_mes:
# #             cur.ctn_qty += 1
# #             cur.total_qty += c.total_qty
# #         else:
# #             groups.append(cur)
# #             cur = c
# #     groups.append(cur)

# #     for i, g in enumerate(groups, 1):
# #         g.ctn_no = str(i) if g.ctn_qty == 1 else f"{i}-{i + g.ctn_qty - 1}"
# #     return groups


# # # ===============================================================
# # # WEIGHTS  (per carton)
# # # ===============================================================
# # def apply_weights(packs):
# #     for pt, by_color in packs.items():
# #         addon  = CARTON_ADDON[pt]
# #         deduct = NET_DEDUCT[pt]
# #         for color, carts in by_color.items():
# #             for c in carts:
# #                 if pt == "SINGLE":
# #                     nw_sum = sum(NW[s] * q * c.units_per_ctn
# #                                  for s, q in c.size_qty.items())
# #                 else:
# #                     nw_sum = sum(NW[s] * q for s, q in c.size_qty.items())
# #                 c.gross_wt = round(nw_sum + addon, 2)
# #                 c.net_wt   = round(c.gross_wt - deduct, 2)


# # # ===============================================================
# # # XLSX WRITER
# # # ===============================================================
# # THIN = Side(style="thin", color="000000")
# # BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
# # HDR_FILL = PatternFill("solid", fgColor="FFF2CC")
# # TITLE_FILL = PatternFill("solid", fgColor="FCE4D6")
# # BOLD = Font(bold=True)
# # CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
# # RIGHT = Alignment(horizontal="right", vertical="center")
# # LEFT = Alignment(horizontal="left", vertical="center")


# # def _c(ws, r, c, v, bold=False, center=False, fill=None):
# #     cell = ws.cell(row=r, column=c, value=v)
# #     if bold: cell.font = BOLD
# #     cell.alignment = CENTER if center else (RIGHT if isinstance(v, (int, float)) else LEFT)
# #     if fill: cell.fill = fill
# #     cell.border = BORDER


# # class Builder:
# #     def __init__(self, header, packs):
# #         self.header = header
# #         self.packs = packs
# #         self.wb = Workbook()
# #         self.ws = self.wb.active
# #         self.ws.title = "Packing List"
# #         self.row = 1

# #     def build(self):
# #         self._company_header()
# #         self._global_table()
# #         for pt in ("MULTI", "SINGLE", "BULK"):
# #             if pt not in self.packs:
# #                 continue
# #             for color, carts in self.packs[pt].items():
# #                 self._pack_table(pt, color, carts)
# #         self._summary()
# #         self._widths()
# #         return self.wb

# #     def _company_header(self):
# #         self.ws.merge_cells("A1:Q1")
# #         self.ws["A1"] = "CREATIVE COLLECTIONS LTD-1A."
# #         self.ws["A1"].font = Font(bold=True, size=14)
# #         self.ws["A1"].alignment = CENTER
# #         self.ws.merge_cells("A2:Q2")
# #         self.ws["A2"] = "Nishat Nagar , Tongi , Gazipur ."
# #         self.ws["A2"].font = Font(bold=True, size=12)
# #         self.ws["A2"].alignment = CENTER
# #         self.row = 4

# #     def _global_table(self):
# #         rows = [
# #             ["SIZE", "", "SIZE", *SIZES],
# #             ["N.W", "", "N.WT", *[NW[s] for s in SIZES]],
# #             ["N.N.W", "", "N.N. WT", *[NNW[s] for s in SIZES]],
# #             ["EMPTY CTN WET", "", "", *[""] * len(SIZES)],
# #         ]
# #         for i, rd in enumerate(rows):
# #             for c, v in enumerate(rd, 1):
# #                 _c(self.ws, self.row + i, c, v, bold=True, center=True, fill=HDR_FILL)
# #         self.row += 6

# #     def _pack_table(self, pt, color, carts):
# #         title = {"MULTI": "MULTY PACK ( Y )", "SINGLE": "SINGLE  PACK", "BULK": "BULK PACK"}[pt]
# #         total_pcs  = sum(c.total_qty for c in carts)
# #         total_ctns = sum(c.ctn_qty for c in carts)
# #         po = self.header.get("po", "")

# #         self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=17)
# #         _c(self.ws, self.row, 1, title, bold=True, center=True, fill=TITLE_FILL)
# #         self.row += 1

# #         info = [
# #             ("BUYER", ":", "OLD NAVY"),
# #             ("STYLE", ":", "905518"),
# #             ("P. O. #", ":", po),
# #             ("O/QTY", ":", total_pcs, "PCS"),
# #             ("SHIP QTY", ":", total_pcs, "PCS"),
# #             ("EX/SHORT", ":", 0, "PCS"),
# #             ("CTN QTY", ":", total_ctns, "CTNS"),
# #             ("CTN MEAS.", ":", "58.67 X 38.48 X 29.71 CM.G8_SL"),
# #             ("CTN MEAS.", ":", "58.67 X 38.48 X 14.86 CM.G8S_SL"),
# #             ("CTN MEAS.", ":", "38.48 X 29.33 X 14.86 CM.G-8M"),
# #             ("CTN MEAS.", ":", "29.33 X 19.25 X 14.86 CM.G-8XS"),
# #         ]
# #         for i, tup in enumerate(info):
# #             for c, v in enumerate(tup, 1):
# #                 _c(self.ws, self.row + i, c, v, bold=(c <= 2))

# #         right = [
# #             ("PO #", po),
# #             ("SKU / Item", "905518"),
# #             ("Unit / Prepack", {"MULTI": "184/23", "SINGLE": "80/40", "BULK": "80/40"}[pt]),
# #             ("CARTON", f"01 of {total_ctns}"),
# #         ]
# #         for i, (k, v) in enumerate(right):
# #             _c(self.ws, self.row + i, 13, k, bold=True)
# #             _c(self.ws, self.row + i, 15, v)

# #         self.row += len(info) + 2

# #         if pt == "MULTI":
# #             headers = ["CTN NO", "CTN\nMES:", "CTN\nQTY", "COLOR", "PREPACK\nSTECKER",
# #                        *SIZES, "PER\nBLST", "QTY\nPER\nCTN", "TOTAL\nQTY",
# #                        "GROSS\nWEIGHT", "NET\nWEIGHT"]
# #         else:
# #             headers = ["CTN NO", "CTN\nMES:", "CTN\nQTY", "COLOR", "PREPACK\nSTECKER",
# #                        *SIZES, "QTY\nPER\nCTN", "TOTAL\nQTY",
# #                        "GROSS\nWEIGHT", "NET\nWEIGHT"]
# #         for c, h in enumerate(headers, 1):
# #             _c(self.ws, self.row, c, h, bold=True, center=True, fill=HDR_FILL)
# #         self.row += 1

# #         size_totals = {s: 0 for s in SIZES}
# #         for ct in carts:
# #             vals = [ct.ctn_no, ct.ctn_mes, ct.ctn_qty, ct.color, "",
# #                     *[ct.size_qty.get(s, "") for s in SIZES]]
# #             if pt == "MULTI":
# #                 vals += ["7 X 1", ct.qty_per_ctn, ct.total_qty, ct.gross_wt, ct.net_wt]
# #             else:
# #                 vals += [ct.qty_per_ctn, ct.total_qty, ct.gross_wt, ct.net_wt]
# #             for i, v in enumerate(vals, 1):
# #                 _c(self.ws, self.row, i, v, center=(i <= 5))
# #             for s in SIZES:
# #                 size_totals[s] += (ct.size_qty.get(s, 0) or 0) * ct.ctn_qty
# #             self.row += 1

# #         if pt == "MULTI":
# #             tot = ["TOTAL", "", total_ctns, "CTNS", "",
# #                    *[size_totals[s] for s in SIZES],
# #                    "", "", total_pcs, "PCS", ""]
# #         else:
# #             tot = ["TOTAL", "", total_ctns, "CTN", "",
# #                    *[size_totals[s] for s in SIZES],
# #                    "", total_pcs, "PCS", ""]
# #         for c, v in enumerate(tot, 1):
# #             _c(self.ws, self.row, c, v, bold=True, center=True, fill=HDR_FILL)
# #         self.row += 2

# #     def _summary(self):
# #         agg = defaultdict(lambda: {s: 0 for s in SIZES})
# #         for pt, by_color in self.packs.items():
# #             for color, carts in by_color.items():
# #                 for ct in carts:
# #                     for sz, q in ct.size_qty.items():
# #                         if pt == "MULTI":
# #                             pcs = q * ct.ctn_qty
# #                         else:
# #                             pcs = q * ct.ctn_qty * ct.units_per_ctn
# #                         agg[color][sz] += pcs

# #         grand = {s: 0 for s in SIZES}
# #         for color in list(agg.keys()) + ["G. Total Summery"]:
# #             self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=10)
# #             _c(self.ws, self.row, 1, color, bold=True, center=True, fill=TITLE_FILL)
# #             self.row += 1

# #             headers = ["SIZE", *SIZES, "", "TOTAL"]
# #             for c, h in enumerate(headers, 1):
# #                 _c(self.ws, self.row, c, h, bold=True, center=True, fill=HDR_FILL)
# #             self.row += 1

# #             vals = ["SHIP QTY"]
# #             row_total = 0
# #             for s in SIZES:
# #                 v = grand[s] if color == "G. Total Summery" else agg[color][s]
# #                 vals.append(v)
# #                 row_total += v
# #             vals += ["", row_total]
# #             for c, v in enumerate(vals, 1):
# #                 _c(self.ws, self.row, c, v, bold=(c == 1), center=(c == 1))
# #             self.row += 1

# #             if color != "G. Total Summery":
# #                 for s in SIZES:
# #                     grand[s] += agg[color][s]
# #             self.row += 1

# #     def _widths(self):
# #         for col, w in {"A": 12, "B": 8, "C": 6, "D": 14, "E": 14, "F": 5,
# #                        "G": 5, "H": 5, "I": 5, "J": 5, "K": 5, "L": 7,
# #                        "M": 8, "N": 9, "O": 9, "P": 9, "Q": 9}.items():
# #             self.ws.column_dimensions[col].width = w


# # # ===============================================================
# # # STREAMLIT UI
# # # ===============================================================
# # st.set_page_config(page_title="PO → Packing List", page_icon="📦", layout="centered")
# # st.title("📦 PO → Packing List Generator")
# # st.caption("Upload a Destination Purchase Order PDF. Download a full packing list XLSX.")

# # pdf_file = st.file_uploader("**PO PDF**", type=["pdf"])

# # if pdf_file is None:
# #     st.info("⬆️ Upload a PO PDF to begin.")
# #     st.stop()

# # if st.button("🚀 Generate Packing List", type="primary"):
# #     with st.spinner("Parsing PO and generating packing list..."):
# #         header, items, color_map = parse_po(pdf_file.read())
# #         if not items:
# #             st.error("No line items found.")
# #             st.stop()
# #         st.info(f"🎨 Colors detected: **{', '.join(color_map.values())}**")
# #         packs = plan_cartons(items)
# #         apply_weights(packs)
# #         wb = Builder(header, packs).build()
# #         buf = io.BytesIO()
# #         wb.save(buf)
# #         st.session_state["out_bytes"] = buf.getvalue()
# #         st.session_state["header"]    = header
# #         st.session_state["packs"]     = packs
# #         st.session_state["color_map"] = color_map
# #         st.session_state["filename"]  = f"Packing_List_{header.get('po','PO')}.xlsx"

# # if "out_bytes" in st.session_state:
# #     header    = st.session_state["header"]
# #     packs     = st.session_state["packs"]
# #     color_map = st.session_state["color_map"]

# #     total_pcs  = sum(c.total_qty for pt in packs.values() for carts in pt.values() for c in carts)
# #     total_ctns = sum(c.ctn_qty  for pt in packs.values() for carts in pt.values() for c in carts)

# #     st.success(f"✅ Generated — {total_pcs:,} pcs across {total_ctns} cartons")

# #     col1, col2, col3 = st.columns(3)
# #     col1.metric("PO #", header.get("po", "—"))
# #     col2.metric("Total Pieces", f"{total_pcs:,}")
# #     col3.metric("Total Cartons", total_ctns)

# #     with st.expander("🎨 Color mapping detected", expanded=True):
# #         for code, name in color_map.items():
# #             st.write(f"`{code}` → **{name}**")

# #     st.subheader("Section breakdown")
# #     for pt, by_color in packs.items():
# #         for color, carts in by_color.items():
# #             pcs = sum(c.total_qty for c in carts)
# #             st.write(f"**{pt}** · {color} · {sum(c.ctn_qty for c in carts)} cartons · {pcs} pcs")

# #     st.download_button(
# #         "⬇️  Download Packing List XLSX",
# #         data=st.session_state["out_bytes"],
# #         file_name=st.session_state["filename"],
# #         mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
# #         type="primary",
# #     )







# """
# PO PDF -> Packing List XLSX
# Dynamic pack-type detection (MULTI / SINGLE / BULK — any combination)

# Run: streamlit run app.py
# """
# import io
# import re
# from collections import defaultdict
# from dataclasses import dataclass

# import streamlit as st
# import pdfplumber
# from openpyxl import Workbook
# from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


# # ===============================================================
# # CONFIG
# # ===============================================================
# SIZES = ["XS", "S", "M", "L", "XL", "XXL", "XXL+"]

# NW  = {"XS": 0.510, "S": 0.520, "M": 0.550, "L": 0.560,
#        "XL": 0.610, "XXL": 0.650, "XXL+": 0.690}
# NNW = {"XS": 0.49,  "S": 0.50,  "M": 0.53,  "L": 0.54,
#        "XL": 0.59,  "XXL": 0.63,  "XXL+": 0.67}

# CARTON_ADDON = {"MULTI": 0.71, "SINGLE": 0.71, "BULK": 0.42}
# NET_DEDUCT   = {"MULTI": 0.42, "SINGLE": 0.24, "BULK": 0.24}

# # Display order: whichever of these exist will be shown in this order
# PACK_ORDER = ["MULTI", "SINGLE", "BULK"]

# # Pack-type display titles
# PACK_TITLES = {
#     "MULTI":  "MULTY PACK ( Y )",
#     "SINGLE": "SINGLE  PACK",
#     "BULK":   "BULK PACK",
# }

# # Unit / Prepack labels
# UNIT_PREPACK = {
#     "MULTI":  "184/23",
#     "SINGLE": "80/40",
#     "BULK":   "80/40",
# }

# # Anchor: 3-digit prefix + style-color + size + unit cost + total cost
# ANCHOR_RE = re.compile(
#     r"(\d{3})(\d{6})-(\d{3})\s+"
#     r"(XXL\+|XXL|XL|XS|S|M|L)\s+"
#     r"([\d.]+)\s+([\d,]+\.\d{2})"
# )

# COLOR_STOPWORDS = {"SIZE", "TOTAL", "COLOR", "PCS", "CTNS"}

# # Pack markers we know about
# KNOWN_PACK_MARKERS = {
#     "Multi":  "MULTI",
#     "Single": "SINGLE",
#     "Bulk":   "BULK",
# }

# # Common "other pack type" hints for warnings (extend as you encounter them)
# UNKNOWN_HINTS = ("Prepack", "Assorted", "Master", "Mixed", "Combo")


# # ===============================================================
# # MODELS
# # ===============================================================
# @dataclass
# class Item:
#     pack: str
#     color_code: str
#     color_name: str
#     size: str
#     qty: int


# @dataclass
# class Carton:
#     ctn_no: str
#     ctn_mes: str
#     ctn_qty: int
#     color: str
#     size_qty: dict
#     units_per_ctn: int
#     qty_per_ctn: int
#     total_qty: int
#     gross_wt: float = 0.0
#     net_wt: float = 0.0


# # ===============================================================
# # PDF PARSER
# # ===============================================================
# def _search(text, pattern):
#     m = re.search(pattern, text, re.I)
#     return m.group(1).strip() if m else ""


# def _looks_like_color(line: str) -> bool:
#     if not line:
#         return False
#     s = line.strip()
#     if not (3 <= len(s) <= 40):
#         return False
#     if not s.isupper():
#         return False
#     if s in COLOR_STOPWORDS:
#         return False
#     if not all(ch.isalpha() or ch == " " for ch in s):
#         return False
#     if sum(ch.isalpha() for ch in s) < 3:
#         return False
#     if any(w in COLOR_STOPWORDS for w in s.split()):
#         return False
#     return True


# def parse_po(pdf_bytes: bytes):
#     """
#     Returns:
#         header: dict
#         items: list of Item
#         color_map: {color_code: name}
#         unknown_markers: set of unknown pack-type markers seen
#     """
#     header, items, color_map = {}, [], {}
#     unknown_markers = set()

#     with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
#         # ---- Header from page 1 ----
#         t0 = pdf.pages[0].extract_text() or ""
#         header["po"]          = _search(t0, r"Destination\s*Purchase\s*Order\s*#?\s*(\d{6,})")
#         header["ship_before"] = _search(t0, r"Do Not Ship Before Date\s*(\S+)")
#         header["ship_cancel"] = _search(t0, r"Ship Cancel Date\s*(\S+)")
#         header["in_dc"]       = _search(t0, r"In DC Date\s*(\S+)")
#         header["stock"]       = _search(t0, r"Planned Stock Date\s*(\S+)")

#         # ---- Line items ----
#         for page in pdf.pages[1:]:
#             text = page.extract_text() or ""
#             lines = [l.rstrip() for l in text.splitlines()]

#             for i, line in enumerate(lines):
#                 ls = line.strip()
#                 if not ls.startswith("3340"):
#                     continue
#                 m = ANCHOR_RE.search(ls)
#                 if not m:
#                     continue
#                 color_code = m.group(3)
#                 size       = m.group(4)
#                 cost       = float(m.group(5))
#                 total      = float(m.group(6).replace(",", ""))
#                 if cost <= 0:
#                     continue
#                 qty = round(total / cost)

#                 # ---------- DYNAMIC pack-type detection ----------
#                 pt = None
#                 for marker, canonical in KNOWN_PACK_MARKERS.items():
#                     if marker in ls:
#                         pt = canonical
#                         break

#                 if pt is None:
#                     # Look for hint of an unknown marker for a warning
#                     for hint in UNKNOWN_HINTS:
#                         if hint in ls:
#                             unknown_markers.add(hint)
#                             break
#                     continue

#                 # ---------- Dynamic color ----------
#                 color_name = color_map.get(color_code, "")
#                 if i + 1 < len(lines):
#                     nxt = lines[i + 1].strip()
#                     if _looks_like_color(nxt):
#                         color_name = nxt
#                         color_map[color_code] = nxt
#                 if not color_name:
#                     color_name = f"COLOR-{color_code}"

#                 items.append(Item(pt, color_code, color_name, size, qty))

#     # ---- De-dup ----
#     seen, clean = set(), []
#     for it in items:
#         k = (it.pack, it.color_code, it.size, it.qty)
#         if k in seen:
#             continue
#         seen.add(k)
#         clean.append(it)

#     for it in clean:
#         if it.color_name.startswith("COLOR-") and it.color_code in color_map:
#             it.color_name = color_map[it.color_code]

#     return header, clean, color_map, unknown_markers


# # ===============================================================
# # CARTON PLANNING
# # ===============================================================
# def plan_cartons(items):
#     grouped = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
#     for it in items:
#         grouped[it.pack][it.color_name][it.size] += it.qty

#     # Result preserves only pack types that actually exist
#     result = {}
#     for pt in PACK_ORDER:                     # stable order
#         if pt not in grouped:
#             continue
#         result[pt] = {}
#         for color, sizes in grouped[pt].items():
#             if pt == "MULTI":
#                 result[pt][color] = _plan_multi(color, sizes)
#             elif pt == "SINGLE":
#                 result[pt][color] = _plan_single(color, sizes)
#             elif pt == "BULK":
#                 result[pt][color] = _plan_bulk(color, sizes)
#     return result


# def _plan_multi(color, sizes):
#     total = sum(sizes.values())
#     n = total // 7
#     if n == 0:
#         return []
#     per = {s: sizes.get(s, 0) // n for s in SIZES}
#     return [Carton(
#         ctn_no="1", ctn_mes="G82", ctn_qty=n, color=color,
#         size_qty=dict(per), units_per_ctn=7, qty_per_ctn=7,
#         total_qty=n * 7,
#     )]


# def _plan_single(color, sizes):
#     raw = []
#     for sz in SIZES:
#         pcs = sizes.get(sz, 0)
#         while pcs > 0:
#             take = min(20, pcs)
#             pcs -= take
#             raw.append(Carton(
#                 ctn_no="",
#                 ctn_mes="G81" if take == 20 else "G82",
#                 ctn_qty=1, color=color,
#                 size_qty={sz: take // 2},
#                 units_per_ctn=2,
#                 qty_per_ctn=take, total_qty=take,
#             ))
#     return _group_cartons(raw)


# def _plan_bulk(color, sizes):
#     raw = []
#     for sz in SIZES:
#         qty = sizes.get(sz, 0)
#         while qty > 0:
#             take = min(20, qty)
#             qty -= take
#             raw.append(Carton(
#                 ctn_no="",
#                 ctn_mes="G81" if take == 20 else "G82",
#                 ctn_qty=1, color=color,
#                 size_qty={sz: take},
#                 units_per_ctn=1,
#                 qty_per_ctn=take, total_qty=take,
#             ))
#     return _group_cartons(raw)


# def _group_cartons(raw):
#     if not raw:
#         return []
#     groups = []
#     cur = raw[0]
#     for c in raw[1:]:
#         if c.size_qty == cur.size_qty and c.ctn_mes == cur.ctn_mes:
#             cur.ctn_qty += 1
#             cur.total_qty += c.total_qty
#         else:
#             groups.append(cur)
#             cur = c
#     groups.append(cur)

#     for i, g in enumerate(groups, 1):
#         g.ctn_no = str(i) if g.ctn_qty == 1 else f"{i}-{i + g.ctn_qty - 1}"
#     return groups


# # ===============================================================
# # WEIGHTS  (per carton)
# # ===============================================================
# def apply_weights(packs):
#     for pt, by_color in packs.items():
#         addon  = CARTON_ADDON[pt]
#         deduct = NET_DEDUCT[pt]
#         for color, carts in by_color.items():
#             for c in carts:
#                 if pt == "SINGLE":
#                     nw_sum = sum(NW[s] * q * c.units_per_ctn
#                                  for s, q in c.size_qty.items())
#                 else:
#                     nw_sum = sum(NW[s] * q for s, q in c.size_qty.items())
#                 c.gross_wt = round(nw_sum + addon, 2)
#                 c.net_wt   = round(c.gross_wt - deduct, 2)


# # ===============================================================
# # XLSX WRITER
# # ===============================================================
# THIN = Side(style="thin", color="000000")
# BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
# HDR_FILL = PatternFill("solid", fgColor="FFF2CC")
# TITLE_FILL = PatternFill("solid", fgColor="FCE4D6")
# BOLD = Font(bold=True)
# CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
# RIGHT = Alignment(horizontal="right", vertical="center")
# LEFT = Alignment(horizontal="left", vertical="center")


# def _c(ws, r, c, v, bold=False, center=False, fill=None):
#     cell = ws.cell(row=r, column=c, value=v)
#     if bold: cell.font = BOLD
#     cell.alignment = CENTER if center else (RIGHT if isinstance(v, (int, float)) else LEFT)
#     if fill: cell.fill = fill
#     cell.border = BORDER


# class Builder:
#     def __init__(self, header, packs):
#         self.header = header
#         self.packs = packs
#         self.wb = Workbook()
#         self.ws = self.wb.active
#         self.ws.title = "Packing List"
#         self.row = 1

#     def build(self):
#         self._company_header()
#         self._global_table()
#         # Iterate over PACK_ORDER but only render those that exist
#         for pt in PACK_ORDER:
#             if pt not in self.packs:
#                 continue
#             for color, carts in self.packs[pt].items():
#                 self._pack_table(pt, color, carts)
#         self._summary()
#         self._widths()
#         return self.wb

#     def _company_header(self):
#         self.ws.merge_cells("A1:Q1")
#         self.ws["A1"] = "CREATIVE COLLECTIONS LTD-1A."
#         self.ws["A1"].font = Font(bold=True, size=14)
#         self.ws["A1"].alignment = CENTER
#         self.ws.merge_cells("A2:Q2")
#         self.ws["A2"] = "Nishat Nagar , Tongi , Gazipur ."
#         self.ws["A2"].font = Font(bold=True, size=12)
#         self.ws["A2"].alignment = CENTER
#         self.row = 4

#     def _global_table(self):
#         rows = [
#             ["SIZE", "", "SIZE", *SIZES],
#             ["N.W", "", "N.WT", *[NW[s] for s in SIZES]],
#             ["N.N.W", "", "N.N. WT", *[NNW[s] for s in SIZES]],
#             ["EMPTY CTN WET", "", "", *[""] * len(SIZES)],
#         ]
#         for i, rd in enumerate(rows):
#             for c, v in enumerate(rd, 1):
#                 _c(self.ws, self.row + i, c, v, bold=True, center=True, fill=HDR_FILL)
#         self.row += 6

#     def _pack_table(self, pt, color, carts):
#         title = PACK_TITLES.get(pt, pt)
#         total_pcs  = sum(c.total_qty for c in carts)
#         total_ctns = sum(c.ctn_qty for c in carts)
#         po = self.header.get("po", "")

#         self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=17)
#         _c(self.ws, self.row, 1, title, bold=True, center=True, fill=TITLE_FILL)
#         self.row += 1

#         info = [
#             ("BUYER", ":", "OLD NAVY"),
#             ("STYLE", ":", "905518"),
#             ("P. O. #", ":", po),
#             ("O/QTY", ":", total_pcs, "PCS"),
#             ("SHIP QTY", ":", total_pcs, "PCS"),
#             ("EX/SHORT", ":", 0, "PCS"),
#             ("CTN QTY", ":", total_ctns, "CTNS"),
#             ("CTN MEAS.", ":", "58.67 X 38.48 X 29.71 CM.G8_SL"),
#             ("CTN MEAS.", ":", "58.67 X 38.48 X 14.86 CM.G8S_SL"),
#             ("CTN MEAS.", ":", "38.48 X 29.33 X 14.86 CM.G-8M"),
#             ("CTN MEAS.", ":", "29.33 X 19.25 X 14.86 CM.G-8XS"),
#         ]
#         for i, tup in enumerate(info):
#             for c, v in enumerate(tup, 1):
#                 _c(self.ws, self.row + i, c, v, bold=(c <= 2))

#         right = [
#             ("PO #", po),
#             ("SKU / Item", "905518"),
#             ("Unit / Prepack", UNIT_PREPACK.get(pt, "")),
#             ("CARTON", f"01 of {total_ctns}"),
#         ]
#         for i, (k, v) in enumerate(right):
#             _c(self.ws, self.row + i, 13, k, bold=True)
#             _c(self.ws, self.row + i, 15, v)

#         self.row += len(info) + 2

#         if pt == "MULTI":
#             headers = ["CTN NO", "CTN\nMES:", "CTN\nQTY", "COLOR", "PREPACK\nSTECKER",
#                        *SIZES, "PER\nBLST", "QTY\nPER\nCTN", "TOTAL\nQTY",
#                        "GROSS\nWEIGHT", "NET\nWEIGHT"]
#         else:
#             headers = ["CTN NO", "CTN\nMES:", "CTN\nQTY", "COLOR", "PREPACK\nSTECKER",
#                        *SIZES, "QTY\nPER\nCTN", "TOTAL\nQTY",
#                        "GROSS\nWEIGHT", "NET\nWEIGHT"]
#         for c, h in enumerate(headers, 1):
#             _c(self.ws, self.row, c, h, bold=True, center=True, fill=HDR_FILL)
#         self.row += 1

#         size_totals = {s: 0 for s in SIZES}
#         for ct in carts:
#             vals = [ct.ctn_no, ct.ctn_mes, ct.ctn_qty, ct.color, "",
#                     *[ct.size_qty.get(s, "") for s in SIZES]]
#             if pt == "MULTI":
#                 vals += ["7 X 1", ct.qty_per_ctn, ct.total_qty, ct.gross_wt, ct.net_wt]
#             else:
#                 vals += [ct.qty_per_ctn, ct.total_qty, ct.gross_wt, ct.net_wt]
#             for i, v in enumerate(vals, 1):
#                 _c(self.ws, self.row, i, v, center=(i <= 5))
#             for s in SIZES:
#                 size_totals[s] += (ct.size_qty.get(s, 0) or 0) * ct.ctn_qty
#             self.row += 1

#         if pt == "MULTI":
#             tot = ["TOTAL", "", total_ctns, "CTNS", "",
#                    *[size_totals[s] for s in SIZES],
#                    "", "", total_pcs, "PCS", ""]
#         else:
#             tot = ["TOTAL", "", total_ctns, "CTN", "",
#                    *[size_totals[s] for s in SIZES],
#                    "", total_pcs, "PCS", ""]
#         for c, v in enumerate(tot, 1):
#             _c(self.ws, self.row, c, v, bold=True, center=True, fill=HDR_FILL)
#         self.row += 2

#     def _summary(self):
#         """
#         Build summary tables dynamically. If a pack type is absent,
#         its rows simply don't appear. Colors are discovered from packs.
#         """
#         agg = defaultdict(lambda: {s: 0 for s in SIZES})
#         for pt, by_color in self.packs.items():
#             for color, carts in by_color.items():
#                 for ct in carts:
#                     for sz, q in ct.size_qty.items():
#                         if pt == "MULTI":
#                             pcs = q * ct.ctn_qty
#                         else:
#                             pcs = q * ct.ctn_qty * ct.units_per_ctn
#                         agg[color][sz] += pcs

#         grand = {s: 0 for s in SIZES}
#         colors = list(agg.keys())
#         for color in colors + ["G. Total Summery"]:
#             self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=10)
#             _c(self.ws, self.row, 1, color, bold=True, center=True, fill=TITLE_FILL)
#             self.row += 1

#             headers = ["SIZE", *SIZES, "", "TOTAL"]
#             for c, h in enumerate(headers, 1):
#                 _c(self.ws, self.row, c, h, bold=True, center=True, fill=HDR_FILL)
#             self.row += 1

#             vals = ["SHIP QTY"]
#             row_total = 0
#             for s in SIZES:
#                 v = grand[s] if color == "G. Total Summery" else agg[color][s]
#                 vals.append(v)
#                 row_total += v
#             vals += ["", row_total]
#             for c, v in enumerate(vals, 1):
#                 _c(self.ws, self.row, c, v, bold=(c == 1), center=(c == 1))
#             self.row += 1

#             if color != "G. Total Summery":
#                 for s in SIZES:
#                     grand[s] += agg[color][s]
#             self.row += 1

#     def _widths(self):
#         for col, w in {"A": 12, "B": 8, "C": 6, "D": 14, "E": 14, "F": 5,
#                        "G": 5, "H": 5, "I": 5, "J": 5, "K": 5, "L": 7,
#                        "M": 8, "N": 9, "O": 9, "P": 9, "Q": 9}.items():
#             self.ws.column_dimensions[col].width = w


# # ===============================================================
# # STREAMLIT UI
# # ===============================================================
# st.set_page_config(page_title="PO → Packing List", page_icon="📦", layout="centered")
# st.title("📦 PO → Packing List Generator")
# st.caption("Upload a Destination Purchase Order PDF. Download a full packing list XLSX.")

# pdf_file = st.file_uploader("**PO PDF**", type=["pdf"])

# if pdf_file is None:
#     st.info("⬆️ Upload a PO PDF to begin.")
#     st.stop()

# if st.button("🚀 Generate Packing List", type="primary"):
#     with st.spinner("Parsing PO and generating packing list..."):
#         header, items, color_map, unknown_markers = parse_po(pdf_file.read())

#         if not items:
#             st.error("No line items found.")
#             st.stop()

#         # --- Discovery banner ---
#         detected_packs = list({it.pack for it in items})
#         # keep PACK_ORDER preference
#         detected_packs = [p for p in PACK_ORDER if p in detected_packs]

#         st.info(
#             f"📦 Pack types detected: **{', '.join(detected_packs) or 'none'}**\n\n"
#             f"🎨 Colors detected: **{', '.join(color_map.values()) or 'none'}**"
#         )

#         # --- Unknown marker warning ---
#         if unknown_markers:
#             st.warning(
#                 "⚠️ The PDF contains lines with markers we don't yet handle: "
#                 f"**{', '.join(sorted(unknown_markers))}**. "
#                 "These lines were skipped. Tell the developer to add support."
#             )

#         # --- Build and export ---
#         packs = plan_cartons(items)
#         apply_weights(packs)
#         wb = Builder(header, packs).build()

#         buf = io.BytesIO()
#         wb.save(buf)

#         st.session_state["out_bytes"]      = buf.getvalue()
#         st.session_state["header"]         = header
#         st.session_state["packs"]          = packs
#         st.session_state["color_map"]      = color_map
#         st.session_state["detected_packs"] = detected_packs
#         st.session_state["filename"]       = f"Packing_List_{header.get('po','PO')}.xlsx"


# if "out_bytes" in st.session_state:
#     header         = st.session_state["header"]
#     packs          = st.session_state["packs"]
#     color_map      = st.session_state["color_map"]
#     detected_packs = st.session_state["detected_packs"]

#     total_pcs  = sum(c.total_qty for pt in packs.values() for carts in pt.values() for c in carts)
#     total_ctns = sum(c.ctn_qty  for pt in packs.values() for carts in pt.values() for c in carts)

#     st.success(f"✅ Generated — {total_pcs:,} pcs across {total_ctns} cartons")

#     col1, col2, col3 = st.columns(3)
#     col1.metric("PO #", header.get("po", "—"))
#     col2.metric("Total Pieces", f"{total_pcs:,}")
#     col3.metric("Total Cartons", total_ctns)

#     with st.expander("🔍 Detection summary", expanded=True):
#         st.write(f"**Pack types:** {', '.join(detected_packs) or '—'}")
#         st.write(f"**Colors:** {', '.join(color_map.values()) or '—'}")

#     st.subheader("Section breakdown")
#     for pt, by_color in packs.items():
#         for color, carts in by_color.items():
#             pcs = sum(c.total_qty for c in carts)
#             st.write(f"**{pt}** · {color} · {sum(c.ctn_qty for c in carts)} cartons · {pcs} pcs")

#     st.download_button(
#         "⬇️  Download Packing List XLSX",
#         data=st.session_state["out_bytes"],
#         file_name=st.session_state["filename"],
#         mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
#         type="primary",
#     )



"""
PO PDF -> Packing List XLSX  (fully dynamic)
- Handles multiple PO PDF formats (dept 3340, 3361, ...)
- Discovers sizes, pack types, colors from the PDF
- Splits tables by (pack_type, color)
- Unknown sizes default to 0.01

Run: streamlit run app.py
"""
import io
import re
from collections import defaultdict, OrderedDict
from dataclasses import dataclass

import streamlit as st
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


# ===============================================================
# STATIC CONFIG
# ===============================================================

# Known N.W per size. Anything NOT in this map defaults to 0.01.
NW = {
    "XS": 0.510, "S": 0.520, "M": 0.550, "L": 0.560,
    "XL": 0.610, "XXL": 0.650, "XXL+": 0.690,
}

UNKNOWN_NW = 0.01
NNW_OFFSET = 0.02   # N.N.W = N.W - 0.02 for known sizes

# Carton-level addons / deductions by pack type
CARTON_ADDON = {"MULTI": 0.71, "SINGLE": 0.71, "BULK": 0.42}
NET_DEDUCT   = {"MULTI": 0.42, "SINGLE": 0.24, "BULK": 0.24}

# Display order for pack sections
PACK_ORDER = ["MULTI", "SINGLE", "BULK"]

PACK_TITLES = {
    "MULTI":  "MULTY PACK ( Y )",
    "SINGLE": "SINGLE  PACK",
    "BULK":   "BULK PACK",
}

UNIT_PREPACK = {
    "MULTI":  "184/23",
    "SINGLE": "80/40",
    "BULK":   "80/40",
}

# Pack-type keyword mapping (keyword in PDF line -> canonical)
KNOWN_PACK_MARKERS = {
    "Multi":  "MULTI",
    "Single": "SINGLE",
    "Bulk":   "BULK",
}
UNKNOWN_HINTS = ("Prepack", "Assorted", "Master", "Mixed", "Combo")


# ===============================================================
# Line-item anchor (very generic — matches any size token)
# Example matches:
#   3340 905518 Bulk Bulk 9055180210000 Bulk MR PLEATED WID 26 29 000905518-002 XS 8.39 243.31
#   3361 1183764 Bulk Bulk 323912932 Bulk SUPER BAGGY CA 32 0 001183764-000 8 6.06 0.00
# ===============================================================
ANCHOR_RE = re.compile(
    r"^(?P<dept>\d{4})\s+"                     # dept
    r"(?P<style>\d{6,7})\s+"                   # style (6 or 7 digits)
    r"(?P<rest>.*?)"                           # everything up to the CC#
    r"(?P<cc>\d{9,12}-\d{3})\s+"               # Universal CC #
    r"(?P<size>[A-Za-z0-9+]{1,6})\s+"          # size token (alpha OR numeric)
    r"(?P<cost>\d+\.\d{2})\s+"                 # unit cost
    r"(?P<total>[\d,]+\.\d{2})\s*$"            # total cost
)

COLOR_STOPWORDS = {"SIZE", "TOTAL", "COLOR", "PCS", "CTNS"}


# ===============================================================
# MODELS
# ===============================================================
@dataclass
class Item:
    pack: str
    color_code: str
    color_name: str
    size: str
    qty: int


@dataclass
class Carton:
    ctn_no: str
    ctn_mes: str
    ctn_qty: int
    color: str
    size_qty: dict
    units_per_ctn: int
    qty_per_ctn: int
    total_qty: int
    gross_wt: float = 0.0
    net_wt: float = 0.0


# ===============================================================
# HELPERS
# ===============================================================
def _search(text, pattern):
    m = re.search(pattern, text, re.I)
    return m.group(1).strip() if m else ""


def _nw(size: str) -> float:
    return NW.get(size, UNKNOWN_NW)


def _nnw(size: str) -> float:
    return NW.get(size, UNKNOWN_NW) - NNW_OFFSET if size in NW else UNKNOWN_NW


def _looks_like_color(line: str) -> bool:
    if not line:
        return False
    s = line.strip()
    if not (3 <= len(s) <= 40):
        return False
    if not s.isupper():
        return False
    if s in COLOR_STOPWORDS:
        return False
    if not all(ch.isalpha() or ch == " " for ch in s):
        return False
    if sum(ch.isalpha() for ch in s) < 3:
        return False
    if any(w in COLOR_STOPWORDS for w in s.split()):
        return False
    return True


# ===============================================================
# PDF PARSER
# ===============================================================
def parse_po(pdf_bytes: bytes):
    """
    Returns:
        header: dict
        items: list[Item]
        color_map: {color_code: name}
        unknown_markers: set[str]
        sizes_seen: list[str]  (in first-seen order)
    """
    header, items, color_map = {}, [], {}
    unknown_markers = set()
    sizes_seen = []

    def track_size(sz):
        if sz not in sizes_seen:
            sizes_seen.append(sz)

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        # ---- header ----
        t0 = pdf.pages[0].extract_text() or ""
        header["po"]          = _search(t0, r"Destination\s*Purchase\s*Order\s*#?\s*(\d{6,})")
        header["ship_before"] = _search(t0, r"Do Not Ship Before Date\s*(\S+)")
        header["ship_cancel"] = _search(t0, r"Ship Cancel Date\s*(\S+)")
        header["in_dc"]       = _search(t0, r"In DC Date\s*(\S+)")
        header["stock"]       = _search(t0, r"Planned Stock Date\s*(\S+)")
        header["dept"]        = _search(t0, r"Department:?\s*(\d{4})")

        # ---- line items ----
        for page in pdf.pages[1:]:
            text = page.extract_text() or ""
            lines = [l.rstrip() for l in text.splitlines()]

            for i, line in enumerate(lines):
                ls = line.strip()
                m = ANCHOR_RE.match(ls)
                if not m:
                    continue

                color_code = m.group("cc").split("-")[-1]
                size       = m.group("size")
                cost       = float(m.group("cost"))
                total      = float(m.group("total").replace(",", ""))
                if cost <= 0:
                    continue

                qty = round(total / cost) if total > 0 else 0

                # ---------- pack type ----------
                pt = None
                for marker, canonical in KNOWN_PACK_MARKERS.items():
                    if marker in m.group("rest"):
                        pt = canonical
                        break
                if pt is None:
                    for hint in UNKNOWN_HINTS:
                        if hint in m.group("rest"):
                            unknown_markers.add(hint)
                    continue

                # ---------- color (from next line) ----------
                color_name = color_map.get(color_code, "")
                if i + 1 < len(lines):
                    nxt = lines[i + 1].strip()
                    if _looks_like_color(nxt):
                        color_name = nxt
                        color_map[color_code] = nxt
                if not color_name:
                    color_name = f"COLOR-{color_code}"

                track_size(size)

                if qty > 0:  # skip qty=0 rows
                    items.append(Item(pt, color_code, color_name, size, qty))

    # ---- de-dup ----
    seen, clean = set(), []
    for it in items:
        k = (it.pack, it.color_code, it.size, it.qty)
        if k in seen:
            continue
        seen.add(k)
        clean.append(it)

    # ---- backfill color names ----
    for it in clean:
        if it.color_name.startswith("COLOR-") and it.color_code in color_map:
            it.color_name = color_map[it.color_code]

    return header, clean, color_map, unknown_markers, sizes_seen


# ===============================================================
# PLANNING
# ===============================================================
def plan_cartons(items):
    grouped = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
    for it in items:
        grouped[it.pack][it.color_name][it.size] += it.qty

    # Preserve order: MULTI -> SINGLE -> BULK, then color first-seen
    result = OrderedDict()
    for pt in PACK_ORDER:
        if pt not in grouped:
            continue
        result[pt] = OrderedDict()
        for color in grouped[pt]:
            sizes = grouped[pt][color]
            if pt == "MULTI":
                result[pt][color] = _plan_multi(color, sizes)
            elif pt == "SINGLE":
                result[pt][color] = _plan_single(color, sizes)
            else:
                result[pt][color] = _plan_bulk(color, sizes)
    return result


def _plan_multi(color, sizes):
    total = sum(sizes.values())
    n = total // 7
    if n == 0:
        return []
    per = {s: sizes.get(s, 0) // n for s in sizes}
    return [Carton("1", "G82", n, color, per, 7, 7, n * 7)]


def _plan_single(color, sizes):
    raw = []
    for sz, pcs in sizes.items():
        while pcs > 0:
            take = min(20, pcs)
            pcs -= take
            raw.append(Carton(
                "", "G81" if take == 20 else "G82", 1, color,
                {sz: take // 2}, 2, take, take,
            ))
    return _group_cartons(raw)


def _plan_bulk(color, sizes):
    raw = []
    for sz, pcs in sizes.items():
        while pcs > 0:
            take = min(20, pcs)
            pcs -= take
            raw.append(Carton(
                "", "G81" if take == 20 else "G82", 1, color,
                {sz: take}, 1, take, take,
            ))
    return _group_cartons(raw)


def _group_cartons(raw):
    if not raw:
        return []
    groups = []
    cur = raw[0]
    for c in raw[1:]:
        if c.size_qty == cur.size_qty and c.ctn_mes == cur.ctn_mes:
            cur.ctn_qty += 1
            cur.total_qty += c.total_qty
        else:
            groups.append(cur)
            cur = c
    groups.append(cur)
    for i, g in enumerate(groups, 1):
        g.ctn_no = str(i) if g.ctn_qty == 1 else f"{i}-{i + g.ctn_qty - 1}"
    return groups


def apply_weights(packs):
    for pt, by_color in packs.items():
        addon  = CARTON_ADDON[pt]
        deduct = NET_DEDUCT[pt]
        for color, carts in by_color.items():
            for c in carts:
                if pt == "SINGLE":
                    nw_sum = sum(_nw(s) * q * c.units_per_ctn
                                 for s, q in c.size_qty.items())
                else:
                    nw_sum = sum(_nw(s) * q for s, q in c.size_qty.items())
                c.gross_wt = round(nw_sum + addon, 2)
                c.net_wt   = round(c.gross_wt - deduct, 2)


# ===============================================================
# XLSX WRITER
# ===============================================================
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE_FILL = PatternFill("solid", fgColor="FCE4D6")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def _c(ws, r, c, v, bold=False, center=False, fill=None):
    cell = ws.cell(row=r, column=c, value=v)
    if bold: cell.font = BOLD
    cell.alignment = CENTER if center else (RIGHT if isinstance(v, (int, float)) else LEFT)
    if fill: cell.fill = fill
    cell.border = BORDER


class Builder:
    def __init__(self, header, packs, sizes):
        self.header = header
        self.packs = packs
        self.sizes = sizes           # discovered sizes, in order
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "Packing List"
        self.row = 1

    def build(self):
        self._company_header()
        self._global_table()
        for pt in PACK_ORDER:
            if pt not in self.packs:
                continue
            for color, carts in self.packs[pt].items():
                self._pack_table(pt, color, carts)
        self._summary()
        self._widths()
        return self.wb

    def _company_header(self):
        self.ws.merge_cells("A1:Q1")
        self.ws["A1"] = "CREATIVE COLLECTIONS LTD-1A."
        self.ws["A1"].font = Font(bold=True, size=14)
        self.ws["A1"].alignment = CENTER
        self.ws.merge_cells("A2:Q2")
        self.ws["A2"] = "Nishat Nagar , Tongi , Gazipur ."
        self.ws["A2"].font = Font(bold=True, size=12)
        self.ws["A2"].alignment = CENTER
        self.row = 4

    def _global_table(self):
        rows = [
            ["SIZE", "", "SIZE", *self.sizes],
            ["N.W",   "", "N.WT",   *[_nw(s)  for s in self.sizes]],
            ["N.N.W", "", "N.N. WT", *[_nnw(s) for s in self.sizes]],
            ["EMPTY CTN WET", "", "", *[""] * len(self.sizes)],
        ]
        for i, rd in enumerate(rows):
            for c, v in enumerate(rd, 1):
                _c(self.ws, self.row + i, c, v, bold=True, center=True, fill=HDR_FILL)
        self.row += 6

    def _pack_table(self, pt, color, carts):
        title = PACK_TITLES.get(pt, pt)
        total_pcs  = sum(c.total_qty for c in carts)
        total_ctns = sum(c.ctn_qty for c in carts)
        po = self.header.get("po", "")

        self.ws.merge_cells(start_row=self.row, start_column=1,
                            end_row=self.row, end_column=17)
        _c(self.ws, self.row, 1, title, bold=True, center=True, fill=TITLE_FILL)
        self.row += 1

        info = [
            ("BUYER", ":", "OLD NAVY"),
            ("STYLE", ":", self.header.get("dept", "")),
            ("P. O. #", ":", po),
            ("O/QTY", ":", total_pcs, "PCS"),
            ("SHIP QTY", ":", total_pcs, "PCS"),
            ("EX/SHORT", ":", 0, "PCS"),
            ("CTN QTY", ":", total_ctns, "CTNS"),
            ("CTN MEAS.", ":", "58.67 X 38.48 X 29.71 CM.G8_SL"),
            ("CTN MEAS.", ":", "58.67 X 38.48 X 14.86 CM.G8S_SL"),
            ("CTN MEAS.", ":", "38.48 X 29.33 X 14.86 CM.G-8M"),
            ("CTN MEAS.", ":", "29.33 X 19.25 X 14.86 CM.G-8XS"),
        ]
        for i, tup in enumerate(info):
            for c, v in enumerate(tup, 1):
                _c(self.ws, self.row + i, c, v, bold=(c <= 2))

        right = [
            ("PO #", po),
            ("SKU / Item", self.header.get("dept", "")),
            ("Unit / Prepack", UNIT_PREPACK.get(pt, "")),
            ("CARTON", f"01 of {total_ctns}"),
        ]
        for i, (k, v) in enumerate(right):
            _c(self.ws, self.row + i, 13, k, bold=True)
            _c(self.ws, self.row + i, 15, v)

        self.row += len(info) + 2

        if pt == "MULTI":
            headers = ["CTN NO", "CTN\nMES:", "CTN\nQTY", "COLOR", "PREPACK\nSTECKER",
                       *self.sizes, "PER\nBLST", "QTY\nPER\nCTN", "TOTAL\nQTY",
                       "GROSS\nWEIGHT", "NET\nWEIGHT"]
        else:
            headers = ["CTN NO", "CTN\nMES:", "CTN\nQTY", "COLOR", "PREPACK\nSTECKER",
                       *self.sizes, "QTY\nPER\nCTN", "TOTAL\nQTY",
                       "GROSS\nWEIGHT", "NET\nWEIGHT"]
        for c, h in enumerate(headers, 1):
            _c(self.ws, self.row, c, h, bold=True, center=True, fill=HDR_FILL)
        self.row += 1

        size_totals = {s: 0 for s in self.sizes}
        for ct in carts:
            vals = [ct.ctn_no, ct.ctn_mes, ct.ctn_qty, ct.color, "",
                    *[ct.size_qty.get(s, "") for s in self.sizes]]
            if pt == "MULTI":
                vals += ["7 X 1", ct.qty_per_ctn, ct.total_qty, ct.gross_wt, ct.net_wt]
            else:
                vals += [ct.qty_per_ctn, ct.total_qty, ct.gross_wt, ct.net_wt]
            for i, v in enumerate(vals, 1):
                _c(self.ws, self.row, i, v, center=(i <= 5))
            for s in self.sizes:
                size_totals[s] += (ct.size_qty.get(s, 0) or 0) * ct.ctn_qty
            self.row += 1

        if pt == "MULTI":
            tot = ["TOTAL", "", total_ctns, "CTNS", "",
                   *[size_totals[s] for s in self.sizes],
                   "", "", total_pcs, "PCS", ""]
        else:
            tot = ["TOTAL", "", total_ctns, "CTN", "",
                   *[size_totals[s] for s in self.sizes],
                   "", total_pcs, "PCS", ""]
        for c, v in enumerate(tot, 1):
            _c(self.ws, self.row, c, v, bold=True, center=True, fill=HDR_FILL)
        self.row += 2

    def _summary(self):
        agg = defaultdict(lambda: {s: 0 for s in self.sizes})
        for pt, by_color in self.packs.items():
            for color, carts in by_color.items():
                for ct in carts:
                    for sz, q in ct.size_qty.items():
                        if pt == "MULTI":
                            pcs = q * ct.ctn_qty
                        else:
                            pcs = q * ct.ctn_qty * ct.units_per_ctn
                        agg[color][sz] += pcs

        grand = {s: 0 for s in self.sizes}
        colors = list(agg.keys())
        for color in colors + ["G. Total Summery"]:
            self.ws.merge_cells(start_row=self.row, start_column=1,
                                end_row=self.row, end_column=10)
            _c(self.ws, self.row, 1, color, bold=True, center=True, fill=TITLE_FILL)
            self.row += 1
            headers = ["SIZE", *self.sizes, "", "TOTAL"]
            for c, h in enumerate(headers, 1):
                _c(self.ws, self.row, c, h, bold=True, center=True, fill=HDR_FILL)
            self.row += 1
            vals = ["SHIP QTY"]; row_total = 0
            for s in self.sizes:
                v = grand[s] if color == "G. Total Summery" else agg[color][s]
                vals.append(v); row_total += v
            vals += ["", row_total]
            for c, v in enumerate(vals, 1):
                _c(self.ws, self.row, c, v, bold=(c == 1), center=(c == 1))
            self.row += 1
            if color != "G. Total Summery":
                for s in self.sizes:
                    grand[s] += agg[color][s]
            self.row += 1

    def _widths(self):
        for col, w in {"A": 12, "B": 8, "C": 6, "D": 14, "E": 14, "F": 5,
                       "G": 5, "H": 5, "I": 5, "J": 5, "K": 5, "L": 7,
                       "M": 8, "N": 9, "O": 9, "P": 9, "Q": 9}.items():
            self.ws.column_dimensions[col].width = w


# ===============================================================
# STREAMLIT UI
# ===============================================================
st.set_page_config(page_title="PO → Packing List", page_icon="📦", layout="centered")
st.title("📦 PO → Packing List Generator")
st.caption("Upload a Destination Purchase Order PDF. Download a full packing list XLSX.")

pdf_file = st.file_uploader("**PO PDF**", type=["pdf"])

if pdf_file is None:
    st.info("⬆️ Upload a PO PDF to begin.")
    st.stop()

if st.button("🚀 Generate Packing List", type="primary"):
    with st.spinner("Parsing PO and generating packing list..."):
        header, items, color_map, unknown_markers, sizes_seen = parse_po(pdf_file.read())

        if not items:
            st.error("No line items found. Check the PDF format.")
            st.stop()

        detected_packs = [p for p in PACK_ORDER if any(it.pack == p for it in items)]

        st.info(
            f"📦 Pack types detected: **{', '.join(detected_packs) or 'none'}**\n\n"
            f"🎨 Colors detected: **{', '.join(color_map.values()) or 'none'}**\n\n"
            f"📏 Sizes detected: **{', '.join(sizes_seen) or 'none'}**"
        )

        if unknown_markers:
            st.warning(
                "⚠️ Unknown pack-type markers seen in PDF: "
                f"**{', '.join(sorted(unknown_markers))}**. Those lines were skipped."
            )

        packs = plan_cartons(items)
        apply_weights(packs)

        # Attach sizes list to header for later reference
        header["sizes"] = sizes_seen

        wb = Builder(header, packs, sizes_seen).build()
        buf = io.BytesIO()
        wb.save(buf)

        st.session_state["out_bytes"]      = buf.getvalue()
        st.session_state["header"]         = header
        st.session_state["packs"]          = packs
        st.session_state["color_map"]      = color_map
        st.session_state["sizes_seen"]     = sizes_seen
        st.session_state["detected_packs"] = detected_packs
        st.session_state["filename"]       = f"Packing_List_{header.get('po','PO')}.xlsx"


if "out_bytes" in st.session_state:
    header         = st.session_state["header"]
    packs          = st.session_state["packs"]
    color_map      = st.session_state["color_map"]
    sizes_seen     = st.session_state["sizes_seen"]
    detected_packs = st.session_state["detected_packs"]

    total_pcs  = sum(c.total_qty for pt in packs.values() for carts in pt.values() for c in carts)
    total_ctns = sum(c.ctn_qty  for pt in packs.values() for carts in pt.values() for c in carts)

    st.success(f"✅ Generated — {total_pcs:,} pcs across {total_ctns} cartons")

    col1, col2, col3 = st.columns(3)
    col1.metric("PO #", header.get("po", "—"))
    col2.metric("Total Pieces", f"{total_pcs:,}")
    col3.metric("Total Cartons", total_ctns)

    with st.expander("🔍 Detection summary", expanded=True):
        st.write(f"**Pack types:** {', '.join(detected_packs) or '—'}")
        st.write(f"**Colors:** {', '.join(color_map.values()) or '—'}")
        st.write(f"**Sizes:** {', '.join(sizes_seen) or '—'}")

    st.subheader("Section breakdown")
    for pt, by_color in packs.items():
        for color, carts in by_color.items():
            pcs = sum(c.total_qty for c in carts)
            st.write(f"**{pt}** · {color} · {sum(c.ctn_qty for c in carts)} cartons · {pcs} pcs")

    st.download_button(
        "⬇️  Download Packing List XLSX",
        data=st.session_state["out_bytes"],
        file_name=st.session_state["filename"],
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )