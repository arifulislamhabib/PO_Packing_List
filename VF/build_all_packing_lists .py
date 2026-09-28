"""
Build a packing-list-style sheet for EVERY PO line found in
po_extracted.json -- one worksheet per line, in a single workbook.

HOW TO RUN (after extract_po.py has created po_extracted.json):
    python3 build_all_packing_lists.py

Requires:
    pip install openpyxl
"""
import json
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

with open("po_extracted.json") as f:
    data = json.load(f)

FONT = "Arial"
bold = Font(name=FONT, bold=True, size=10)
normal = Font(name=FONT, size=10)
title_font = Font(name=FONT, bold=True, size=14, color="1F4E79")
subtitle_font = Font(name=FONT, bold=True, size=11, color="2E75B6")
header_font = Font(name=FONT, bold=True, size=9, color="1F4E79")
label_font = Font(name=FONT, bold=True, size=10, color="333333")
thin = Side(style="thin", color="B4C6E7")
medium = Side(style="medium", color="5B9BD5")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
header_box = Border(left=thin, right=thin, top=medium, bottom=medium)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_center = Alignment(horizontal="left", vertical="center", wrap_text=True)

# Light color palette
header_fill = PatternFill("solid", fgColor="D6E3F0")      # soft blue header
alt_row_fill = PatternFill("solid", fgColor="F2F7FB")     # very light blue alt rows
spec_fill = PatternFill("solid", fgColor="E2EFDA")        # soft green for size-spec
yellow = PatternFill("solid", fgColor="FFF2CC")           # soft yellow for inputs
title_fill = PatternFill("solid", fgColor="D6E3F0")
total_fill = PatternFill("solid", fgColor="DDEBF7")       # light blue for totals
summary_header_fill = PatternFill("solid", fgColor="C6EFCE")  # soft green summary


def sheet_name_for(line):
    raw = f"{line['style']}"
    return re.sub(r'[\[\]:*?/\\]', '-', raw)[:31]


def compute_carton_rows(sublines):
    """
    Splits each subline's qty into full-size cartons (qty // pack) plus a
    shared remainder pool that gets packed into mixed cartons, filled in
    size order until each hits pack capacity; the final leftover carton
    may be under capacity.
    Returns a list of dicts: {"sizes": {size: qty, ...}, "total_ctn": int,
    "ctn_pcs": int, "color": str or None, "upc": str or None, "mixed": bool}
    """
    if not sublines:
        return []
    pack = sublines[0].get("items_per_outer_pack") or 1

    full_rows = []
    remainder_pool = []  # [size, qty_left, color, upc]
    for s in sublines:
        qty = s["qty"]
        size = s["size"]
        full_ctn = qty // pack
        rem = qty % pack
        if full_ctn > 0:
            full_rows.append({
                "sizes": {size: pack},
                "total_ctn": full_ctn,
                "ctn_pcs": pack,
                "color": s["color"],
                "upc": s["upc"],
                "mixed": False,
            })
        if rem > 0:
            remainder_pool.append([size, rem, s["color"], s["upc"]])

    mixed_rows = []
    current = {}
    current_total = 0
    idx = 0
    while idx < len(remainder_pool):
        size, rem, color, upc = remainder_pool[idx]
        space_left = pack - current_total
        take = min(rem, space_left)
        if take > 0:
            current[size] = current.get(size, 0) + take
            current_total += take
            remainder_pool[idx][1] -= take
        if remainder_pool[idx][1] == 0:
            idx += 1
        if current_total == pack:
            mixed_rows.append({
                "sizes": dict(current), "total_ctn": 1, "ctn_pcs": pack,
                "color": None, "upc": None, "mixed": True,
            })
            current = {}
            current_total = 0
    if current:
        mixed_rows.append({
            "sizes": dict(current), "total_ctn": 1, "ctn_pcs": current_total,
            "color": None, "upc": None, "mixed": True,
        })

    return full_rows + mixed_rows


def style_cell(cell, font=None, fill=None, border=None, alignment=None):
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if border is not None:
        cell.border = border
    if alignment is not None:
        cell.alignment = alignment


def build_sheet(ws, line, sublines):
    sizes = [s["size"] for s in sublines]
    n_sizes = len(sizes)

    # ---- Title block ----
    ws.merge_cells("A1:H1")
    ws["A1"] = "CREATIVE COLLECTIONS LTD. U-1-A"
    style_cell(ws["A1"], font=title_font, fill=title_fill, alignment=center)
    ws.row_dimensions[1].height = 22

    ws.merge_cells("A2:H2")
    ws["A2"] = "PACKING LIST DETAILS"
    style_cell(ws["A2"], font=subtitle_font, fill=title_fill, alignment=center)
    ws.row_dimensions[2].height = 18

    # ---- Small size-spec table: SIZE / N.W. / N.N.W. / EMPTY CTN ----
    spec_header_row = 4
    c = ws.cell(row=spec_header_row, column=1, value="SIZE")
    style_cell(c, font=header_font, fill=spec_fill, border=box, alignment=center)
    for sc_idx, size_name in enumerate(sizes):
        c = ws.cell(row=spec_header_row, column=2 + sc_idx, value=size_name)
        style_cell(c, font=header_font, fill=spec_fill, border=box, alignment=center)

    spec_rows = ["N.W.", "N.N.W.", "EMPTY CTN"]
    for i, label in enumerate(spec_rows):
        rr = spec_header_row + 1 + i
        lc = ws.cell(row=rr, column=1, value=label)
        style_cell(lc, font=bold, fill=spec_fill, border=box, alignment=center)
        for sc_idx, size_name in enumerate(sizes):
            c = ws.cell(row=rr, column=2 + sc_idx, value=1)
            style_cell(c, font=normal, border=box, alignment=center)

    # ---- Info block (horizontal layout) ----
    r = spec_header_row + len(spec_rows) + 2

    # Row 1 labels
    info_labels_1 = ["BUYER", "LOT NO", "P.O NO", "ORDER QTY", "PACK QTY", "EXCESS/SHORT QTY", "PERCENTAGE"]
    for col_idx, label in enumerate(info_labels_1, start=1):
        c = ws.cell(row=r, column=col_idx, value=label)
        style_cell(c, font=header_font, fill=header_fill, border=box, alignment=center)
    r += 1

    # Row 1 values
    info_vals_1 = [
        "VANS",
        line["style"],
        line["po_line_no"],
        f'{line["order_qty"]} PCS',
        f'{line["order_qty"]} PCS',
        None,  # EXCESS/SHORT – filled later with formula
        None,  # PERCENTAGE – filled later with formula
    ]
    excess_short_row = r
    percentage_col = 7  # column for PERCENTAGE value
    excess_col = 6      # column for EXCESS/SHORT value
    for col_idx, val in enumerate(info_vals_1, start=1):
        c = ws.cell(row=r, column=col_idx, value=val)
        style_cell(c, font=normal, border=box, alignment=center)
    percentage_row = r
    r += 1

    # Row 2 labels
    info_labels_2 = ["CRD", "COUNTRY", "DESCRIPTION"]
    for col_idx, label in enumerate(info_labels_2, start=1):
        c = ws.cell(row=r, column=col_idx, value=label)
        style_cell(c, font=header_font, fill=header_fill, border=box, alignment=center)
    r += 1

    # Row 2 values
    info_vals_2 = [
        line.get("crd"),
        line.get("destination_country"),
        line["description"],
    ]
    for col_idx, val in enumerate(info_vals_2, start=1):
        c = ws.cell(row=r, column=col_idx, value=val)
        style_cell(c, font=normal, border=box, alignment=center)
        if col_idx == 3:
            # DESCRIPTION can be long – merge a few columns
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
            c.alignment = center
    r += 1

    # ---- Main data header ----
    header_row = r + 2
    headers = ["CASE LABEL NO.", "CARTON NO.", "COLOR", "UPC Number"] + sizes + \
              ["CTN PCS", "TOTAL CTN", "TOTAL PCS", "Grs.wt.pr ctn", "Net.wt.pr ctn",
               "total.Grs.wt", "total.net.wt"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        style_cell(c, font=header_font, fill=header_fill, border=header_box, alignment=center)
    n_cols = len(headers)
    ws.row_dimensions[header_row].height = 30

    # ---- Data rows (carton rows) ----
    data_row_start = header_row + 1
    carton_rows = compute_carton_rows(sublines)
    for idx, row_data in enumerate(carton_rows):
        rr = data_row_start + idx
        row_fill = alt_row_fill if idx % 2 == 1 else None

        c = ws.cell(row=rr, column=1, value=idx + 1)
        style_cell(c, font=normal, fill=row_fill, border=box, alignment=center)

        # CARTON NO. set after TOTAL CTN column is known (cumulative formula)
        color_val = row_data["color"] if row_data["color"] is not None else "MIXED"
        color_cell = ws.cell(row=rr, column=3, value=color_val)
        style_cell(color_cell, font=normal, border=box, alignment=center,
                   fill=yellow if row_data["mixed"] else row_fill)

        upc_cell = ws.cell(row=rr, column=4, value=row_data["upc"])
        style_cell(upc_cell, font=normal, fill=row_fill, border=box, alignment=center)

        for sc_idx, size_name in enumerate(sizes):
            col = 5 + sc_idx
            val = row_data["sizes"].get(size_name)
            cell = ws.cell(row=rr, column=col, value=val)
            style_cell(cell, font=normal, fill=row_fill, border=box, alignment=center)

        # CTN PCS
        ctn_pcs_col = 5 + n_sizes
        first_size_letter = get_column_letter(5)
        last_size_letter = get_column_letter(4 + n_sizes)
        ctn_cell = ws.cell(
            row=rr, column=ctn_pcs_col,
            value=f"=SUM({first_size_letter}{rr}:{last_size_letter}{rr})"
        )
        style_cell(ctn_cell, font=bold, fill=row_fill, border=box, alignment=center)

        total_ctn_col = ctn_pcs_col + 1
        total_pcs_col = ctn_pcs_col + 2
        total_ctn_letter = get_column_letter(total_ctn_col)

        tc_cell = ws.cell(row=rr, column=total_ctn_col, value=row_data["total_ctn"])
        style_cell(tc_cell, font=normal, fill=row_fill, border=box, alignment=center)

        # CARTON NO. = cumulative sum of TOTAL CTN (running total)
        if idx == 0:
            carton_formula = f"={total_ctn_letter}{rr}"
        else:
            carton_formula = f"={get_column_letter(2)}{rr - 1}+{total_ctn_letter}{rr}"
        cB = ws.cell(row=rr, column=2, value=carton_formula)
        style_cell(cB, font=normal, fill=row_fill, border=box, alignment=center)

        ctn_pcs_letter = get_column_letter(ctn_pcs_col)
        tp_cell = ws.cell(
            row=rr, column=total_pcs_col,
            value=f"={total_ctn_letter}{rr}*{ctn_pcs_letter}{rr}"
        )
        style_cell(tp_cell, font=normal, fill=row_fill, border=box, alignment=center)

        # Grs.wt.pr ctn
        grs_wt_col = total_pcs_col + 1
        nw_row = spec_header_row + 1
        empty_row = spec_header_row + 3
        first_spec_letter = get_column_letter(2)
        last_spec_letter = get_column_letter(1 + n_sizes)
        first_main_size_letter = get_column_letter(5)
        last_main_size_letter = get_column_letter(4 + n_sizes)

        grs_formula = (
            f"=SUMPRODUCT({first_main_size_letter}{rr}:{last_main_size_letter}{rr},"
            f"{first_spec_letter}{nw_row}:{last_spec_letter}{nw_row})"
            f"+{first_spec_letter}{empty_row}"
        )
        grs_cell = ws.cell(row=rr, column=grs_wt_col, value=grs_formula)
        style_cell(grs_cell, font=normal, fill=row_fill, border=box, alignment=center)

        # Net.wt.pr ctn
        net_wt_col = grs_wt_col + 1
        grs_wt_letter_for_net = get_column_letter(grs_wt_col)
        net_cell = ws.cell(
            row=rr, column=net_wt_col,
            value=f"={grs_wt_letter_for_net}{rr}-{first_spec_letter}{empty_row}"
        )
        style_cell(net_cell, font=normal, fill=row_fill, border=box, alignment=center)

        # total.Grs.wt / total.net.wt
        total_grs_col = net_wt_col + 1
        grs_wt_letter = get_column_letter(grs_wt_col)
        tg_cell = ws.cell(
            row=rr, column=total_grs_col,
            value=f"={total_ctn_letter}{rr}*{grs_wt_letter}{rr}"
        )
        style_cell(tg_cell, font=normal, fill=row_fill, border=box, alignment=center)

        total_net_col = total_grs_col + 1
        net_wt_letter = get_column_letter(net_wt_col)
        tn_cell = ws.cell(
            row=rr, column=total_net_col,
            value=f"={total_ctn_letter}{rr}*{net_wt_letter}{rr}"
        )
        style_cell(tn_cell, font=normal, fill=row_fill, border=box, alignment=center)

    # ---- TOTAL row ----
    total_row = data_row_start + len(carton_rows)
    total_label = ws.cell(row=total_row, column=3, value="TOTAL")
    style_cell(total_label, font=bold, fill=total_fill, alignment=center)

    total_ctn_letter_for_sum = get_column_letter(total_ctn_col)
    for sc_idx, size_name in enumerate(sizes):
        col = 5 + sc_idx
        col_letter = get_column_letter(col)
        cell = ws.cell(
            row=total_row, column=col,
            value=f"=SUMPRODUCT({col_letter}{data_row_start}:{col_letter}{total_row-1},"
                  f"{total_ctn_letter_for_sum}{data_row_start}:{total_ctn_letter_for_sum}{total_row-1})"
        )
        style_cell(cell, font=bold, fill=total_fill, border=box, alignment=center)

    total_ctn_col = ctn_pcs_col + 1
    total_pcs_col = ctn_pcs_col + 2
    total_ctn_letter = get_column_letter(total_ctn_col)
    total_pcs_letter = get_column_letter(total_pcs_col)

    for col, formula in [
        (total_ctn_col,
         f'=IF(COUNT({total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1})=0,"",'
         f'SUM({total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1}))'),
        (total_pcs_col,
         f'=IF(COUNT({total_pcs_letter}{data_row_start}:{total_pcs_letter}{total_row-1})=0,"",'
         f'SUM({total_pcs_letter}{data_row_start}:{total_pcs_letter}{total_row-1}))'),
    ]:
        c = ws.cell(row=total_row, column=col, value=formula)
        style_cell(c, font=bold, fill=total_fill, border=box, alignment=center)

    grs_wt_col = total_pcs_col + 1
    net_wt_col = total_pcs_col + 2
    grs_wt_letter = get_column_letter(grs_wt_col)
    net_wt_letter = get_column_letter(net_wt_col)

    # SUMPRODUCT (not SUM): these are PER-CARTON weights, so a row with
    # TOTAL CTN=2 must count double toward the grand total.
    c = ws.cell(
        row=total_row, column=grs_wt_col,
        value=f"=SUMPRODUCT({grs_wt_letter}{data_row_start}:{grs_wt_letter}{total_row-1},"
              f"{total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1})"
    )
    style_cell(c, font=bold, fill=total_fill, border=box, alignment=center)
    c = ws.cell(
        row=total_row, column=net_wt_col,
        value=f"=SUMPRODUCT({net_wt_letter}{data_row_start}:{net_wt_letter}{total_row-1},"
              f"{total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1})"
    )
    style_cell(c, font=bold, fill=total_fill, border=box, alignment=center)

    total_grs_col = total_pcs_col + 3
    total_grs_letter = get_column_letter(total_grs_col)
    c = ws.cell(row=total_row, column=total_grs_col,
                value=f'=IF(COUNT({total_grs_letter}{data_row_start}:{total_grs_letter}{total_row-1})=0,"",'
                      f'SUM({total_grs_letter}{data_row_start}:{total_grs_letter}{total_row-1}))')
    style_cell(c, font=bold, fill=total_fill, border=box, alignment=center)

    total_net_col = total_grs_col + 1
    total_net_letter = get_column_letter(total_net_col)
    c = ws.cell(row=total_row, column=total_net_col,
                value=f'=IF(COUNT({total_net_letter}{data_row_start}:{total_net_letter}{total_row-1})=0,"",'
                      f'SUM({total_net_letter}{data_row_start}:{total_net_letter}{total_row-1}))')
    style_cell(c, font=bold, fill=total_fill, border=box, alignment=center)

    # Apply fill/border to empty total-row cells for visual consistency
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=total_row, column=col)
        if cell.fill.fgColor is None or cell.fill.fgColor.rgb == "00000000":
            cell.fill = total_fill
        if cell.border.left.style is None:
            cell.border = box
        cell.alignment = center

    # ---- Summary block (SIZE / CUT QTY / ORDER QTY ...) ----
    sum_header_row = total_row + 3
    c = ws.cell(row=sum_header_row, column=1, value="SIZE")
    style_cell(c, font=header_font, fill=summary_header_fill, border=box, alignment=center)
    for sc_idx, size_name in enumerate(sizes):
        c = ws.cell(row=sum_header_row, column=2 + sc_idx, value=size_name)
        style_cell(c, font=header_font, fill=summary_header_fill, border=box, alignment=center)
    c = ws.cell(row=sum_header_row, column=2 + n_sizes, value="G TOTAL")
    style_cell(c, font=header_font, fill=summary_header_fill, border=box, alignment=center)

    gw_col = 3 + n_sizes
    nw_col = 4 + n_sizes
    ws.merge_cells(start_row=sum_header_row, start_column=gw_col,
                    end_row=sum_header_row, end_column=gw_col)
    c = ws.cell(row=sum_header_row, column=gw_col, value="Grand Total gross weight (kg)")
    style_cell(c, font=header_font, fill=summary_header_fill, border=box, alignment=center)
    ws.merge_cells(start_row=sum_header_row, start_column=nw_col,
                    end_row=sum_header_row, end_column=nw_col)
    c = ws.cell(row=sum_header_row, column=nw_col, value="Grand Total net weight (kg)")
    style_cell(c, font=header_font, fill=summary_header_fill, border=box, alignment=center)

    rows_needed = ["CUT QTY", "ORDER QTY", "SHIP QTY", "EXS/SHT QTY", "PERCENTAGE"]
    for i, label in enumerate(rows_needed):
        rr = sum_header_row + 1 + i
        lc = ws.cell(row=rr, column=1, value=label)
        style_cell(lc, font=bold, border=box, alignment=center)
        for sc_idx, size_name in enumerate(sizes):
            col = 2 + sc_idx
            col_letter = get_column_letter(col)
            cell = ws.cell(row=rr, column=col)
            if label == "ORDER QTY":
                qty = next(s["qty"] for s in sublines if s["size"] == size_name)
                cell.value = qty
            elif label == "CUT QTY":
                cell.fill = yellow
            elif label == "SHIP QTY":
                # SHIP QTY = total packed qty per size from main table TOTAL row
                # (SUMPRODUCT of size qty × TOTAL CTN already computed there)
                main_size_col = 5 + sc_idx
                main_size_letter = get_column_letter(main_size_col)
                cell.value = f"={main_size_letter}{total_row}"
            elif label == "EXS/SHT QTY":
                ship_row = sum_header_row + 1 + rows_needed.index("SHIP QTY")
                order_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
                cell.value = f'=IF({col_letter}{ship_row}="","",{col_letter}{ship_row}-{col_letter}{order_row})'
            elif label == "PERCENTAGE":
                exs_row = sum_header_row + 1 + rows_needed.index("EXS/SHT QTY")
                order_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
                cell.value = f'=IF({col_letter}{exs_row}="","",IFERROR({col_letter}{exs_row}/{col_letter}{order_row},0))'
                cell.number_format = "0.00%"
            style_cell(cell, font=normal, border=box, alignment=center)
            if label == "CUT QTY":
                cell.fill = yellow

        gt_col = 2 + n_sizes
        gt_letter = get_column_letter(gt_col)
        first_letter = get_column_letter(2)
        last_letter = get_column_letter(1 + n_sizes)
        if label in ("CUT QTY", "EXS/SHT QTY"):
            c = ws.cell(row=rr, column=gt_col,
                        value=f'=IF(COUNT({first_letter}{rr}:{last_letter}{rr})=0,"",SUM({first_letter}{rr}:{last_letter}{rr}))')
            style_cell(c, font=bold, border=box, alignment=center)
        elif label in ("ORDER QTY", "SHIP QTY"):
            c = ws.cell(row=rr, column=gt_col,
                        value=f"=SUM({first_letter}{rr}:{last_letter}{rr})")
            style_cell(c, font=bold, border=box, alignment=center)
        else:  # PERCENTAGE
            exs_row = sum_header_row + 1 + rows_needed.index("EXS/SHT QTY")
            order_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
            c = ws.cell(row=rr, column=gt_col,
                        value=f'=IF({gt_letter}{exs_row}="","",IFERROR({gt_letter}{exs_row}/{gt_letter}{order_row},0))')
            c.number_format = "0.00%"
            style_cell(c, font=bold, border=box, alignment=center)

    # Link top header EXCESS/SHORT & PERCENTAGE (horizontal info row)
    exs_gt_row = sum_header_row + 1 + rows_needed.index("EXS/SHT QTY")
    order_gt_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
    gt_col_letter = get_column_letter(2 + n_sizes)

    c = ws.cell(row=excess_short_row, column=excess_col,
                value=f'={gt_col_letter}{exs_gt_row}')
    style_cell(c, font=normal, border=box, alignment=center)
    c = ws.cell(row=percentage_row, column=percentage_col,
                value=f'=IF({gt_col_letter}{exs_gt_row}="","-",'
                      f'TEXT(IFERROR({gt_col_letter}{exs_gt_row}/{gt_col_letter}{order_gt_row},0),"0.00%"))')
    style_cell(c, font=normal, border=box, alignment=center)

    # ---- Weight / CBM block ----
    # GROSS WEIGHT = Grand Total gross weight (kg)
    # NET WEIGHT   = Grand Total net weight (kg)
    # CTN MEAS     = 2 yellow user-input lines
    # CBM          = (yellow numeric input) × TOTAL CTN
    wt_row = sum_header_row + len(rows_needed) + 2
    total_grs_letter = get_column_letter(total_grs_col)
    total_net_letter = get_column_letter(total_net_col)
    total_ctn_letter_final = get_column_letter(total_ctn_col)

    # GROSS WEIGHT (auto)
    lc = ws.cell(row=wt_row, column=1, value="GROSS WEIGHT :")
    style_cell(lc, font=label_font, alignment=center)
    vc = ws.cell(row=wt_row, column=2, value=f"={total_grs_letter}{total_row}")
    style_cell(vc, font=bold, border=box, alignment=center)

    # NET WEIGHT (auto)
    lc = ws.cell(row=wt_row + 1, column=1, value="NET WEIGHT :")
    style_cell(lc, font=label_font, alignment=center)
    vc = ws.cell(row=wt_row + 1, column=2, value=f"={total_net_letter}{total_row}")
    style_cell(vc, font=bold, border=box, alignment=center)

    # CTN MEAS : user input 1 + user input 2
    lc = ws.cell(row=wt_row + 2, column=1, value="CTN MEAS :")
    style_cell(lc, font=label_font, alignment=center)
    vc1 = ws.cell(row=wt_row + 2, column=2, value=None)  # user input 1
    style_cell(vc1, fill=yellow, border=box, alignment=center)
    vc2 = ws.cell(row=wt_row + 3, column=2, value=None)  # user input 2
    style_cell(vc2, fill=yellow, border=box, alignment=center)

    # CBM : (yellow numeric user input) × TOTAL CTN
    # Column B = factor (user enters CBM per carton or multiplier)
    # Column C = formula = B × grand TOTAL CTN
    lc = ws.cell(row=wt_row + 4, column=1, value="CBM :")
    style_cell(lc, font=label_font, alignment=center)
    cbm_input = ws.cell(row=wt_row + 4, column=2, value=None)  # user numeric input
    style_cell(cbm_input, fill=yellow, border=box, alignment=center)
    cbm_result = ws.cell(
        row=wt_row + 4, column=3,
        value=f'=IF(B{wt_row + 4}="","",B{wt_row + 4}*{total_ctn_letter_final}{total_row})'
    )
    style_cell(cbm_result, font=bold, border=box, alignment=center)

    # Grand Total gross/net weight (kg) - merged in summary block
    first_sum_row = sum_header_row + 1
    last_sum_row = sum_header_row + len(rows_needed)

    ws.merge_cells(start_row=first_sum_row, start_column=gw_col,
                    end_row=last_sum_row, end_column=gw_col)
    gw_cell = ws.cell(row=first_sum_row, column=gw_col,
                       value=f'={total_grs_letter}{total_row}')
    style_cell(gw_cell, font=bold, fill=total_fill, border=box, alignment=center)

    ws.merge_cells(start_row=first_sum_row, start_column=nw_col,
                    end_row=last_sum_row, end_column=nw_col)
    nw_cell = ws.cell(row=first_sum_row, column=nw_col,
                       value=f'={total_net_letter}{total_row}')
    style_cell(nw_cell, font=bold, fill=total_fill, border=box, alignment=center)

    # ---- Column widths ----
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    for i in range(n_sizes):
        ws.column_dimensions[get_column_letter(5 + i)].width = 10
    for col in range(5 + n_sizes, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions[get_column_letter(gw_col)].width = 24
    ws.column_dimensions[get_column_letter(nw_col)].width = 24

    # ---- Legend ----
    legend_row = wt_row + 6
    ws.cell(row=legend_row, column=1, value="Legend:").font = bold
    ws.cell(row=legend_row, column=1).alignment = center
    legend = ws.cell(row=legend_row + 1, column=1,
                     value="Yellow cells = data NOT available in the PO PDF (fill in from factory cutting/packing records)")
    legend.font = Font(name=FONT, size=9, italic=True, color="666666")
    legend.alignment = left_center


wb = Workbook()
wb.remove(wb.active)

used_names = set()
for line in data["lines"]:
    sublines = [s for s in data["sublines"] if s["style"] == line["style"]]
    if not sublines:
        continue
    name = sheet_name_for(line)
    base_name, i = name, 1
    while name in used_names:
        i += 1
        name = f"{base_name[:28]}_{i}"
    used_names.add(name)

    ws = wb.create_sheet(title=name)
    build_sheet(ws, line, sublines)

out_path = "packing_list_all_PO_lines.xlsx"
wb.save(out_path)
print("Saved:", out_path)
print("Sheets created:", wb.sheetnames)
