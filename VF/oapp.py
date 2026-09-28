"""
PO -> Packing List Streamlit App
=================================
Upload a VF/Vans-style PO PDF, review the extracted data, and download
a packing-list-style Excel workbook (one sheet per PO line) - all in
the browser, no terminal steps needed.

HOW TO RUN:
    pip install -r requirements.txt
    streamlit run app.py

Then open the local URL it prints (usually http://localhost:8501).
"""
import io
import re

import pdfplumber
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

# ============================================================
# PART 1 - PDF EXTRACTION
# ============================================================

LINE_PATTERN = re.compile(
    r"(\d{15})\s+([A-Z0-9]+)\s+([A-Z0-9 \-/]+?)\s+(\d+)\s+PIECE\s+([\d.]+)\s+([\d,]+\.\d{2})"
)

SUBLINE_PATTERN = re.compile(
    r"(\d{12,13})\s+([A-Z0-9]+)\s+[A-Z0-9 \-/]+?\s+(\d+)\s+([\d.]+)\s*\n"
    r"[A-Z0-9 %\-/]+?\s*\n"
    r"UnitOfMeasureCode\s+PC\s+Color\s+([A-Z /]+?)\s*\n"
    r"Size\s+(\S+(?:\s\S+)?)\s+Dimension\s+1\s*\n"
    r"SKU Number\s+(\S+)\s+UPC Number\s+(\d+)\s*\n"
    r"Packing Method\s+CASE\s+Items Per Outer\s+(\d+)\s*\nPack"
)

CRD_PATTERN = re.compile(r"Brand Requested CRD\s+(\d{4}-\d{2}-\d{2})")
COUNTRY_PATTERN = re.compile(r"([A-Z]{3,})\s*\na\s*\na\s*\nBrand Buyer Code")


def extract_data(pdf_file):
    """pdf_file: a file-like object (e.g. Streamlit's UploadedFile)."""
    with pdfplumber.open(pdf_file) as pdf:
        full_text = "\n".join(page.extract_text() for page in pdf.pages)

    lines = []
    for m in LINE_PATTERN.finditer(full_text):
        lines.append({
            "po_line_no": m.group(1),
            "style": m.group(2),
            "description": m.group(3).strip(),
            "order_qty": int(m.group(4)),
            "unit_price": float(m.group(5)),
            "line_cost": float(m.group(6).replace(",", "")),
        })

    sublines = []
    for m in SUBLINE_PATTERN.finditer(full_text):
        sublines.append({
            "subline_no": m.group(1),
            "style": m.group(2),
            "qty": int(m.group(3)),
            "unit_cost": float(m.group(4)),
            "color": m.group(5).strip(),
            "size": m.group(6).strip(),
            "sku": m.group(7),
            "upc": m.group(8),
            "items_per_outer_pack": int(m.group(9)),
        })

    crd_matches = CRD_PATTERN.findall(full_text)
    for i, l in enumerate(lines):
        l["crd"] = crd_matches[i] if i < len(crd_matches) else None

    country_match = COUNTRY_PATTERN.search(full_text)
    destination_country = country_match.group(1) if country_match else None
    for l in lines:
        l["destination_country"] = destination_country

    return {"lines": lines, "sublines": sublines}


# ============================================================
# PART 2 - EXCEL BUILDING
# ============================================================

FONT = "Arial"
bold = Font(name=FONT, bold=True, size=10)
normal = Font(name=FONT, size=10)
title_font = Font(name=FONT, bold=True, size=13)
thin = Side(style="thin")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
yellow = PatternFill("solid", fgColor="FFFF00")
grey = PatternFill("solid", fgColor="D9D9D9")


def sheet_name_for(line):
    raw = f"{line['style']}"
    return re.sub(r'[\[\]:*?/\\]', '-', raw)[:31]


def compute_carton_rows(sublines):
    if not sublines:
        return []
    pack = sublines[0].get("items_per_outer_pack") or 1

    full_rows = []
    remainder_pool = []
    for s in sublines:
        qty = s["qty"]
        size = s["size"]
        full_ctn = qty // pack
        rem = qty % pack
        if full_ctn > 0:
            full_rows.append({
                "sizes": {size: pack},
                "total_ctn": full_ctn,
                "ctn_pcs": pack,
                "color": s["color"],
                "upc": s["upc"],
                "mixed": False,
            })
        if rem > 0:
            remainder_pool.append([size, rem, s["color"], s["upc"]])

    mixed_rows = []
    current = {}
    current_total = 0
    idx = 0
    while idx < len(remainder_pool):
        size, rem, color, upc = remainder_pool[idx]
        space_left = pack - current_total
        take = min(rem, space_left)
        if take > 0:
            current[size] = current.get(size, 0) + take
            current_total += take
            remainder_pool[idx][1] -= take
        if remainder_pool[idx][1] == 0:
            idx += 1
        if current_total == pack:
            mixed_rows.append({
                "sizes": dict(current), "total_ctn": 1, "ctn_pcs": pack,
                "color": None, "upc": None, "mixed": True,
            })
            current = {}
            current_total = 0
    if current:
        mixed_rows.append({
            "sizes": dict(current), "total_ctn": 1, "ctn_pcs": current_total,
            "color": None, "upc": None, "mixed": True,
        })

    return full_rows + mixed_rows


def build_sheet(ws, line, sublines):
    sizes = [s["size"] for s in sublines]
    n_sizes = len(sizes)

    ws.merge_cells("A1:H1")
    ws["A1"] = "CREATIVE COLLECTIONS LTD. U-1-A"
    ws["A1"].font = title_font
    ws["A1"].alignment = center

    ws.merge_cells("A2:H2")
    ws["A2"] = "PACKING LIST DETAILS"
    ws["A2"].font = Font(name=FONT, bold=True, size=10)
    ws["A2"].alignment = center

    # ---- Size-spec table ----
    spec_header_row = 4
    ws.cell(row=spec_header_row, column=1, value="SIZE").font = bold
    ws.cell(row=spec_header_row, column=1).border = box
    for sc_idx, size_name in enumerate(sizes):
        c = ws.cell(row=spec_header_row, column=2 + sc_idx, value=size_name)
        c.font = bold
        c.alignment = center
        c.border = box
        c.fill = grey

    spec_rows = ["N.W.", "N.N.W.", "EMPTY CTN"]
    for i, label in enumerate(spec_rows):
        rr = spec_header_row + 1 + i
        lc = ws.cell(row=rr, column=1, value=label)
        lc.font = bold
        lc.border = box
        for sc_idx, size_name in enumerate(sizes):
            c = ws.cell(row=rr, column=2 + sc_idx, value=1)
            c.alignment = center
            c.border = box

    r = spec_header_row + len(spec_rows) + 2
    ws.cell(row=r, column=1, value="BUYER").font = bold
    ws.cell(row=r, column=2, value="VANS").font = normal
    r += 1
    ws.cell(row=r, column=1, value="LOT NO").font = bold
    ws.cell(row=r, column=2, value=line["style"]).font = normal
    r += 1
    ws.cell(row=r, column=1, value="P.O NO").font = bold
    ws.cell(row=r, column=2, value=line["po_line_no"]).font = normal
    r += 1
    ws.cell(row=r, column=1, value="ORDER QTY").font = bold
    ws.cell(row=r, column=2, value=line["order_qty"]).font = normal
    ws.cell(row=r, column=3, value="PCS").font = normal
    r += 1
    ws.cell(row=r, column=1, value="PACK QTY").font = bold
    ws.cell(row=r, column=2, value=line["order_qty"]).font = normal
    ws.cell(row=r, column=3, value="PCS").font = normal
    r += 1
    excess_short_row = r
    ws.cell(row=r, column=1, value="EXCESS/SHORT QTY").font = bold
    ws.cell(row=r, column=3, value="PCS").font = normal
    r += 1
    percentage_row = r
    ws.cell(row=r, column=1, value="PERCENTAGE").font = bold
    r += 1
    ws.cell(row=r, column=1, value="CRD").font = bold
    ws.cell(row=r, column=2, value=line.get("crd")).font = normal
    r += 1
    ws.cell(row=r, column=1, value="COUNTRY").font = bold
    ws.cell(row=r, column=2, value=line.get("destination_country")).font = normal
    r += 1
    ws.cell(row=r, column=1, value="DESCRIPTION").font = bold
    ws.cell(row=r, column=2, value=line["description"]).font = normal

    header_row = r + 2
    headers = ["CASE LABEL NO.", "CARTON NO.", "COLOR", "UPC Number"] + sizes + \
              ["CTN PCS", "TOTAL CTN", "TOTAL PCS", "Grs.wt.pr ctn", "Net.wt.pr ctn",
               "total.Grs.wt", "total.net.wt"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = bold
        c.alignment = center
        c.border = box
        c.fill = grey
    n_cols = len(headers)

    data_row_start = header_row + 1
    carton_rows = compute_carton_rows(sublines)
    for idx, row_data in enumerate(carton_rows):
        rr = data_row_start + idx
        ws.cell(row=rr, column=1, value=idx + 1).border = box

        color_val = row_data["color"] if row_data["color"] is not None else "MIXED"
        color_cell = ws.cell(row=rr, column=3, value=color_val)
        color_cell.border = box
        if row_data["mixed"]:
            color_cell.fill = yellow
        ws.cell(row=rr, column=4, value=row_data["upc"]).border = box
        for sc_idx, size_name in enumerate(sizes):
            col = 5 + sc_idx
            val = row_data["sizes"].get(size_name)
            cell = ws.cell(row=rr, column=col, value=val)
            cell.border = box
            cell.alignment = center

        ctn_pcs_col = 5 + n_sizes
        first_size_letter = get_column_letter(5)
        last_size_letter = get_column_letter(4 + n_sizes)
        ctn_cell = ws.cell(
            row=rr, column=ctn_pcs_col,
            value=f"=SUM({first_size_letter}{rr}:{last_size_letter}{rr})"
        )
        ctn_cell.border = box
        ctn_cell.alignment = center
        ctn_cell.font = bold

        total_ctn_col = ctn_pcs_col + 1
        total_pcs_col = ctn_pcs_col + 2
        total_ctn_letter = get_column_letter(total_ctn_col)

        tc_cell = ws.cell(row=rr, column=total_ctn_col, value=row_data["total_ctn"])
        tc_cell.border = box
        tc_cell.alignment = center

        # CARTON NO. = cumulative TOTAL CTN
        if idx == 0:
            carton_formula = f"={total_ctn_letter}{rr}"
        else:
            carton_formula = f"={get_column_letter(2)}{rr - 1}+{total_ctn_letter}{rr}"
        cB = ws.cell(row=rr, column=2, value=carton_formula)
        cB.border = box
        cB.alignment = center

        ctn_pcs_letter = get_column_letter(ctn_pcs_col)
        tp_cell = ws.cell(
            row=rr, column=total_pcs_col,
            value=f"={total_ctn_letter}{rr}*{ctn_pcs_letter}{rr}"
        )
        tp_cell.border = box
        tp_cell.alignment = center

        grs_wt_col = total_pcs_col + 1
        nw_row = spec_header_row + 1
        empty_row = spec_header_row + 3
        first_spec_letter = get_column_letter(2)
        last_spec_letter = get_column_letter(1 + n_sizes)
        first_main_size_letter = get_column_letter(5)
        last_main_size_letter = get_column_letter(4 + n_sizes)

        grs_formula = (
            f"=SUMPRODUCT({first_main_size_letter}{rr}:{last_main_size_letter}{rr},"
            f"{first_spec_letter}{nw_row}:{last_spec_letter}{nw_row})"
            f"+{first_spec_letter}{empty_row}"
        )
        grs_cell = ws.cell(row=rr, column=grs_wt_col, value=grs_formula)
        grs_cell.border = box
        grs_cell.alignment = center

        net_wt_col = grs_wt_col + 1
        grs_wt_letter_for_net = get_column_letter(grs_wt_col)
        net_cell = ws.cell(
            row=rr, column=net_wt_col,
            value=f"={grs_wt_letter_for_net}{rr}-{first_spec_letter}{empty_row}"
        )
        net_cell.border = box
        net_cell.alignment = center

        total_grs_col = net_wt_col + 1
        grs_wt_letter = get_column_letter(grs_wt_col)
        tg_cell = ws.cell(
            row=rr, column=total_grs_col,
            value=f"={total_ctn_letter}{rr}*{grs_wt_letter}{rr}"
        )
        tg_cell.border = box
        tg_cell.alignment = center

        total_net_col = total_grs_col + 1
        net_wt_letter = get_column_letter(net_wt_col)
        tn_cell = ws.cell(
            row=rr, column=total_net_col,
            value=f"={total_ctn_letter}{rr}*{net_wt_letter}{rr}"
        )
        tn_cell.border = box
        tn_cell.alignment = center

    total_row = data_row_start + len(carton_rows)
    ws.cell(row=total_row, column=3, value="TOTAL").font = bold
    total_ctn_letter_for_sum = get_column_letter(total_ctn_col)
    for sc_idx, size_name in enumerate(sizes):
        col = 5 + sc_idx
        col_letter = get_column_letter(col)
        cell = ws.cell(
            row=total_row, column=col,
            value=f"=SUMPRODUCT({col_letter}{data_row_start}:{col_letter}{total_row-1},"
                  f"{total_ctn_letter_for_sum}{data_row_start}:{total_ctn_letter_for_sum}{total_row-1})"
        )
        cell.font = bold
        cell.border = box

    total_ctn_col = ctn_pcs_col + 1
    total_pcs_col = ctn_pcs_col + 2
    total_ctn_letter = get_column_letter(total_ctn_col)
    total_pcs_letter = get_column_letter(total_pcs_col)

    ws.cell(row=total_row, column=total_ctn_col,
            value=f'=IF(COUNT({total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1})=0,"",'
                  f'SUM({total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1}))').font = bold

    ws.cell(row=total_row, column=total_pcs_col,
            value=f'=IF(COUNT({total_pcs_letter}{data_row_start}:{total_pcs_letter}{total_row-1})=0,"",'
                  f'SUM({total_pcs_letter}{data_row_start}:{total_pcs_letter}{total_row-1}))').font = bold

    grs_wt_col = total_pcs_col + 1
    net_wt_col = total_pcs_col + 2
    grs_wt_letter = get_column_letter(grs_wt_col)
    net_wt_letter = get_column_letter(net_wt_col)

    # SUMPRODUCT (not SUM): these are PER-CARTON weights, so a row with
    # TOTAL CTN=2 must count double toward the grand total.
    ws.cell(row=total_row, column=grs_wt_col,
            value=f"=SUMPRODUCT({grs_wt_letter}{data_row_start}:{grs_wt_letter}{total_row-1},"
                  f"{total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1})").font = bold
    ws.cell(row=total_row, column=net_wt_col,
            value=f"=SUMPRODUCT({net_wt_letter}{data_row_start}:{net_wt_letter}{total_row-1},"
                  f"{total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1})").font = bold

    total_grs_col = total_pcs_col + 3
    total_grs_letter = get_column_letter(total_grs_col)
    ws.cell(row=total_row, column=total_grs_col,
            value=f'=IF(COUNT({total_grs_letter}{data_row_start}:{total_grs_letter}{total_row-1})=0,"",'
                  f'SUM({total_grs_letter}{data_row_start}:{total_grs_letter}{total_row-1}))').font = bold

    total_net_col = total_grs_col + 1
    total_net_letter = get_column_letter(total_net_col)
    ws.cell(row=total_row, column=total_net_col,
            value=f'=IF(COUNT({total_net_letter}{data_row_start}:{total_net_letter}{total_row-1})=0,"",'
                  f'SUM({total_net_letter}{data_row_start}:{total_net_letter}{total_row-1}))').font = bold

    # ---- Summary block ----
    sum_header_row = total_row + 3
    ws.cell(row=sum_header_row, column=1, value="SIZE").font = bold
    for sc_idx, size_name in enumerate(sizes):
        ws.cell(row=sum_header_row, column=2 + sc_idx, value=size_name).font = bold
    ws.cell(row=sum_header_row, column=2 + n_sizes, value="G TOTAL").font = bold

    gw_col = 3 + n_sizes
    nw_col = 4 + n_sizes
    ws.merge_cells(start_row=sum_header_row, start_column=gw_col,
                    end_row=sum_header_row, end_column=gw_col)
    ws.cell(row=sum_header_row, column=gw_col, value="Grand Total gross weight (kg)").font = bold
    ws.cell(row=sum_header_row, column=gw_col).alignment = center
    ws.merge_cells(start_row=sum_header_row, start_column=nw_col,
                    end_row=sum_header_row, end_column=nw_col)
    ws.cell(row=sum_header_row, column=nw_col, value="Grand Total net weight (kg)").font = bold
    ws.cell(row=sum_header_row, column=nw_col).alignment = center

    rows_needed = ["CUT QTY", "ORDER QTY", "SHIP QTY", "EXS/SHT QTY", "PERCENTAGE"]
    for i, label in enumerate(rows_needed):
        rr = sum_header_row + 1 + i
        ws.cell(row=rr, column=1, value=label).font = bold
        for sc_idx, size_name in enumerate(sizes):
            col = 2 + sc_idx
            col_letter = get_column_letter(col)
            cell = ws.cell(row=rr, column=col)
            if label == "ORDER QTY":
                qty = next(s["qty"] for s in sublines if s["size"] == size_name)
                cell.value = qty
            elif label == "CUT QTY":
                cell.fill = yellow
            elif label == "SHIP QTY":
                # packet-wise total from main table TOTAL row
                main_size_col = 5 + sc_idx
                main_size_letter = get_column_letter(main_size_col)
                cell.value = f"={main_size_letter}{total_row}"
            elif label == "EXS/SHT QTY":
                ship_row = sum_header_row + 1 + rows_needed.index("SHIP QTY")
                order_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
                cell.value = f'=IF({col_letter}{ship_row}="","",{col_letter}{ship_row}-{col_letter}{order_row})'
            elif label == "PERCENTAGE":
                exs_row = sum_header_row + 1 + rows_needed.index("EXS/SHT QTY")
                order_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
                cell.value = f'=IF({col_letter}{exs_row}="","",IFERROR({col_letter}{exs_row}/{col_letter}{order_row},0))'
                cell.number_format = "0.00%"
            cell.border = box
            cell.alignment = center

        gt_col = 2 + n_sizes
        gt_letter = get_column_letter(gt_col)
        first_letter = get_column_letter(2)
        last_letter = get_column_letter(1 + n_sizes)
        if label in ("CUT QTY", "EXS/SHT QTY"):
            ws.cell(row=rr, column=gt_col,
                    value=f'=IF(COUNT({first_letter}{rr}:{last_letter}{rr})=0,"",SUM({first_letter}{rr}:{last_letter}{rr}))').font = bold
        elif label in ("ORDER QTY", "SHIP QTY"):
            ws.cell(row=rr, column=gt_col,
                    value=f"=SUM({first_letter}{rr}:{last_letter}{rr})").font = bold
        else:
            exs_row = sum_header_row + 1 + rows_needed.index("EXS/SHT QTY")
            order_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
            ws.cell(row=rr, column=gt_col,
                    value=f'=IF({gt_letter}{exs_row}="","",IFERROR({gt_letter}{exs_row}/{gt_letter}{order_row},0))').number_format = "0.00%"
        ws.cell(row=rr, column=gt_col).border = box

    exs_gt_row = sum_header_row + 1 + rows_needed.index("EXS/SHT QTY")
    order_gt_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
    gt_col_letter = get_column_letter(2 + n_sizes)

    ws.cell(row=excess_short_row, column=2,
            value=f'={gt_col_letter}{exs_gt_row}').font = normal
    ws.cell(row=percentage_row, column=2,
            value=f'=IF({gt_col_letter}{exs_gt_row}="","-",'
                  f'TEXT(IFERROR({gt_col_letter}{exs_gt_row}/{gt_col_letter}{order_gt_row},0),"0.00%"))'
            ).font = normal

    # ---- Weight / CBM block ----
    wt_row = sum_header_row + len(rows_needed) + 2
    total_grs_letter = get_column_letter(total_grs_col)
    total_net_letter = get_column_letter(total_net_col)
    total_ctn_letter_final = get_column_letter(total_ctn_col)

    # GROSS WEIGHT = Grand Total gross weight
    ws.cell(row=wt_row, column=1, value="GROSS WEIGHT :").font = bold
    c = ws.cell(row=wt_row, column=2, value=f"={total_grs_letter}{total_row}")
    c.font = bold
    c.border = box
    c.alignment = center

    # NET WEIGHT = Grand Total net weight
    ws.cell(row=wt_row + 1, column=1, value="NET WEIGHT :").font = bold
    c = ws.cell(row=wt_row + 1, column=2, value=f"={total_net_letter}{total_row}")
    c.font = bold
    c.border = box
    c.alignment = center

    # CTN MEAS : user input 1 + user input 2
    ws.cell(row=wt_row + 2, column=1, value="CTN MEAS :").font = bold
    ws.cell(row=wt_row + 2, column=2).fill = yellow
    ws.cell(row=wt_row + 2, column=2).border = box
    ws.cell(row=wt_row + 3, column=2).fill = yellow
    ws.cell(row=wt_row + 3, column=2).border = box

    # CBM : (yellow numeric input) × TOTAL CTN
    ws.cell(row=wt_row + 4, column=1, value="CBM :").font = bold
    ws.cell(row=wt_row + 4, column=2).fill = yellow
    ws.cell(row=wt_row + 4, column=2).border = box
    cbm_result = ws.cell(
        row=wt_row + 4, column=3,
        value=f'=IF(B{wt_row + 4}="","",B{wt_row + 4}*{total_ctn_letter_final}{total_row})'
    )
    cbm_result.font = bold
    cbm_result.border = box
    cbm_result.alignment = center

    # Grand Total gross/net weight (kg) - merged in summary
    first_sum_row = sum_header_row + 1
    last_sum_row = sum_header_row + len(rows_needed)

    ws.merge_cells(start_row=first_sum_row, start_column=gw_col,
                    end_row=last_sum_row, end_column=gw_col)
    gw_cell = ws.cell(row=first_sum_row, column=gw_col,
                       value=f'={total_grs_letter}{total_row}')
    gw_cell.alignment = center
    gw_cell.font = bold
    gw_cell.border = box

    ws.merge_cells(start_row=first_sum_row, start_column=nw_col,
                    end_row=last_sum_row, end_column=nw_col)
    nw_cell = ws.cell(row=first_sum_row, column=nw_col,
                       value=f'={total_net_letter}{total_row}')
    nw_cell.alignment = center
    nw_cell.font = bold
    nw_cell.border = box

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    for i in range(n_sizes):
        ws.column_dimensions[get_column_letter(5 + i)].width = 10
    for col in range(5 + n_sizes, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions[get_column_letter(gw_col)].width = 22
    ws.column_dimensions[get_column_letter(nw_col)].width = 22

    legend_row = wt_row + 6
    ws.cell(row=legend_row, column=1, value="Legend:").font = bold
    ws.cell(row=legend_row + 1, column=1,
            value="Yellow cells = data NOT available in the PO PDF (fill in from factory cutting/packing records)"
            ).font = Font(name=FONT, size=9, italic=True)


def build_workbook(data):
    wb = Workbook()
    wb.remove(wb.active)
    used_names = set()
    for line in data["lines"]:
        sublines = [s for s in data["sublines"] if s["style"] == line["style"]]
        if not sublines:
            continue
        name = sheet_name_for(line)
        base_name, i = name, 1
        while name in used_names:
            i += 1
            name = f"{base_name[:28]}_{i}"
        used_names.add(name)
        ws = wb.create_sheet(title=name)
        build_sheet(ws, line, sublines)
    return wb


# ============================================================
# HTML PACKING-LIST PREVIEW
# ============================================================

_PLIST_CSS = """
<style>
.plist-wrap { font-family: Arial, sans-serif; font-size: 13px; color: #000; }
table.plist { border-collapse: collapse; width: 100%; margin-bottom: 6px; }
table.plist td, table.plist th { border: 1px solid #000; padding: 3px 8px; text-align: center; }
.hdr { font-weight: bold; background: #fff; }
.yellow { background: #FFFF66; }
.gray { background: #D9D9D9; font-weight: bold; }
.bold { font-weight: bold; }
.noborder td { border: none; padding: 1px 4px; }
.desc-cell { text-align: left; font-weight: bold; }
.title-row td { border: none; font-weight: bold; padding: 2px 4px; }
</style>
"""


def render_size_spec_table(style_key, full_sizes):
    st.markdown(_PLIST_CSS, unsafe_allow_html=True)
    st.markdown(
        '<div class="plist-wrap">'
        '<table class="plist noborder">'
        '<tr class="title-row"><td style="text-align:center; font-size:16px;">'
        '<b>CREATIVE COLLECTIONS LTD. U-1-A</b></td></tr>'
        '<tr class="title-row"><td style="text-align:center;">PACKING LIST DETAILS.</td></tr>'
        '</table></div>',
        unsafe_allow_html=True,
    )

    header_cols = st.columns([1] + [1] * len(full_sizes))
    header_cols[0].markdown("**SIZE**")
    for col, size in zip(header_cols[1:], full_sizes):
        col.markdown(f"**{size}**")

    nw_by_size, nnw_by_size, empty_by_size = {}, {}, {}

    # Use v2 keys to avoid stale session_state from older app versions
    row_cols = st.columns([1] + [1] * len(full_sizes))
    row_cols[0].markdown("**N.W.**")
    for col, size in zip(row_cols[1:], full_sizes):
        nw_by_size[size] = col.number_input(
            f"N.W. — {size}", min_value=0.0, value=1.0, step=0.01, format="%.2f",
            key=f"nw_v2_{style_key}_{size}",
            label_visibility="collapsed",
        )

    row_cols = st.columns([1] + [1] * len(full_sizes))
    row_cols[0].markdown("**N.N.W.**")
    for col, size in zip(row_cols[1:], full_sizes):
        nnw_by_size[size] = col.number_input(
            f"N.N.W. — {size}", min_value=0.0, value=1.0, step=0.01, format="%.2f",
            key=f"nnw_v2_{style_key}_{size}",
            label_visibility="collapsed",
        )

    row_cols = st.columns([1] + [1] * len(full_sizes))
    row_cols[0].markdown("**EMPTY CTN**")
    for col, size in zip(row_cols[1:], full_sizes):
        empty_by_size[size] = col.number_input(
            f"EMPTY CTN — {size}", min_value=0.0, value=1.0, step=0.01, format="%.2f",
            key=f"empty_v2_{style_key}_{size}",
            label_visibility="collapsed",
        )

    return nw_by_size, nnw_by_size, empty_by_size


def render_packing_list_preview(line, sublines):
    style_key = line["style"]
    full_sizes = [s["size"] for s in sublines]
    carton_rows = compute_carton_rows(sublines)
    n_rows = len(carton_rows)

    nw_by_size, nnw_by_size, empty_by_size = render_size_spec_table(style_key, full_sizes)
    empty_ctn_wt = float(empty_by_size[full_sizes[0]]) if full_sizes else 0.0

    st.markdown(_PLIST_CSS, unsafe_allow_html=True)
    st.caption(
        "Edit **SIZE qty** and **CTN PCS** in each carton row below. "
        "CTN PCS auto-syncs to Σ(sizes) when sizes change; you can still override it. "
        "Weights, TOTAL PCS, SHIP QTY and totals recalculate automatically."
    )

    # ---- Editable carton rows ----
    st.markdown("##### Main carton table (editable SIZE + CTN PCS)")

    # Header labels
    hdr_cols = st.columns([0.6, 0.7, 1.0, 1.0] + [0.7] * len(full_sizes) + [0.7, 0.7, 0.8, 0.8, 0.8, 0.8, 0.8])
    headers = ["#", "Ctn No", "Color", "UPC"] + full_sizes + [
        "CTN PCS", "TOTAL CTN", "TOTAL PCS", "Grs/ctn", "Net/ctn", "tot Grs", "tot Net"
    ]
    for col, h in zip(hdr_cols, headers):
        col.markdown(f"**{h}**")

    edited_sizes = []   # list of dict size -> qty
    edited_ctn_pcs = []
    edited_total_ctn = []
    grs_wts, net_wts = [], []
    row_total_pcs = []
    cum = 0

    for i, r in enumerate(carton_rows):
        color = "MIXED" if r["mixed"] else (r["color"] or "")
        upc = r["upc"] or ""
        total_ctn_val = int(r["total_ctn"])

        cols = st.columns([0.6, 0.7, 1.0, 1.0] + [0.7] * len(full_sizes) + [0.7, 0.7, 0.8, 0.8, 0.8, 0.8, 0.8])

        cols[0].markdown(str(i + 1))
        # Carton No shown after we know total_ctn (cumulative)
        cols[2].markdown(color)
        cols[3].markdown(str(upc) if upc else "")

        # Editable SIZE quantities
        size_vals = {}
        for j, sz in enumerate(full_sizes):
            default_qty = int(r["sizes"].get(sz, 0) or 0)
            size_vals[sz] = cols[4 + j].number_input(
                f"sz_{i}_{sz}",
                min_value=0,
                value=default_qty,
                step=1,
                key=f"sz_v3_{style_key}_{i}_{sz}",
                label_visibility="collapsed",
            )

        size_sum = int(sum(size_vals.values()))

        # CTN PCS: auto-sync to size sum when sizes change; still editable
        pcs_key = f"pcs_v3_{style_key}_{i}"
        sum_key = f"pcs_sum_v3_{style_key}_{i}"
        if sum_key not in st.session_state or st.session_state[sum_key] != size_sum:
            st.session_state[pcs_key] = size_sum
            st.session_state[sum_key] = size_sum

        ctn_pcs_val = cols[4 + len(full_sizes)].number_input(
            f"pcs_{i}",
            min_value=0,
            step=1,
            key=pcs_key,
            label_visibility="collapsed",
        )

        # TOTAL CTN (editable too so packing can be adjusted)
        tctn_key = f"tctn_v3_{style_key}_{i}"
        if tctn_key not in st.session_state:
            st.session_state[tctn_key] = total_ctn_val
        total_ctn_edit = cols[5 + len(full_sizes)].number_input(
            f"tctn_{i}",
            min_value=0,
            step=1,
            key=tctn_key,
            label_visibility="collapsed",
        )

        cum += total_ctn_edit
        cols[1].markdown(str(cum))

        # Weights from edited sizes
        piece_wt = sum(
            float(size_vals.get(sz, 0) or 0) * float(nw_by_size.get(sz, 0.0) or 0.0)
            for sz in full_sizes
        )
        grs = piece_wt + empty_ctn_wt
        net = piece_wt
        tot_pcs = total_ctn_edit * ctn_pcs_val
        tot_grs = total_ctn_edit * grs
        tot_net = total_ctn_edit * net

        cols[6 + len(full_sizes)].markdown(f"**{tot_pcs}**")
        cols[7 + len(full_sizes)].markdown(f"{grs:.2f}")
        cols[8 + len(full_sizes)].markdown(f"{net:.2f}")
        cols[9 + len(full_sizes)].markdown(f"{tot_grs:.2f}")
        cols[10 + len(full_sizes)].markdown(f"{tot_net:.2f}")

        edited_sizes.append(size_vals)
        edited_ctn_pcs.append(ctn_pcs_val)
        edited_total_ctn.append(total_ctn_edit)
        grs_wts.append(grs)
        net_wts.append(net)
        row_total_pcs.append(tot_pcs)

    # ---- Aggregates from edited values ----
    total_ctn = sum(edited_total_ctn)
    total_pcs = sum(row_total_pcs)
    total_grs = sum(tc * g for tc, g in zip(edited_total_ctn, grs_wts))
    total_net = sum(tc * n for tc, n in zip(edited_total_ctn, net_wts))

    ship_qty_by_size = {sz: 0 for sz in full_sizes}
    for size_vals, tc in zip(edited_sizes, edited_total_ctn):
        for sz in full_sizes:
            ship_qty_by_size[sz] += int(size_vals.get(sz, 0) or 0) * tc

    size_packet_totals = {
        sz: sum(int(edited_sizes[i].get(sz, 0) or 0) * edited_total_ctn[i] for i in range(n_rows))
        for sz in full_sizes
    }

    with st.expander(f"Edit cut qty / CTN MEAS / CBM for {style_key}", expanded=False):
        cut_qty = {}
        cut_cols = st.columns(len(full_sizes)) if full_sizes else []
        for col, size in zip(cut_cols, full_sizes):
            with col:
                cut_qty[size] = st.number_input(
                    f"Cut qty — {size}", value=0, key=f"cut_v3_{style_key}_{size}"
                )
        ctn_meas_1 = st.text_input("Ctn meas — line 1", "", key=f"meas1_v3_{style_key}")
        ctn_meas_2 = st.text_input("Ctn meas — line 2", "", key=f"meas2_v3_{style_key}")
        cbm_per_ctn = st.number_input(
            "CBM per carton (m³)", value=0.0, format="%.8f", key=f"cbm_v3_{style_key}"
        )

    order_qty_by_size = {s["size"]: s["qty"] for s in sublines}
    order_total = sum(order_qty_by_size.get(sz, 0) for sz in full_sizes)
    ship_total = sum(ship_qty_by_size.get(sz, 0) for sz in full_sizes)
    cut_total = sum(cut_qty.get(sz, 0) for sz in full_sizes)
    exs_by_size = {sz: ship_qty_by_size.get(sz, 0) - order_qty_by_size.get(sz, 0) for sz in full_sizes}
    exs_total = ship_total - order_total
    pct_by_size = {
        sz: (exs_by_size[sz] / order_qty_by_size[sz]) if order_qty_by_size.get(sz) else None
        for sz in full_sizes
    }
    pct_total = (exs_total / order_total) if order_total else None
    cbm_total = cbm_per_ctn * total_ctn

    # TOTAL row strip
    st.markdown(
        f"**TOTAL** &nbsp;|&nbsp; "
        + " &nbsp; ".join(f"{sz}: **{size_packet_totals[sz]}**" for sz in full_sizes)
        + f" &nbsp;|&nbsp; CTN: **{total_ctn}** &nbsp;|&nbsp; PCS: **{total_pcs}** "
        + f"&nbsp;|&nbsp; Grs: **{total_grs:.2f}** &nbsp;|&nbsp; Net: **{total_net:.2f}**"
    )

    st.markdown('<div class="plist-wrap">', unsafe_allow_html=True)

    st.markdown(f"""
<table class="plist noborder">
<tr class="title-row">
  <td style="width:16%"><b>BUYER</b></td><td style="width:14%">VANS</td>
  <td style="width:14%"></td>
  <td style="width:16%"><b>ORDER QTY</b></td><td style="width:8%">{line['order_qty']}</td><td style="width:6%">PCS</td>
  <td style="width:8%"></td>
  <td style="width:6%"><u><b>CRD</b></u></td><td>{line.get('crd','')}</td>
</tr>
<tr class="title-row">
  <td><b>LOT NO</b></td><td>{line['style']}</td><td></td>
  <td><b>PACK QTY</b></td><td>{line['order_qty']}</td><td>PCS</td><td></td><td></td><td></td>
</tr>
<tr class="title-row">
  <td><b>P.O NO</b></td><td>{line['po_line_no']}</td><td></td>
  <td><b>EXCESS/SHORT QTY</b></td><td>{exs_total}</td><td>PCS</td><td></td><td></td><td></td>
</tr>
<tr class="title-row">
  <td colspan="2" class="desc-cell">DESCRIPTION:<br>{line['description']}</td><td></td>
  <td><b>PERCENTAGE</b></td>
  <td>{'-' if pct_total is None else f'{pct_total*100:.2f}'}</td><td>%</td><td></td>
  <td colspan="2">{line.get('destination_country','')}</td>
</tr>
</table>
""", unsafe_allow_html=True)

    def fmt_pct(v):
        return "" if v is None else f"{v*100:.2f}%"

    n_data_rows = 5
    size_th = "".join(f"<th>{sz}</th>" for sz in full_sizes)

    summary_body = ""
    for i, (label, vals, gt, cls) in enumerate([
        ("CUT QTY", [cut_qty.get(sz, "") for sz in full_sizes], cut_total, "yellow"),
        ("ORDER QTY", [order_qty_by_size.get(sz, "") for sz in full_sizes], order_total, "gray"),
        ("SHIP QTY", [ship_qty_by_size.get(sz, "") for sz in full_sizes], ship_total, ""),
        ("EXS/SHT QTY", [exs_by_size.get(sz, "") for sz in full_sizes], exs_total, ""),
        ("PERCENTAGE", [fmt_pct(pct_by_size.get(sz)) for sz in full_sizes], fmt_pct(pct_total), ""),
    ]):
        cells = "".join(f'<td class="{cls}">{v}</td>' for v in vals)
        if i == 0:
            weight_cells = (
                f'<td class="bold" rowspan="{n_data_rows}" style="vertical-align:middle;">'
                f'{total_grs:.2f}</td>'
                f'<td class="bold" rowspan="{n_data_rows}" style="vertical-align:middle;">'
                f'{total_net:.2f}</td>'
            )
        else:
            weight_cells = ""
        summary_body += (
            f'<tr><td class="{cls} bold">{label}</td>{cells}'
            f'<td class="bold">{gt}</td>{weight_cells}</tr>'
        )

    st.markdown(f"""
<table class="plist">
<tr class="hdr">
  <th>SIZE</th>{size_th}
  <th>G Total</th>
  <th>Grand Total gross weight kg</th>
  <th>Grand Total Net weight kg</th>
</tr>
{summary_body}
</table>
""", unsafe_allow_html=True)

    st.markdown(f"""
<table class="plist noborder">
<tr class="title-row"><td style="width:15%"><b>GROSS WEIGHT :</b></td><td>{total_grs:.2f}</td></tr>
<tr class="title-row"><td><b>NET WEIGHT :</b></td><td>{total_net:.2f}</td></tr>
<tr class="title-row"><td><b>CTN MEAS :</b></td><td>{ctn_meas_1}<br>{ctn_meas_2}</td></tr>
<tr class="title-row"><td><b>CBM :</b></td><td>{cbm_total:.8f}</td></tr>
</table>
<br><br>
<table class="plist noborder">
<tr class="title-row"><td style="width:30%; border-top:1px solid #000; text-align:center;">INCHARGE</td>
<td style="width:30%"></td>
<td style="width:30%; border-top:1px solid #000; text-align:center;">FM</td></tr>
</table>
""", unsafe_allow_html=True)

    st.markdown('</div>', unsafe_allow_html=True)


# ============================================================
# PART 3 - STREAMLIT UI
# ============================================================

st.set_page_config(page_title="PO -> Packing List", layout="wide")
st.title("PO to Packing List Converter")
st.write(
    "Upload a VF/Vans-style Purchase Order PDF. This extracts the line items "
    "and sizes, then builds a packing-list-style Excel workbook (one sheet per "
    "PO line) with live formulas for cartons, weights, and totals."
)

uploaded_file = st.file_uploader("Upload PO PDF", type=["pdf"])

if uploaded_file is not None:
    with st.spinner("Reading PDF..."):
        try:
            data = extract_data(uploaded_file)
        except Exception as e:
            st.error(f"Failed to read the PDF: {e}")
            st.stop()

    if not data["lines"]:
        st.warning(
            "No PO lines were found. This PDF's layout may differ from the "
            "expected format - the extraction patterns may need adjusting."
        )
        st.stop()

    st.success(f"Found {len(data['lines'])} PO line(s) and {len(data['sublines'])} size/color subline(s).")

    st.subheader("PO Lines")
    st.caption("Click a row to filter the sublines below and download just that line's sheet.")
    po_line_rows = [
        {
            "PO Line No.": l["po_line_no"],
            "Style": l["style"],
            "Description": l["description"],
            "Order Qty": l["order_qty"],
            "Unit Price": l["unit_price"],
            "Line Cost": l["line_cost"],
            "CRD": l.get("crd"),
            "Country": l.get("destination_country"),
        }
        for l in data["lines"]
    ]
    selection = st.dataframe(
        po_line_rows,
        use_container_width=True,
        on_select="rerun",
        selection_mode="single-row",
        key="po_lines_table",
    )

    selected_idx = None
    rows_selected = selection.selection.rows if selection is not None else []
    if rows_selected:
        selected_idx = rows_selected[0]
        selected_line = data["lines"][selected_idx]
        st.info(f"Selected: **{selected_line['style']}** (PO Line {selected_line['po_line_no']})")

    st.subheader("Sublines (size / color / UPC)")
    if selected_idx is not None:
        filtered_sublines = [s for s in data["sublines"] if s["style"] == selected_line["style"]]
        st.dataframe(filtered_sublines, use_container_width=True)
    else:
        st.dataframe(data["sublines"], use_container_width=True)

    col1, col2 = st.columns(2)

    with col1:
        if st.button("Build full workbook (all PO lines)", type="primary"):
            with st.spinner("Building Excel workbook..."):
                wb = build_workbook(data)
                buf = io.BytesIO()
                wb.save(buf)
                buf.seek(0)

            st.success(f"Workbook built with {len(wb.sheetnames)} sheet(s): {', '.join(wb.sheetnames)}")
            st.download_button(
                label="Download packing_list_all_PO_lines.xlsx",
                data=buf,
                file_name="packing_list_all_PO_lines.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

    with col2:
        if selected_idx is not None:
            if st.button(f"Build sheet for {selected_line['style']} only"):
                single_line_data = {
                    "lines": [selected_line],
                    "sublines": [s for s in data["sublines"] if s["style"] == selected_line["style"]],
                }
                with st.spinner("Building Excel sheet..."):
                    wb = build_workbook(single_line_data)
                    buf = io.BytesIO()
                    wb.save(buf)
                    buf.seek(0)

                st.success(f"Sheet built for {selected_line['style']}")
                st.download_button(
                    label=f"Download {selected_line['style']}.xlsx",
                    data=buf,
                    file_name=f"packing_list_{selected_line['style']}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                st.session_state["preview_style"] = selected_line["style"]
        else:
            st.caption("Select a PO line above to download just that sheet.")

    if selected_idx is not None and st.session_state.get("preview_style") == selected_line["style"]:
        st.subheader(f"Packing List Preview — {selected_line['style']}")
        preview_sublines = [s for s in data["sublines"] if s["style"] == selected_line["style"]]
        render_packing_list_preview(selected_line, preview_sublines)

    st.info(
        "Yellow cells in the downloaded file still need manual entry: "
        "CUT QTY, CTN MEAS, and CBM factor - these come from your factory's "
        "cutting/packing records, not the PO. "
        "CARTON NO., SHIP QTY, GROSS/NET WEIGHT are auto-calculated."
    )
else:
    st.info("Upload a PDF to get started.")
