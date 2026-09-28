# """
# PO -> Packing List Streamlit App
# =================================
# Upload a VF/Vans-style PO PDF, review the extracted data, and download
# a packing-list-style Excel workbook (one sheet per PO line) - all in
# the browser, no terminal steps needed.

# HOW TO RUN:
#     pip install -r requirements.txt
#     (requirements.txt should include: streamlit, pdfplumber, openpyxl, reportlab)
#     streamlit run app.py
# Then open the local URL it prints (usually http://localhost:8501).
# """
# import io
# import json
# import re
# import uuid
# import datetime

# import pdfplumber
# import streamlit as st
# from openpyxl import Workbook
# from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
# from openpyxl.utils import get_column_letter
# from reportlab.lib import colors as rl_colors
# from reportlab.lib.pagesizes import A4, landscape
# from reportlab.lib.units import mm
# from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
# from reportlab.lib.styles import getSampleStyleSheet
# from html import escape

# # ============================================================
# # PART 1 - PDF EXTRACTION
# # ============================================================

# LINE_PATTERN = re.compile(
#     r"""
#     (?P<po_line_no>\d{15})
#     \s+
#     (?P<style>[A-Z0-9]+)
#     \s+
#     (?P<description>[A-Z0-9 \-/]+?)
#     \s+
#     (?P<order_qty>[\d,]+)
#     \s+
#     (?:PIECE|EACH)
#     \s+
#     (?P<unit_price>[\d.]+)
#     \s+
#     (?P<line_cost>[\d,]+\.\d{2})
#     """,
#     re.IGNORECASE | re.VERBOSE,
# )


# SUBLINE_START_PATTERN = re.compile(
#     r"(?P<subline_no>\d{12,13})\s+(?P<style>VN[A-Z0-9]+)",
#     re.IGNORECASE,
# )


# def extract_sublines(full_text):
#     """Extract VF/Vans sublines from the PDF text."""
#     text = full_text.replace("\r\n", "\n").replace("\r", "\n")
#     text = text.replace("\xa0", " ")

#     matches = list(SUBLINE_START_PATTERN.finditer(text))
#     sublines = []

#     for i, start_match in enumerate(matches):
#         block_start = start_match.start()
#         block_end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
#         block = text[block_start:block_end]

#         if not re.search(r"UnitOfMeasureCode\s+(?:PC|EA)\b", block, re.I):
#             continue

#         qty_cost = re.search(
#             r"(?P<style>VN[A-Z0-9]+)\s+.*?"
#             r"(?P<qty>[\d,]+)\s+(?P<unit_cost>\d+(?:\.\d+)?)\s+.*?"
#             r"UnitOfMeasureCode",
#             block,
#             re.IGNORECASE | re.DOTALL,
#         )

#         color = re.search(
#             r"UnitOfMeasureCode\s+(?:PC|EA)\s+Color\s+(?P<color>.*?)\s+Size",
#             block,
#             re.IGNORECASE | re.DOTALL,
#         )

#         size = re.search(
#             r"Size\s+(?P<size>.*?)\s+Dimension\s+1",
#             block,
#             re.IGNORECASE | re.DOTALL,
#         )

#         sku = re.search(
#             r"SKU\s+Number\s+(?P<sku>\S+)",
#             block,
#             re.IGNORECASE,
#         )

#         upc = re.search(
#             r"UPC\s+Number\s+(?P<upc>\d{12,13})",
#             block,
#             re.IGNORECASE,
#         )

#         pack = re.search(
#             r"Items\s+Per\s+Outer\s+(?:Pack\s+)?(?P<pack>[\d,]+)",
#             block,
#             re.IGNORECASE | re.DOTALL,
#         )

#         if not all((qty_cost, color, size, sku, upc, pack)):
#             continue

#         sublines.append({
#             "subline_no": start_match.group("subline_no"),
#             "style": start_match.group("style"),
#             "qty": int(qty_cost.group("qty").replace(",", "")),
#             "unit_cost": float(qty_cost.group("unit_cost").replace(",", "")),
#             "color": re.sub(r"\s+", " ", color.group("color")).strip(),
#             "size": re.sub(r"\s+", " ", size.group("size")).strip(),
#             "sku": sku.group("sku"),
#             "upc": upc.group("upc"),
#             "items_per_outer_pack": int(pack.group("pack").replace(",", "")),
#         })

#     return sublines


# CRD_PATTERN = re.compile(r"Brand Requested CRD\s+(\d{4}-\d{2}-\d{2})")
# COUNTRY_PATTERN = re.compile(r"([A-Z]{3,})\s*\na\s*\na\s*\nBrand Buyer Code")


# def extract_data(pdf_file):
#     """pdf_file: a file-like object (e.g. Streamlit's UploadedFile)."""
#     with pdfplumber.open(pdf_file) as pdf:
#         full_text = "\n".join((page.extract_text() or "") for page in pdf.pages)

#     lines = []
#     for m in LINE_PATTERN.finditer(full_text):
#         lines.append({
#             "po_line_no": m.group("po_line_no"),
#             "style": m.group("style"),
#             "description": re.sub(r"\s+", " ", m.group("description")).strip(),
#             "order_qty": int(m.group("order_qty").replace(",", "")),
#             "unit_price": float(m.group("unit_price").replace(",", "")),
#             "line_cost": float(m.group("line_cost").replace(",", "")),
#         })

#     sublines = extract_sublines(full_text)

#     crd_matches = CRD_PATTERN.findall(full_text)
#     for i, l in enumerate(lines):
#         l["crd"] = crd_matches[i] if i < len(crd_matches) else None

#     country_match = COUNTRY_PATTERN.search(full_text)
#     destination_country = country_match.group(1) if country_match else None
#     for l in lines:
#         l["destination_country"] = destination_country

#     print(f"[PDF EXTRACTION] PO lines found: {len(lines)}")
#     print(f"[PDF EXTRACTION] Sublines found: {len(sublines)}")
#     return {"lines": lines, "sublines": sublines}


# # ============================================================
# # PART 2 - EXCEL BUILDING
# # ============================================================

# FONT = "Arial"
# bold = Font(name=FONT, bold=True, size=10)
# normal = Font(name=FONT, size=10)
# title_font = Font(name=FONT, bold=True, size=13)
# thin = Side(style="thin")
# box = Border(left=thin, right=thin, top=thin, bottom=thin)
# center = Alignment(horizontal="center", vertical="center", wrap_text=True)
# yellow = PatternFill("solid", fgColor="FFFF00")
# grey = PatternFill("solid", fgColor="D9D9D9")


# def sheet_name_for(line):
#     raw = f"{line['style']}"
#     return re.sub(r'[\[\]:*?/\\]', '-', raw)[:31]


# def compute_carton_rows(sublines, pack_override=None):
#     if not sublines:
#         return []
#     pack = pack_override if pack_override is not None else (
#         sublines[0].get("items_per_outer_pack") or 1
#     )

#     full_rows = []
#     remainder_pool = []
#     for s in sublines:
#         qty = s["qty"]
#         size = s["size"]
#         full_ctn = qty // pack
#         rem = qty % pack
#         if full_ctn > 0:
#             full_rows.append({
#                 "sizes": {size: pack},
#                 "total_ctn": full_ctn,
#                 "ctn_pcs": pack,
#                 "color": s["color"],
#                 "upc": s["upc"],
#                 "mixed": False,
#             })
#         if rem > 0:
#             remainder_pool.append([size, rem, s["color"], s["upc"]])

#     mixed_rows = []
#     current = {}
#     current_total = 0
#     idx = 0
#     while idx < len(remainder_pool):
#         size, rem, color, upc = remainder_pool[idx]
#         space_left = pack - current_total
#         take = min(rem, space_left)
#         if take > 0:
#             current[size] = current.get(size, 0) + take
#             current_total += take
#             remainder_pool[idx][1] -= take
#         if remainder_pool[idx][1] == 0:
#             idx += 1
#         if current_total == pack:
#             mixed_rows.append({
#                 "sizes": dict(current), "total_ctn": 1, "ctn_pcs": pack,
#                 "color": None, "upc": None, "mixed": True,
#             })
#             current = {}
#             current_total = 0
#     if current:
#         mixed_rows.append({
#             "sizes": dict(current), "total_ctn": 1, "ctn_pcs": current_total,
#             "color": None, "upc": None, "mixed": True,
#         })

#     return full_rows + mixed_rows


# def build_sheet(ws, line, sublines):
#     sizes = [s["size"] for s in sublines]
#     n_sizes = len(sizes)

#     ws.merge_cells("A1:H1")
#     ws["A1"] = "CREATIVE COLLECTIONS LTD. U-1-A"
#     ws["A1"].font = title_font
#     ws["A1"].alignment = center

#     ws.merge_cells("A2:H2")
#     ws["A2"] = "PACKING LIST DETAILS"
#     ws["A2"].font = Font(name=FONT, bold=True, size=10)
#     ws["A2"].alignment = center

#     # ---- Size-spec table ----
#     spec_header_row = 4
#     ws.cell(row=spec_header_row, column=1, value="SIZE").font = bold
#     ws.cell(row=spec_header_row, column=1).border = box
#     for sc_idx, size_name in enumerate(sizes):
#         c = ws.cell(row=spec_header_row, column=2 + sc_idx, value=size_name)
#         c.font = bold
#         c.alignment = center
#         c.border = box
#         c.fill = grey

#     spec_rows = ["N.W.", "N.N.W.", "EMPTY CTN"]
#     for i, label in enumerate(spec_rows):
#         rr = spec_header_row + 1 + i
#         lc = ws.cell(row=rr, column=1, value=label)
#         lc.font = bold
#         lc.border = box
#         for sc_idx, size_name in enumerate(sizes):
#             c = ws.cell(row=rr, column=2 + sc_idx, value=1)
#             c.alignment = center
#             c.border = box

#     r = spec_header_row + len(spec_rows) + 2
#     ws.cell(row=r, column=1, value="BUYER").font = bold
#     ws.cell(row=r, column=2, value="VANS").font = normal
#     r += 1
#     ws.cell(row=r, column=1, value="LOT NO").font = bold
#     ws.cell(row=r, column=2, value=line["style"]).font = normal
#     r += 1
#     ws.cell(row=r, column=1, value="P.O NO").font = bold
#     ws.cell(row=r, column=2, value=line["po_line_no"]).font = normal
#     r += 1
#     ws.cell(row=r, column=1, value="ORDER QTY").font = bold
#     ws.cell(row=r, column=2, value=line["order_qty"]).font = normal
#     ws.cell(row=r, column=3, value="PCS").font = normal
#     r += 1
#     ws.cell(row=r, column=1, value="PACK QTY").font = bold
#     ws.cell(row=r, column=2, value=line["order_qty"]).font = normal
#     ws.cell(row=r, column=3, value="PCS").font = normal
#     r += 1
#     excess_short_row = r
#     ws.cell(row=r, column=1, value="EXCESS/SHORT QTY").font = bold
#     ws.cell(row=r, column=3, value="PCS").font = normal
#     r += 1
#     percentage_row = r
#     ws.cell(row=r, column=1, value="PERCENTAGE").font = bold
#     r += 1
#     ws.cell(row=r, column=1, value="CRD").font = bold
#     ws.cell(row=r, column=2, value=line.get("crd")).font = normal
#     r += 1
#     ws.cell(row=r, column=1, value="COUNTRY").font = bold
#     ws.cell(row=r, column=2, value=line.get("destination_country")).font = normal
#     r += 1
#     ws.cell(row=r, column=1, value="DESCRIPTION").font = bold
#     ws.cell(row=r, column=2, value=line["description"]).font = normal

#     header_row = r + 2
#     headers = ["CASE LABEL NO.", "CARTON NO.", "COLOR", "UPC Number"] + sizes + \
#               ["CTN PCS", "TOTAL CTN", "TOTAL PCS", "Grs.wt.pr ctn", "Net.wt.pr ctn",
#                "total.Grs.wt", "total.net.wt"]
#     for i, h in enumerate(headers, start=1):
#         c = ws.cell(row=header_row, column=i, value=h)
#         c.font = bold
#         c.alignment = center
#         c.border = box
#         c.fill = grey
#     n_cols = len(headers)

#     data_row_start = header_row + 1
#     carton_rows = compute_carton_rows(sublines)
#     for idx, row_data in enumerate(carton_rows):
#         rr = data_row_start + idx
#         ws.cell(row=rr, column=1, value=idx + 1).border = box

#         color_val = row_data["color"] if row_data["color"] is not None else "MIXED"
#         color_cell = ws.cell(row=rr, column=3, value=color_val)
#         color_cell.border = box
#         if row_data["mixed"]:
#             color_cell.fill = yellow
#         ws.cell(row=rr, column=4, value=row_data["upc"]).border = box
#         for sc_idx, size_name in enumerate(sizes):
#             col = 5 + sc_idx
#             val = row_data["sizes"].get(size_name)
#             cell = ws.cell(row=rr, column=col, value=val)
#             cell.border = box
#             cell.alignment = center

#         ctn_pcs_col = 5 + n_sizes
#         first_size_letter = get_column_letter(5)
#         last_size_letter = get_column_letter(4 + n_sizes)
#         ctn_cell = ws.cell(
#             row=rr, column=ctn_pcs_col,
#             value=f"=SUM({first_size_letter}{rr}:{last_size_letter}{rr})"
#         )
#         ctn_cell.border = box
#         ctn_cell.alignment = center
#         ctn_cell.font = bold

#         total_ctn_col = ctn_pcs_col + 1
#         total_pcs_col = ctn_pcs_col + 2
#         total_ctn_letter = get_column_letter(total_ctn_col)

#         tc_cell = ws.cell(row=rr, column=total_ctn_col, value=row_data["total_ctn"])
#         tc_cell.border = box
#         tc_cell.alignment = center

#         if idx == 0:
#             carton_formula = f"={total_ctn_letter}{rr}"
#         else:
#             carton_formula = f"={get_column_letter(2)}{rr - 1}+{total_ctn_letter}{rr}"
#         cB = ws.cell(row=rr, column=2, value=carton_formula)
#         cB.border = box
#         cB.alignment = center

#         ctn_pcs_letter = get_column_letter(ctn_pcs_col)
#         tp_cell = ws.cell(
#             row=rr, column=total_pcs_col,
#             value=f"={total_ctn_letter}{rr}*{ctn_pcs_letter}{rr}"
#         )
#         tp_cell.border = box
#         tp_cell.alignment = center

#         grs_wt_col = total_pcs_col + 1
#         nw_row = spec_header_row + 1
#         empty_row = spec_header_row + 3
#         first_spec_letter = get_column_letter(2)
#         last_spec_letter = get_column_letter(1 + n_sizes)
#         first_main_size_letter = get_column_letter(5)
#         last_main_size_letter = get_column_letter(4 + n_sizes)

#         grs_formula = (
#             f"=SUMPRODUCT({first_main_size_letter}{rr}:{last_main_size_letter}{rr},"
#             f"{first_spec_letter}{nw_row}:{last_spec_letter}{nw_row})"
#             f"+{first_spec_letter}{empty_row}"
#         )
#         grs_cell = ws.cell(row=rr, column=grs_wt_col, value=grs_formula)
#         grs_cell.border = box
#         grs_cell.alignment = center

#         net_wt_col = grs_wt_col + 1
#         grs_wt_letter_for_net = get_column_letter(grs_wt_col)
#         net_cell = ws.cell(
#             row=rr, column=net_wt_col,
#             value=f"={grs_wt_letter_for_net}{rr}-{first_spec_letter}{empty_row}"
#         )
#         net_cell.border = box
#         net_cell.alignment = center

#         total_grs_col = net_wt_col + 1
#         grs_wt_letter = get_column_letter(grs_wt_col)
#         tg_cell = ws.cell(
#             row=rr, column=total_grs_col,
#             value=f"={total_ctn_letter}{rr}*{grs_wt_letter}{rr}"
#         )
#         tg_cell.border = box
#         tg_cell.alignment = center

#         total_net_col = total_grs_col + 1
#         net_wt_letter = get_column_letter(net_wt_col)
#         tn_cell = ws.cell(
#             row=rr, column=total_net_col,
#             value=f"={total_ctn_letter}{rr}*{net_wt_letter}{rr}"
#         )
#         tn_cell.border = box
#         tn_cell.alignment = center

#     total_row = data_row_start + len(carton_rows)
#     ws.cell(row=total_row, column=3, value="TOTAL").font = bold
#     total_ctn_letter_for_sum = get_column_letter(total_ctn_col)
#     for sc_idx, size_name in enumerate(sizes):
#         col = 5 + sc_idx
#         col_letter = get_column_letter(col)
#         cell = ws.cell(
#             row=total_row, column=col,
#             value=f"=SUMPRODUCT({col_letter}{data_row_start}:{col_letter}{total_row-1},"
#                   f"{total_ctn_letter_for_sum}{data_row_start}:{total_ctn_letter_for_sum}{total_row-1})"
#         )
#         cell.font = bold
#         cell.border = box

#     total_ctn_col = ctn_pcs_col + 1
#     total_pcs_col = ctn_pcs_col + 2
#     total_ctn_letter = get_column_letter(total_ctn_col)
#     total_pcs_letter = get_column_letter(total_pcs_col)

#     ws.cell(row=total_row, column=total_ctn_col,
#             value=f'=IF(COUNT({total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1})=0,"",'
#                   f'SUM({total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1}))').font = bold

#     ws.cell(row=total_row, column=total_pcs_col,
#             value=f'=IF(COUNT({total_pcs_letter}{data_row_start}:{total_pcs_letter}{total_row-1})=0,"",'
#                   f'SUM({total_pcs_letter}{data_row_start}:{total_pcs_letter}{total_row-1}))').font = bold

#     grs_wt_col = total_pcs_col + 1
#     net_wt_col = total_pcs_col + 2
#     grs_wt_letter = get_column_letter(grs_wt_col)
#     net_wt_letter = get_column_letter(net_wt_col)

#     ws.cell(row=total_row, column=grs_wt_col,
#             value=f"=SUMPRODUCT({grs_wt_letter}{data_row_start}:{grs_wt_letter}{total_row-1},"
#                   f"{total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1})").font = bold
#     ws.cell(row=total_row, column=net_wt_col,
#             value=f"=SUMPRODUCT({net_wt_letter}{data_row_start}:{net_wt_letter}{total_row-1},"
#                   f"{total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1})").font = bold

#     total_grs_col = total_pcs_col + 3
#     total_grs_letter = get_column_letter(total_grs_col)
#     ws.cell(row=total_row, column=total_grs_col,
#             value=f'=IF(COUNT({total_grs_letter}{data_row_start}:{total_grs_letter}{total_row-1})=0,"",'
#                   f'SUM({total_grs_letter}{data_row_start}:{total_grs_letter}{total_row-1}))').font = bold

#     total_net_col = total_grs_col + 1
#     total_net_letter = get_column_letter(total_net_col)
#     ws.cell(row=total_row, column=total_net_col,
#             value=f'=IF(COUNT({total_net_letter}{data_row_start}:{total_net_letter}{total_row-1})=0,"",'
#                   f'SUM({total_net_letter}{data_row_start}:{total_net_letter}{total_row-1}))').font = bold

#     # ---- Summary block ----
#     sum_header_row = total_row + 3
#     ws.cell(row=sum_header_row, column=1, value="SIZE").font = bold
#     for sc_idx, size_name in enumerate(sizes):
#         ws.cell(row=sum_header_row, column=2 + sc_idx, value=size_name).font = bold
#     ws.cell(row=sum_header_row, column=2 + n_sizes, value="G TOTAL").font = bold

#     gw_col = 3 + n_sizes
#     nw_col = 4 + n_sizes
#     ws.merge_cells(start_row=sum_header_row, start_column=gw_col,
#                     end_row=sum_header_row, end_column=gw_col)
#     ws.cell(row=sum_header_row, column=gw_col, value="Grand Total gross weight (kg)").font = bold
#     ws.cell(row=sum_header_row, column=gw_col).alignment = center
#     ws.merge_cells(start_row=sum_header_row, start_column=nw_col,
#                     end_row=sum_header_row, end_column=nw_col)
#     ws.cell(row=sum_header_row, column=nw_col, value="Grand Total net weight (kg)").font = bold
#     ws.cell(row=sum_header_row, column=nw_col).alignment = center

#     rows_needed = ["CUT QTY", "ORDER QTY", "SHIP QTY", "EXS/SHT QTY", "PERCENTAGE"]
#     for i, label in enumerate(rows_needed):
#         rr = sum_header_row + 1 + i
#         ws.cell(row=rr, column=1, value=label).font = bold
#         for sc_idx, size_name in enumerate(sizes):
#             col = 2 + sc_idx
#             col_letter = get_column_letter(col)
#             cell = ws.cell(row=rr, column=col)
#             if label == "ORDER QTY":
#                 qty = next(s["qty"] for s in sublines if s["size"] == size_name)
#                 cell.value = qty
#             elif label == "CUT QTY":
#                 cell.fill = yellow
#             elif label == "SHIP QTY":
#                 main_size_col = 5 + sc_idx
#                 main_size_letter = get_column_letter(main_size_col)
#                 cell.value = f"={main_size_letter}{total_row}"
#             elif label == "EXS/SHT QTY":
#                 ship_row = sum_header_row + 1 + rows_needed.index("SHIP QTY")
#                 order_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
#                 cell.value = f'=IF({col_letter}{ship_row}="","",{col_letter}{ship_row}-{col_letter}{order_row})'
#             elif label == "PERCENTAGE":
#                 exs_row = sum_header_row + 1 + rows_needed.index("EXS/SHT QTY")
#                 order_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
#                 cell.value = f'=IF({col_letter}{exs_row}="","",IFERROR({col_letter}{exs_row}/{col_letter}{order_row},0))'
#                 cell.number_format = "0.00%"
#             cell.border = box
#             cell.alignment = center

#         gt_col = 2 + n_sizes
#         gt_letter = get_column_letter(gt_col)
#         first_letter = get_column_letter(2)
#         last_letter = get_column_letter(1 + n_sizes)
#         if label in ("CUT QTY", "EXS/SHT QTY"):
#             ws.cell(row=rr, column=gt_col,
#                     value=f'=IF(COUNT({first_letter}{rr}:{last_letter}{rr})=0,"",SUM({first_letter}{rr}:{last_letter}{rr}))').font = bold
#         elif label in ("ORDER QTY", "SHIP QTY"):
#             ws.cell(row=rr, column=gt_col,
#                     value=f"=SUM({first_letter}{rr}:{last_letter}{rr})").font = bold
#         else:
#             exs_row = sum_header_row + 1 + rows_needed.index("EXS/SHT QTY")
#             order_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
#             ws.cell(row=rr, column=gt_col,
#                     value=f'=IF({gt_letter}{exs_row}="","",IFERROR({gt_letter}{exs_row}/{gt_letter}{order_row},0))').number_format = "0.00%"
#         ws.cell(row=rr, column=gt_col).border = box

#     exs_gt_row = sum_header_row + 1 + rows_needed.index("EXS/SHT QTY")
#     order_gt_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
#     gt_col_letter = get_column_letter(2 + n_sizes)

#     ws.cell(row=excess_short_row, column=2,
#             value=f'={gt_col_letter}{exs_gt_row}').font = normal
#     ws.cell(row=percentage_row, column=2,
#             value=f'=IF({gt_col_letter}{exs_gt_row}="","-",'
#                   f'TEXT(IFERROR({gt_col_letter}{exs_gt_row}/{gt_col_letter}{order_gt_row},0),"0.00%"))'
#             ).font = normal

#     # ---- Weight / CBM block ----
#     wt_row = sum_header_row + len(rows_needed) + 2
#     total_grs_letter = get_column_letter(total_grs_col)
#     total_net_letter = get_column_letter(total_net_col)
#     total_ctn_letter_final = get_column_letter(total_ctn_col)

#     ws.cell(row=wt_row, column=1, value="GROSS WEIGHT :").font = bold
#     c = ws.cell(row=wt_row, column=2, value=f"={total_grs_letter}{total_row}")
#     c.font = bold
#     c.border = box
#     c.alignment = center

#     ws.cell(row=wt_row + 1, column=1, value="NET WEIGHT :").font = bold
#     c = ws.cell(row=wt_row + 1, column=2, value=f"={total_net_letter}{total_row}")
#     c.font = bold
#     c.border = box
#     c.alignment = center

#     ws.cell(row=wt_row + 2, column=1, value="CTN MEAS :").font = bold
#     ws.cell(row=wt_row + 2, column=2).fill = yellow
#     ws.cell(row=wt_row + 2, column=2).border = box
#     ws.cell(row=wt_row + 3, column=2).fill = yellow
#     ws.cell(row=wt_row + 3, column=2).border = box

#     ws.cell(row=wt_row + 4, column=1, value="CBM :").font = bold
#     ws.cell(row=wt_row + 4, column=2).fill = yellow
#     ws.cell(row=wt_row + 4, column=2).border = box
#     cbm_result = ws.cell(
#         row=wt_row + 4, column=3,
#         value=f'=IF(B{wt_row + 4}="","",B{wt_row + 4}*{total_ctn_letter_final}{total_row})'
#     )
#     cbm_result.font = bold
#     cbm_result.border = box
#     cbm_result.alignment = center

#     first_sum_row = sum_header_row + 1
#     last_sum_row = sum_header_row + len(rows_needed)

#     ws.merge_cells(start_row=first_sum_row, start_column=gw_col,
#                     end_row=last_sum_row, end_column=gw_col)
#     gw_cell = ws.cell(row=first_sum_row, column=gw_col,
#                        value=f'={total_grs_letter}{total_row}')
#     gw_cell.alignment = center
#     gw_cell.font = bold
#     gw_cell.border = box

#     ws.merge_cells(start_row=first_sum_row, start_column=nw_col,
#                     end_row=last_sum_row, end_column=nw_col)
#     nw_cell = ws.cell(row=first_sum_row, column=nw_col,
#                        value=f'={total_net_letter}{total_row}')
#     nw_cell.alignment = center
#     nw_cell.font = bold
#     nw_cell.border = box

#     ws.column_dimensions["A"].width = 16
#     ws.column_dimensions["B"].width = 14
#     ws.column_dimensions["C"].width = 16
#     ws.column_dimensions["D"].width = 16
#     for i in range(n_sizes):
#         ws.column_dimensions[get_column_letter(5 + i)].width = 10
#     for col in range(5 + n_sizes, n_cols + 1):
#         ws.column_dimensions[get_column_letter(col)].width = 13
#     ws.column_dimensions[get_column_letter(gw_col)].width = 22
#     ws.column_dimensions[get_column_letter(nw_col)].width = 22

#     legend_row = wt_row + 6
#     ws.cell(row=legend_row, column=1, value="Legend:").font = bold
#     ws.cell(row=legend_row + 1, column=1,
#             value="Yellow cells = data NOT available in the PO PDF (fill in from factory cutting/packing records)"
#             ).font = Font(name=FONT, size=9, italic=True)


# def build_workbook(data):
#     wb = Workbook()
#     wb.remove(wb.active)
#     used_names = set()
#     for line in data["lines"]:
#         sublines = [s for s in data["sublines"] if s["style"] == line["style"]]
#         if not sublines:
#             continue
#         name = sheet_name_for(line)
#         base_name, i = name, 1
#         while name in used_names:
#             i += 1
#             name = f"{base_name[:28]}_{i}"
#         used_names.add(name)
#         ws = wb.create_sheet(title=name)
#         build_sheet(ws, line, sublines)
#     return wb


# # ============================================================
# # GRID CONSTANTS (Color & UPC columns are wider)
# # ============================================================

# ACTION_WIDTH = 42
# CELL_WIDTH = 72
# COLOR_WIDTH = 120          # ← wider
# UPC_WIDTH = 130            # ← wider
# HEADER_HEIGHT = 36
# CELL_HEIGHT = 34
# TOTAL_HEIGHT = 36
# GRID_HEIGHT = 520


# # ============================================================
# # HTML PACKING-LIST PREVIEW
# # ============================================================

# def render_size_spec_table(style_key, full_sizes):
#     st.markdown(
#         """
#         <style>
#         .plist-wrap { font-family: Arial, sans-serif; font-size: 13px; color: #000; }
#         table.plist { border-collapse: collapse; width: 100%; margin-bottom: 6px; }
#         table.plist td, table.plist th { border: 1px solid #000; padding: 3px 8px; text-align: center; }
#         .hdr { font-weight: bold; background: #fff; }
#         .yellow { background: #FFFF66; }
#         .gray { background: #D9D9D9; font-weight: bold; }
#         .bold { font-weight: bold; }
#         .noborder td { border: none; padding: 1px 4px; }
#         .desc-cell { text-align: left; font-weight: bold; }
#         .title-row td { border: none; font-weight: bold; padding: 2px 4px; }
#         .ctn-total { background:#D9D9D9; color:#000; font-weight:bold; text-align:center; padding:6px 2px; border-top:2px solid #000; border-bottom:2px solid #000; }
#         </style>
#         """,
#         unsafe_allow_html=True,
#     )
#     st.markdown(
#         '<div class="plist-wrap">'
#         '<table class="plist noborder">'
#         '<tr class="title-row"><td style="text-align:center; font-size:16px;">'
#         '<b>CREATIVE COLLECTIONS LTD. U-1-A</b></td></tr>'
#         '<tr class="title-row"><td style="text-align:center;">PACKING LIST DETAILS.</td></tr>'
#         '</table></div>',
#         unsafe_allow_html=True,
#     )

#     header_cols = st.columns([1] + [1] * len(full_sizes))
#     header_cols[0].markdown("**SIZE**")
#     for col, size in zip(header_cols[1:], full_sizes):
#         col.markdown(f"**{size}**")

#     nw_by_size, nnw_by_size, empty_by_size = {}, {}, {}

#     row_cols = st.columns([1] + [1] * len(full_sizes))
#     row_cols[0].markdown("**N.W.**")
#     for col, size in zip(row_cols[1:], full_sizes):
#         nw_by_size[size] = col.number_input(
#             f"N.W. — {size}", min_value=0.0, value=1.0, step=0.01, format="%.2f",
#             key=f"nw_v2_{style_key}_{size}",
#             label_visibility="collapsed",
#         )

#     row_cols = st.columns([1] + [1] * len(full_sizes))
#     row_cols[0].markdown("**N.N.W.**")
#     for col, size in zip(row_cols[1:], full_sizes):
#         nnw_by_size[size] = col.number_input(
#             f"N.N.W. — {size}", min_value=0.0, value=1.0, step=0.01, format="%.2f",
#             key=f"nnw_v2_{style_key}_{size}",
#             label_visibility="collapsed",
#         )

#     row_cols = st.columns([1] + [1] * len(full_sizes))
#     row_cols[0].markdown("**EMPTY CTN**")
#     for col, size in zip(row_cols[1:], full_sizes):
#         empty_by_size[size] = col.number_input(
#             f"EMPTY CTN — {size}", min_value=0.0, value=1.0, step=0.01, format="%.2f",
#             key=f"empty_v2_{style_key}_{size}",
#             label_visibility="collapsed",
#         )

#     return nw_by_size, nnw_by_size, empty_by_size


# def build_edited_workbook(export_state):
#     line = export_state["line"]
#     full_sizes = export_state["full_sizes"]
#     rows = export_state["rows"]
#     n_sizes = len(full_sizes)

#     wb = Workbook()
#     ws = wb.active
#     ws.title = sheet_name_for(line)

#     ws.merge_cells("A1:H1")
#     ws["A1"] = "CREATIVE COLLECTIONS LTD. U-1-A"
#     ws["A1"].font = title_font
#     ws["A1"].alignment = center

#     ws.merge_cells("A2:H2")
#     ws["A2"] = "PACKING LIST DETAILS (edited)"
#     ws["A2"].font = Font(name=FONT, bold=True, size=10)
#     ws["A2"].alignment = center

#     r = 4
#     for label, value in [
#         ("BUYER", "VANS"),
#         ("LOT NO", line["style"]),
#         ("P.O NO", line["po_line_no"]),
#         ("ORDER QTY", line["order_qty"]),
#         ("EXCESS/SHORT QTY", export_state["exs_total"]),
#         ("CRD", line.get("crd")),
#         ("COUNTRY", line.get("destination_country")),
#         ("DESCRIPTION", line["description"]),
#     ]:
#         ws.cell(row=r, column=1, value=label).font = bold
#         ws.cell(row=r, column=2, value=value).font = normal
#         r += 1

#     header_row = r + 1
#     headers = ["Ctn No", "Color", "UPC"] + full_sizes + [
#         "CTN PCS", "TOTAL CTN", "TOTAL PCS", "Grs.wt.pr ctn", "Net.wt.pr ctn",
#         "total.Grs.wt", "total.net.wt"
#     ]
#     for i, h in enumerate(headers, start=1):
#         c = ws.cell(row=header_row, column=i, value=h)
#         c.font = bold
#         c.alignment = center
#         c.border = box
#         c.fill = grey

#     data_row_start = header_row + 1
#     for idx, row_data in enumerate(rows):
#         rr = data_row_start + idx
#         ws.cell(row=rr, column=1, value=row_data["ctn_no"]).border = box
#         ws.cell(row=rr, column=2, value=row_data["color"]).border = box
#         ws.cell(row=rr, column=3, value=row_data["upc"]).border = box
#         for sc_idx, size_name in enumerate(full_sizes):
#             cell = ws.cell(row=rr, column=4 + sc_idx, value=row_data["sizes"].get(size_name))
#             cell.border = box
#             cell.alignment = center
#         base = 4 + n_sizes
#         values = [
#             row_data["ctn_pcs"], row_data["total_ctn"], row_data["total_pcs"],
#             round(row_data["grs_per_ctn"], 2), round(row_data["net_per_ctn"], 2),
#             round(row_data["total_grs"], 2), round(row_data["total_net"], 2),
#         ]
#         for off, val in enumerate(values):
#             cell = ws.cell(row=rr, column=base + off, value=val)
#             cell.border = box
#             cell.alignment = center

#     total_row = data_row_start + len(rows)
#     ws.cell(row=total_row, column=2, value="TOTAL").font = bold
#     for sc_idx, size_name in enumerate(full_sizes):
#         cell = ws.cell(row=total_row, column=4 + sc_idx,
#                         value=export_state["size_packet_totals"].get(size_name, 0))
#         cell.font = bold
#         cell.border = box
#     base = 4 + n_sizes
#     totals_row_values = [
#         "", export_state["total_ctn"], export_state["total_pcs"], "", "",
#         round(export_state["total_grs"], 2), round(export_state["total_net"], 2),
#     ]
#     for off, val in enumerate(totals_row_values):
#         cell = ws.cell(row=total_row, column=base + off, value=val)
#         cell.font = bold
#         cell.border = box

#     sum_header_row = total_row + 3
#     ws.cell(row=sum_header_row, column=1, value="SIZE").font = bold
#     for sc_idx, size_name in enumerate(full_sizes):
#         ws.cell(row=sum_header_row, column=2 + sc_idx, value=size_name).font = bold
#     ws.cell(row=sum_header_row, column=2 + n_sizes, value="G TOTAL").font = bold

#     rows_needed = [
#         ("CUT QTY", export_state["cut_qty"], export_state["cut_total"], yellow),
#         ("ORDER QTY", export_state["order_qty_by_size"], export_state["order_total"], grey),
#         ("SHIP QTY", export_state["ship_qty_by_size"], export_state["ship_total"], None),
#         ("EXS/SHT QTY", export_state["exs_by_size"], export_state["exs_total"], None),
#     ]
#     for ridx, (label, by_size, gt, fill) in enumerate(rows_needed):
#         r_row = sum_header_row + 1 + ridx
#         ws.cell(row=r_row, column=1, value=label).font = bold
#         for sc_idx, size_name in enumerate(full_sizes):
#             cell = ws.cell(row=r_row, column=2 + sc_idx, value=by_size.get(size_name, ""))
#             cell.border = box
#             cell.alignment = center
#             if fill is not None:
#                 cell.fill = fill
#         cell = ws.cell(row=r_row, column=2 + n_sizes, value=gt)
#         cell.font = bold
#         cell.border = box

#     pct_row = sum_header_row + 1 + len(rows_needed)
#     ws.cell(row=pct_row, column=1, value="PERCENTAGE").font = bold
#     for sc_idx, size_name in enumerate(full_sizes):
#         pct = export_state["pct_by_size"].get(size_name)
#         cell = ws.cell(row=pct_row, column=2 + sc_idx,
#                         value="" if pct is None else pct)
#         if pct is not None:
#             cell.number_format = "0.00%"
#         cell.border = box
#         cell.alignment = center
#     pct_total = export_state["pct_total"]
#     cell = ws.cell(row=pct_row, column=2 + n_sizes,
#                     value="" if pct_total is None else pct_total)
#     if pct_total is not None:
#         cell.number_format = "0.00%"
#     cell.font = bold
#     cell.border = box

#     wt_row = pct_row + 3
#     ws.cell(row=wt_row, column=1, value="GROSS WEIGHT :").font = bold
#     ws.cell(row=wt_row, column=2, value=round(export_state["total_grs"], 2)).font = bold
#     ws.cell(row=wt_row + 1, column=1, value="NET WEIGHT :").font = bold
#     ws.cell(row=wt_row + 1, column=2, value=round(export_state["total_net"], 2)).font = bold
#     ws.cell(row=wt_row + 2, column=1, value="CTN MEAS :").font = bold
#     ws.cell(row=wt_row + 2, column=2,
#             value=f'{export_state["ctn_meas_1"]} {export_state["ctn_meas_2"]}'.strip())
#     ws.cell(row=wt_row + 3, column=1, value="CBM :").font = bold
#     ws.cell(row=wt_row + 3, column=2, value=round(export_state["cbm_total"], 8))

#     sign_row = wt_row + 6
#     ws.cell(row=sign_row, column=1, value="INCHARGE").font = bold
#     ws.cell(row=sign_row, column=3, value="FM").font = bold

#     ws.column_dimensions["A"].width = 16
#     ws.column_dimensions["B"].width = 16
#     ws.column_dimensions["C"].width = 16
#     for i in range(n_sizes):
#         ws.column_dimensions[get_column_letter(4 + i)].width = 10

#     return wb


# def build_edited_pdf(export_state):
#     line = export_state["line"]
#     full_sizes = export_state["full_sizes"]
#     rows = export_state["rows"]

#     buf = io.BytesIO()
#     doc = SimpleDocTemplate(
#         buf, pagesize=landscape(A4),
#         leftMargin=5 * mm, rightMargin=5 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
#     )
#     styles = getSampleStyleSheet()
#     elements = []

#     elements.append(Paragraph("CREATIVE COLLECTIONS LTD. U-1-A", styles["Title"]))
#     elements.append(Paragraph("PACKING LIST DETAILS (edited)", styles["Heading3"]))
#     elements.append(Spacer(1, 6))

#     info_data = [
#         ["BUYER", "VANS", "LOT NO", line["style"], "P.O NO", line["po_line_no"]],
#         ["ORDER QTY", line["order_qty"], "EXCESS/SHORT QTY", export_state["exs_total"],
#          "CRD", line.get("crd") or ""],
#         ["COUNTRY", line.get("destination_country") or "", "DESCRIPTION",
#          line["description"], "", ""],
#     ]
#     info_table = Table(info_data, colWidths=[65, 90, 90, 140, 65, 90])
#     info_table.setStyle(TableStyle([
#         ("FONTSIZE", (0, 0), (-1, -1), 8),
#         ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
#         ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.black),
#         ("BACKGROUND", (0, 0), (0, -1), rl_colors.whitesmoke),
#         ("BACKGROUND", (2, 0), (2, -1), rl_colors.whitesmoke),
#         ("BACKGROUND", (4, 0), (4, -1), rl_colors.whitesmoke),
#     ]))
#     elements.append(info_table)
#     elements.append(Spacer(1, 8))

#     headers = ["Ctn No", "Color", "UPC"] + full_sizes + [
#         "CTN PCS", "TOTAL CTN", "TOTAL PCS", "Grs/ctn", "Net/ctn", "tot Grs", "tot Net"
#     ]
#     table_data = [headers]
#     for rdata in rows:
#         table_data.append([
#             rdata["ctn_no"], rdata["color"], rdata["upc"],
#             *[rdata["sizes"].get(sz, "") for sz in full_sizes],
#             rdata["ctn_pcs"], rdata["total_ctn"], rdata["total_pcs"],
#             f'{rdata["grs_per_ctn"]:.2f}', f'{rdata["net_per_ctn"]:.2f}',
#             f'{rdata["total_grs"]:.2f}', f'{rdata["total_net"]:.2f}',
#         ])
#     total_row = (
#         ["", "TOTAL", ""]
#         + [export_state["size_packet_totals"].get(sz, 0) for sz in full_sizes]
#         + ["", export_state["total_ctn"], export_state["total_pcs"], "", "",
#            f'{export_state["total_grs"]:.2f}', f'{export_state["total_net"]:.2f}']
#     )
#     table_data.append(total_row)

#     main_table = Table(table_data, repeatRows=1)
#     main_table.setStyle(TableStyle([
#         ("FONTSIZE", (0, 0), (-1, -1), 6),
#         ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
#         ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
#         ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.black),
#         ("BACKGROUND", (0, 0), (-1, 0), rl_colors.lightgrey),
#         ("BACKGROUND", (0, -1), (-1, -1), rl_colors.lightgrey),
#         ("ALIGN", (0, 0), (-1, -1), "CENTER"),
#     ]))
#     elements.append(main_table)
#     elements.append(Spacer(1, 10))

#     summary_headers = ["SIZE"] + full_sizes + ["G TOTAL"]
#     summary_data = [summary_headers]

#     def fmt_pct(v):
#         return "" if v is None else f"{v * 100:.2f}%"

#     summary_data.append(
#         ["CUT QTY"] + [export_state["cut_qty"].get(sz, "") for sz in full_sizes]
#         + [export_state["cut_total"]]
#     )
#     summary_data.append(
#         ["ORDER QTY"] + [export_state["order_qty_by_size"].get(sz, "") for sz in full_sizes]
#         + [export_state["order_total"]]
#     )
#     summary_data.append(
#         ["SHIP QTY"] + [export_state["ship_qty_by_size"].get(sz, "") for sz in full_sizes]
#         + [export_state["ship_total"]]
#     )
#     summary_data.append(
#         ["EXS/SHT QTY"] + [export_state["exs_by_size"].get(sz, "") for sz in full_sizes]
#         + [export_state["exs_total"]]
#     )
#     summary_data.append(
#         ["PERCENTAGE"] + [fmt_pct(export_state["pct_by_size"].get(sz)) for sz in full_sizes]
#         + [fmt_pct(export_state["pct_total"])]
#     )

#     summary_table = Table(summary_data, repeatRows=1)
#     summary_table.setStyle(TableStyle([
#         ("FONTSIZE", (0, 0), (-1, -1), 7),
#         ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
#         ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.black),
#         ("BACKGROUND", (0, 0), (-1, 0), rl_colors.lightgrey),
#         ("ALIGN", (0, 0), (-1, -1), "CENTER"),
#     ]))
#     elements.append(summary_table)
#     elements.append(Spacer(1, 10))

#     wt_data = [
#         ["GROSS WEIGHT :", f'{export_state["total_grs"]:.2f}'],
#         ["NET WEIGHT :", f'{export_state["total_net"]:.2f}'],
#         ["CTN MEAS :", f'{export_state["ctn_meas_1"]} {export_state["ctn_meas_2"]}'.strip()],
#         ["CBM :", f'{export_state["cbm_total"]:.8f}'],
#     ]
#     wt_table = Table(wt_data, colWidths=[100, 150])
#     wt_table.setStyle(TableStyle([
#         ("FONTSIZE", (0, 0), (-1, -1), 8),
#         ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
#     ]))
#     elements.append(wt_table)
#     elements.append(Spacer(1, 20))

#     sign_table = Table([["INCHARGE", "", "FM"]], colWidths=[120, 120, 120])
#     sign_table.setStyle(TableStyle([
#         ("LINEABOVE", (0, 0), (0, 0), 0.75, rl_colors.black),
#         ("LINEABOVE", (2, 0), (2, 0), 0.75, rl_colors.black),
#         ("ALIGN", (0, 0), (-1, -1), "CENTER"),
#         ("FONTSIZE", (0, 0), (-1, -1), 9),
#     ]))
#     elements.append(sign_table)

#     doc.build(elements)
#     buf.seek(0)
#     return buf


# # ============================================================
# # SAMPLE-STYLE WORKBOOK (centered)
# # ============================================================

# def build_sample_style_workbook(export_state):
#     """
#     Produces the exact sample-style packing list xlsx, fully center-aligned:
#       Title box, size-spec table, info block, main carton table
#       (with merged COLOR runs + yellow MIXED), TOTAL row, summary block
#       with merged Grand Total columns, footer, signature lines.
#     """
#     line = export_state["line"]
#     full_sizes = export_state["full_sizes"]
#     rows = export_state["rows"]
#     n_sizes = len(full_sizes)

#     nw_by_size    = export_state.get("nw_by_size",    {sz: 0.0 for sz in full_sizes})
#     nnw_by_size   = export_state.get("nnw_by_size",   {sz: 0.0 for sz in full_sizes})
#     empty_by_size = export_state.get("empty_by_size", {sz: 0.0 for sz in full_sizes})

#     wb = Workbook()
#     ws = wb.active
#     ws.title = sheet_name_for(line)

#     # ---- Column indices ----
#     C_CASE_A  = 1
#     C_CASE_B  = 2
#     C_COLOR   = 3
#     C_UPC     = 4
#     C_SZ0     = 5
#     C_SZ_LAST = C_SZ0 + n_sizes - 1
#     C_CTN_PCS = C_SZ0 + n_sizes
#     C_TOT_CTN = C_CTN_PCS + 1
#     C_TOT_PCS = C_CTN_PCS + 2
#     C_GRS_CTN = C_CTN_PCS + 3
#     C_NET_CTN = C_CTN_PCS + 4
#     C_TOT_GRS = C_CTN_PCS + 5
#     C_TOT_NET = C_CTN_PCS + 6
#     LAST_COL  = C_TOT_NET

#     BORDER   = Border(left=Side(style="thin"), right=Side(style="thin"),
#                       top=Side(style="thin"), bottom=Side(style="thin"))
#     CENTER   = Alignment(horizontal="center", vertical="center", wrap_text=True)
#     HDR_FILL = PatternFill("solid", fgColor="D9D9D9")

#     def _put(r, c, v, bold_=False, border_=True,
#              fill_=None, numfmt=None, font_size=10):
#         cell = ws.cell(row=r, column=c, value=v)
#         cell.font = Font(name=FONT, bold=bold_, size=font_size)
#         cell.alignment = CENTER
#         if border_:
#             cell.border = BORDER
#         if fill_ is not None:
#             cell.fill = fill_
#         if numfmt:
#             cell.number_format = numfmt
#         return cell

#     # ============ TITLE ============
#     ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=LAST_COL)
#     _put(1, 1, "CREATIVE COLLECTIONS LTD. U-1-A",
#          bold_=True, border_=False, font_size=14)

#     ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=LAST_COL)
#     _put(2, 1, "PACKING LIST DETAILS.",
#          bold_=True, border_=False, font_size=11)

#     # ============ SIZE SPEC TABLE ============
#     spec_top = 4
#     _put(spec_top, 1, "SIZE", bold_=True, fill_=HDR_FILL)
#     for i, sz in enumerate(full_sizes):
#         _put(spec_top, 2 + i, sz, bold_=True, fill_=HDR_FILL)

#     for offset, (label, values) in enumerate([
#         ("N.W.",      nw_by_size),
#         ("N.N.W.",    nnw_by_size),
#         ("EMPTY CTN", empty_by_size),
#     ]):
#         rr = spec_top + 1 + offset
#         _put(rr, 1, label, bold_=True)
#         for i, sz in enumerate(full_sizes):
#             _put(rr, 2 + i, values.get(sz, 0), numfmt="0.00")

#     # ============ INFO BLOCK ============
#     info_top = spec_top + 5

#     _put(info_top, 1, "BUYER", bold_=True, border_=False)
#     _put(info_top, 2, "VANS", border_=False)
#     _put(info_top, 4, "ORDER QTY", bold_=True, border_=False)
#     _put(info_top, 5, line["order_qty"], bold_=True, border_=False)
#     _put(info_top, 6, "PCS", border_=False)
#     _put(info_top, 8, "CRD", bold_=True, border_=False)
#     _put(info_top, 9, line.get("crd", "") or "", border_=False)

#     _put(info_top + 1, 1, "LOT NO", bold_=True, border_=False)
#     _put(info_top + 1, 2, line["style"], border_=False)
#     _put(info_top + 1, 4, "PACK QTY", bold_=True, border_=False)
#     _put(info_top + 1, 5, line["order_qty"], border_=False)
#     _put(info_top + 1, 6, "PCS", border_=False)

#     _put(info_top + 2, 1, "P.O NO", bold_=True, border_=False)
#     _put(info_top + 2, 2, line["po_line_no"], border_=False)
#     _put(info_top + 2, 4, "EXCESS/SHORT QTY", bold_=True, border_=False)
#     _put(info_top + 2, 5, export_state.get("exs_total", 0), border_=False)
#     _put(info_top + 2, 6, "PCS", border_=False)

#     ws.merge_cells(start_row=info_top + 3, start_column=1, end_row=info_top + 3, end_column=3)
#     _put(info_top + 3, 1, f"DESCRIPTION:  {line['description']}",
#          bold_=True, border_=False)
#     _put(info_top + 3, 4, "PERCENTAGE", bold_=True, border_=False)
#     pct_total = export_state.get("pct_total")
#     _put(info_top + 3, 5,
#          "-" if pct_total is None else f"{pct_total*100:.2f}",
#          border_=False)
#     _put(info_top + 3, 6, "%", border_=False)
#     _put(info_top + 3, 8, line.get("destination_country") or "",
#          bold_=True, border_=False)

#     # ============ MAIN TABLE ============
#     table_top = info_top + 5

#     ws.merge_cells(start_row=table_top, start_column=C_CASE_A,
#                    end_row=table_top, end_column=C_CASE_B)
#     _put(table_top, C_CASE_A, "CASE LABEL NO.", bold_=True, fill_=HDR_FILL)
#     _put(table_top, C_CASE_B, None, fill_=HDR_FILL)

#     ws.merge_cells(start_row=table_top, start_column=C_COLOR,
#                    end_row=table_top + 1, end_column=C_COLOR)
#     _put(table_top, C_COLOR, "COLOR", bold_=True, fill_=HDR_FILL)
#     _put(table_top + 1, C_COLOR, None, fill_=HDR_FILL)

#     ws.merge_cells(start_row=table_top, start_column=C_UPC,
#                    end_row=table_top + 1, end_column=C_UPC)
#     _put(table_top, C_UPC, "UPC Number", bold_=True, fill_=HDR_FILL)
#     _put(table_top + 1, C_UPC, None, fill_=HDR_FILL)

#     ws.merge_cells(start_row=table_top, start_column=C_SZ0,
#                    end_row=table_top, end_column=C_SZ_LAST)
#     _put(table_top, C_SZ0, "SIZE", bold_=True, fill_=HDR_FILL)
#     for c in range(C_SZ0 + 1, C_SZ_LAST + 1):
#         _put(table_top, c, None, fill_=HDR_FILL)

#     for c, txt in [
#         (C_CTN_PCS, "CTN PCS"),
#         (C_TOT_CTN, "TOTAL CTN"),
#         (C_TOT_PCS, "TOTAL PCS"),
#         (C_GRS_CTN, "Grs.wt.pr ctn"),
#         (C_NET_CTN, "Net.wt.pr ctn"),
#         (C_TOT_GRS, "total.Grs.wt"),
#         (C_TOT_NET, "total.net.wt"),
#     ]:
#         ws.merge_cells(start_row=table_top, start_column=c,
#                        end_row=table_top + 1, end_column=c)
#         _put(table_top, c, txt, bold_=True, fill_=HDR_FILL)
#         _put(table_top + 1, c, None, fill_=HDR_FILL)

#     for i, sz in enumerate(full_sizes):
#         _put(table_top + 1, C_SZ0 + i, sz, bold_=True, fill_=HDR_FILL)

#     # ---- Data rows ----
#     data_top = table_top + 2
#     cum_ctn = 0
#     row_ranges = []
#     for row in rows:
#         start = cum_ctn + 1
#         cum_ctn += int(row["total_ctn"])
#         end = cum_ctn
#         row_ranges.append((start, end))

#     for idx, row in enumerate(rows):
#         rr = data_top + idx
#         start_ctn, end_ctn = row_ranges[idx]

#         _put(rr, C_CASE_A, start_ctn)
#         _put(rr, C_CASE_B, end_ctn)

#         if row.get("mixed"):
#             _put(rr, C_UPC, None)
#         else:
#             _put(rr, C_UPC, row.get("upc"))

#         for s_idx, sz in enumerate(full_sizes):
#             _put(rr, C_SZ0 + s_idx, row["sizes"].get(sz))

#         _put(rr, C_CTN_PCS, row["ctn_pcs"], bold_=True)
#         _put(rr, C_TOT_CTN, row["total_ctn"])
#         _put(rr, C_TOT_PCS, row["total_pcs"])
#         _put(rr, C_GRS_CTN, round(row["grs_per_ctn"], 2), numfmt="0.00")
#         _put(rr, C_NET_CTN, round(row["net_per_ctn"], 2), numfmt="0.00")
#         _put(rr, C_TOT_GRS, round(row["total_grs"], 2), numfmt="0.00")
#         _put(rr, C_TOT_NET, round(row["total_net"], 2), numfmt="0.00")

#     # ---- Merge COLOR cells in same-color runs ----
#     i = 0
#     while i < len(rows):
#         mixed_i = bool(rows[i].get("mixed"))
#         color_i = "MIXED" if mixed_i else (rows[i].get("color") or "MIXED")
#         j = i
#         while j + 1 < len(rows):
#             mixed_j = bool(rows[j + 1].get("mixed"))
#             color_j = "MIXED" if mixed_j else (rows[j + 1].get("color") or "MIXED")
#             if mixed_i != mixed_j or color_i != color_j:
#                 break
#             j += 1
#         top_r = data_top + i
#         bot_r = data_top + j
#         if bot_r > top_r:
#             ws.merge_cells(start_row=top_r, start_column=C_COLOR,
#                            end_row=bot_r, end_column=C_COLOR)
#         cell = ws.cell(row=top_r, column=C_COLOR)
#         cell.value = color_i
#         cell.font = Font(name=FONT, size=10)
#         cell.alignment = CENTER
#         cell.border = BORDER
#         if mixed_i:
#             cell.fill = yellow
#         for rrr in range(top_r, bot_r + 1):
#             c2 = ws.cell(row=rrr, column=C_COLOR)
#             c2.border = BORDER
#             if mixed_i:
#                 c2.fill = yellow
#         i = j + 1

#     # ---- TOTAL row ----
#     total_row = data_top + len(rows)
#     ws.merge_cells(start_row=total_row, start_column=C_CASE_A,
#                    end_row=total_row, end_column=C_CASE_B)
#     _put(total_row, C_CASE_A, "TOTAL", bold_=True)
#     _put(total_row, C_CASE_B, None, bold_=True)

#     _put(total_row, C_COLOR, None)
#     _put(total_row, C_UPC, None)

#     spt = export_state.get("size_packet_totals", {})
#     for s_idx, sz in enumerate(full_sizes):
#         _put(total_row, C_SZ0 + s_idx, spt.get(sz, 0), bold_=True)

#     _put(total_row, C_CTN_PCS, None)
#     _put(total_row, C_TOT_CTN, export_state.get("total_ctn", 0), bold_=True)
#     _put(total_row, C_TOT_PCS, export_state.get("total_pcs", 0), bold_=True)

#     sum_grs = sum(r["grs_per_ctn"] for r in rows) if rows else 0
#     sum_net = sum(r["net_per_ctn"] for r in rows) if rows else 0
#     _put(total_row, C_GRS_CTN, round(sum_grs, 2), bold_=True, numfmt="0.00")
#     _put(total_row, C_NET_CTN, round(sum_net, 2), bold_=True, numfmt="0.00")
#     _put(total_row, C_TOT_GRS, round(export_state.get("total_grs", 0), 2),
#          bold_=True, numfmt="0.00")
#     _put(total_row, C_TOT_NET, round(export_state.get("total_net", 0), 2),
#          bold_=True, numfmt="0.00")

#     # ============ SUMMARY BLOCK ============
#     sum_top = total_row + 3

#     _put(sum_top, 1, "SIZE", bold_=True, fill_=HDR_FILL)
#     for i, sz in enumerate(full_sizes):
#         _put(sum_top, 2 + i, sz, bold_=True, fill_=HDR_FILL)

#     gt_col = 2 + n_sizes
#     gw_col = gt_col + 1
#     nw_col = gt_col + 2

#     _put(sum_top, gt_col, "G Total", bold_=True, fill_=HDR_FILL)
#     _put(sum_top, gw_col, "Grand Total gross weight kg", bold_=True, fill_=HDR_FILL)
#     _put(sum_top, nw_col, "Grand Total Net weight kg",  bold_=True, fill_=HDR_FILL)

#     def _summary_row(rr, label, by_size, total_val, fill=None, numfmt=None):
#         _put(rr, 1, label, bold_=True, fill_=fill)
#         for i, sz in enumerate(full_sizes):
#             _put(rr, 2 + i, by_size.get(sz, ""), fill_=fill, numfmt=numfmt)
#         _put(rr, gt_col, total_val, bold_=True, fill_=fill, numfmt=numfmt)

#     cut_qty = export_state.get("cut_qty", {})
#     order_by = export_state.get("order_qty_by_size", {})
#     ship_by  = export_state.get("ship_qty_by_size", {})
#     exs_by   = export_state.get("exs_by_size", {})
#     pct_by   = export_state.get("pct_by_size", {})

#     _summary_row(sum_top + 1, "CUT QTY", cut_qty, export_state.get("cut_total", 0), fill=yellow)
#     _summary_row(sum_top + 2, "ORDER QTY", order_by, export_state.get("order_total", 0), fill=grey)
#     _summary_row(sum_top + 3, "SHIP QTY", ship_by, export_state.get("ship_total", 0))
#     _summary_row(sum_top + 4, "EXS/SHT QTY", exs_by, export_state.get("exs_total", 0))

#     pct_row = sum_top + 5
#     _put(pct_row, 1, "PERCENTAGE", bold_=True)
#     for i, sz in enumerate(full_sizes):
#         v = pct_by.get(sz)
#         _put(pct_row, 2 + i, "" if v is None else float(v), numfmt="0.00%")
#     v_total = export_state.get("pct_total")
#     _put(pct_row, gt_col, "" if v_total is None else float(v_total),
#          bold_=True, numfmt="0.00%")

#     ws.merge_cells(start_row=sum_top + 1, start_column=gw_col,
#                    end_row=pct_row, end_column=gw_col)
#     _put(sum_top + 1, gw_col, round(export_state.get("total_grs", 0), 2),
#          bold_=True, numfmt="0.00")
#     for rrr in range(sum_top + 1, pct_row + 1):
#         ws.cell(row=rrr, column=gw_col).border = BORDER

#     ws.merge_cells(start_row=sum_top + 1, start_column=nw_col,
#                    end_row=pct_row, end_column=nw_col)
#     _put(sum_top + 1, nw_col, round(export_state.get("total_net", 0), 2),
#          bold_=True, numfmt="0.00")
#     for rrr in range(sum_top + 1, pct_row + 1):
#         ws.cell(row=rrr, column=nw_col).border = BORDER

#     # ============ FOOTER ============
#     foot_top = pct_row + 3
#     _put(foot_top,     1, "GROSS WEIGHT :", bold_=True, border_=False)
#     _put(foot_top,     2, round(export_state.get("total_grs", 0), 2),
#          bold_=True, border_=False)
#     _put(foot_top + 1, 1, "NET WEIGHT :",   bold_=True, border_=False)
#     _put(foot_top + 1, 2, round(export_state.get("total_net", 0), 2),
#          bold_=True, border_=False)
#     _put(foot_top + 2, 1, "CTN MEAS :",     bold_=True, border_=False)
#     _put(foot_top + 2, 2, export_state.get("ctn_meas_1", ""), border_=False)
#     _put(foot_top + 3, 2, export_state.get("ctn_meas_2", ""), border_=False)
#     _put(foot_top + 4, 1, "CBM :",          bold_=True, border_=False)
#     _put(foot_top + 4, 2, round(export_state.get("cbm_total", 0), 8), border_=False)

#     # ---- Signature ----
#     sign_top = foot_top + 7
#     _put(sign_top, 1, "INCHARGE", bold_=True, border_=False)
#     _put(sign_top, 3, "FM",       bold_=True, border_=False)
#     ws.cell(row=sign_top, column=1).border = Border(top=Side(style="thin"))
#     ws.cell(row=sign_top, column=3).border = Border(top=Side(style="thin"))

#     # ============ COLUMN WIDTHS ============
#     ws.column_dimensions[get_column_letter(C_CASE_A)].width = 8
#     ws.column_dimensions[get_column_letter(C_CASE_B)].width = 8
#     ws.column_dimensions[get_column_letter(C_COLOR)].width = 16
#     ws.column_dimensions[get_column_letter(C_UPC)].width = 14
#     for i in range(n_sizes):
#         ws.column_dimensions[get_column_letter(C_SZ0 + i)].width = 9
#     ws.column_dimensions[get_column_letter(C_CTN_PCS)].width = 9
#     ws.column_dimensions[get_column_letter(C_TOT_CTN)].width = 10
#     ws.column_dimensions[get_column_letter(C_TOT_PCS)].width = 10
#     ws.column_dimensions[get_column_letter(C_GRS_CTN)].width = 12
#     ws.column_dimensions[get_column_letter(C_NET_CTN)].width = 12
#     ws.column_dimensions[get_column_letter(C_TOT_GRS)].width = 12
#     ws.column_dimensions[get_column_letter(C_TOT_NET)].width = 12

#     # --- ADDED: freeze header rows + hide gridlines ---
#     ws.freeze_panes = ws.cell(row=table_top + 2, column=C_SZ0).coordinate
#     ws.sheet_view.showGridLines = False
#     # --- END ADDED ---

#     return wb


# def render_packing_list_preview(line, sublines):
#     """Interactive packing-list preview with wider Color & UPC columns."""
#     style_key = line["style"]
#     full_sizes = [s["size"] for s in sublines]
#     n_sizes = len(full_sizes)

#     default_color = next((s["color"] for s in sublines if s.get("color")), "MIXED")

#     _pdf_pack = int(sublines[0].get("items_per_outer_pack") or 1) if sublines else 1
#     _saved_pack_max = st.session_state.get(f"pack_max_v4_{style_key}")
#     effective_pack = int(_saved_pack_max) if _saved_pack_max is not None else _pdf_pack

#     rows_key = f"carton_rows_v8_{style_key}"
#     pack_sig_key = f"pack_sig_v8_{style_key}"
#     _pack_changed = st.session_state.get(pack_sig_key) != effective_pack
#     if rows_key not in st.session_state or _pack_changed:
#         base = compute_carton_rows(sublines, pack_override=effective_pack)
#         st.session_state[rows_key] = [
#             {
#                 "id": uuid.uuid4().hex,
#                 "color": (r["color"] if r["color"] is not None else default_color),
#                 "upc": r["upc"] if not r["mixed"] else "MIXED",
#                 "sizes": {sz: int(r["sizes"].get(sz, 0) or 0) for sz in full_sizes},
#                 "total_ctn": int(r["total_ctn"]),
#                 "mixed": bool(r["mixed"]),
#             }
#             for r in base
#         ]
#         st.session_state[pack_sig_key] = effective_pack

#     carton_rows = st.session_state[rows_key]

#     nw_by_size, nnw_by_size, empty_by_size = render_size_spec_table(style_key, full_sizes)
#     empty_ctn_wt = float(empty_by_size[full_sizes[0]]) if full_sizes else 0.0

#     # ---- CSS (Color = 5th column, UPC = 6th column are wider) ----
#     st.markdown(
#         f"""
#         <style>
#         .st-key-carton_grid {{
#             width: 100% !important;
#             max-width: 100% !important;

#             overflow-x: auto !important;
#             box-sizing: border-box !important;
#             border: 1px solid #cbd5e1;
#             border-radius: 6px;
#             background: white;
#         }}

#         .st-key-carton_grid [data-testid="stHorizontalBlock"] {{
#             display: flex !important;
#             flex-wrap: nowrap !important;
#             gap: 0 !important;
#             width: max-content !important;
#             min-width: max-content !important;
#             max-width: none !important;
#             flex-shrink: 0 !important;
#         }}

#         .st-key-carton_grid [data-testid="stColumn"] {{
#             flex: 0 0 {CELL_WIDTH}px !important;
#             width: {CELL_WIDTH}px !important;
#             min-width: {CELL_WIDTH}px !important;
#             max-width: {CELL_WIDTH}px !important;
#             box-sizing: border-box !important;
#             padding-left: 0 !important;
#             padding-right: 0 !important;
#             margin: 0 !important;
#             flex-shrink: 0 !important;
#         }}

#         .st-key-carton_grid [data-testid="stHorizontalBlock"] [data-testid="stColumn"]:nth-child(1),
#         .st-key-carton_grid [data-testid="stHorizontalBlock"] [data-testid="stColumn"]:nth-child(2) {{
#             flex: 0 0 {ACTION_WIDTH}px !important;
#             width: {ACTION_WIDTH}px !important;
#             min-width: {ACTION_WIDTH}px !important;
#             max-width: {ACTION_WIDTH}px !important;
#         }}

#         .st-key-carton_grid [data-testid="stHorizontalBlock"] [data-testid="stColumn"]:nth-child(5) {{
#             flex: 0 0 {COLOR_WIDTH}px !important;
#             width: {COLOR_WIDTH}px !important;
#             min-width: {COLOR_WIDTH}px !important;
#             max-width: {COLOR_WIDTH}px !important;
#         }}

#         .st-key-carton_grid [data-testid="stHorizontalBlock"] [data-testid="stColumn"]:nth-child(6) {{
#             flex: 0 0 {UPC_WIDTH}px !important;
#             width: {UPC_WIDTH}px !important;
#             min-width: {UPC_WIDTH}px !important;
#             max-width: {UPC_WIDTH}px !important;
#         }}

#         .grid-header {{
#             height: 36px !important;
#             min-height: 38px !important;
#             max-height: 38px !important;
#             box-sizing: border-box !important;
#             display: flex;
#             align-items: center;
#             justify-content: center;
#             background: #1f2937;
#             color: white;
#             border: 1px solid #4b5563 !important;
#             font-size: 12px;
#             font-weight: 600;
#             white-space: nowrap;
#             overflow: hidden;
#         }}

#         .grid-header-action {{
#             width: {ACTION_WIDTH}px !important;
#             min-width: {ACTION_WIDTH}px !important;
#             max-width: {ACTION_WIDTH}px !important;
#         }}

#         .grid-header-color {{
#             width: {COLOR_WIDTH}px !important;
#             min-width: {COLOR_WIDTH}px !important;
#             max-width: {COLOR_WIDTH}px !important;
#         }}

#         .grid-header-upc {{
#             width: {UPC_WIDTH}px !important;
#             min-width: {UPC_WIDTH}px !important;
#             max-width: {UPC_WIDTH}px !important;
#         }}

#         .st-key-carton_grid [data-testid="stButton"] {{
#             width: {ACTION_WIDTH}px !important;
#             min-width: {ACTION_WIDTH}px !important;
#             max-width: {ACTION_WIDTH}px !important;
#             margin: 0 !important;
#             padding: 0 !important;
#         }}

#         .st-key-carton_grid [data-testid="stButton"] button {{
#             width: {ACTION_WIDTH}px !important;
#             min-width: {ACTION_WIDTH}px !important;
#             max-width: {ACTION_WIDTH}px !important;
#             height: {CELL_HEIGHT}px !important;
#             min-height: {CELL_HEIGHT}px !important;
#             max-height: {CELL_HEIGHT}px !important;
#             padding: 0 !important;
#             margin: 0 !important;
#             border-radius: 0 !important;
#             font-size: 16px !important;
#         }}

#         .st-key-carton_grid [data-testid="stNumberInput"] {{
#             width: {CELL_WIDTH}px !important;
#             min-width: {CELL_WIDTH}px !important;
#             max-width: {CELL_WIDTH}px !important;
#             margin: 0 !important;
#             padding: 0 !important;
#         }}

#         .st-key-carton_grid [data-testid="stNumberInput"] input {{
#             width: {CELL_WIDTH}px !important;
#             min-width: {CELL_WIDTH}px !important;
#             max-width: {CELL_WIDTH}px !important;
#             height: {CELL_HEIGHT}px !important;
#             min-height: {CELL_HEIGHT}px !important;
#             max-height: {CELL_HEIGHT}px !important;
#             box-sizing: border-box !important;
#             border-radius: 0 !important;
#             margin: 0 !important;
#             padding: 0 4px !important;
#             font-size: 12px !important;
#             text-align: center !important;
#         }}

#         .st-key-carton_grid [data-testid="stNumberInput"] label {{
#             display: none !important;
#         }}

#         .data-cell {{
#             height: 34px !important;
#             min-height: 34px !important;
#             max-height: 34px !important;
#             box-sizing: border-box !important;
#             display: flex;
#             align-items: center;
#             justify-content: center;
#             border: 1px solid #d1d5db !important;
#             background: #fff;
#             font-size: 12px;
#             white-space: nowrap;
#             overflow: hidden;
#         }}

#         .data-cell.mixed {{
#             background: #FFFF66;
#         }}

#         .total-cell {{
#             height: 36px !important;
#             min-height: 36px !important;
#             max-height: 36px !important;
#             box-sizing: border-box !important;
#             display: flex;
#             align-items: center;
#             justify-content: center;
#             border: 1px solid #9ca3af !important;
#             background: #e5e7eb;
#             color: #111827;
#             font-weight: 700;
#             font-size: 12px;
#             white-space: nowrap;
#             overflow: hidden;
#         }}

#         .total-action-cell {{
#             width: {ACTION_WIDTH}px !important;
#             min-width: {ACTION_WIDTH}px !important;
#             max-width: {ACTION_WIDTH}px !important;
#             height: {TOTAL_HEIGHT}px !important;
#             min-height: {TOTAL_HEIGHT}px !important;
#             max-height: {TOTAL_HEIGHT}px !important;
#             box-sizing: border-box !important;
#             display: flex;
#             align-items: center;
#             justify-content: center;
#             background: #e5e7eb;
#             border-right: 1px solid #9ca3af;
#             border-bottom: 2px solid #6b7280;
#             font-weight: 700;
#         }}
#         </style>
#         """,
#         unsafe_allow_html=True,
#     )

#     st.caption(
#         "Main carton table is fully editable. Color & UPC columns are wider. "
#         "Change size quantities or TOTAL CTN directly in the cells."
#     )

#     # ---- Handle insert / delete ----
#     insert_key = f"insert_after_v8_{style_key}"
#     delete_key = f"delete_row_v8_{style_key}"

#     if insert_key in st.session_state and st.session_state[insert_key] is not None:
#         idx = st.session_state[insert_key]
#         new_row = {
#             "id": uuid.uuid4().hex,
#             "color": default_color,
#             "upc": "MIXED",
#             "sizes": {sz: 0 for sz in full_sizes},
#             "total_ctn": 1,
#             "mixed": True,
#         }
#         carton_rows.insert(idx + 1, new_row)
#         st.session_state[rows_key] = carton_rows
#         st.session_state[insert_key] = None
#         st.rerun()

#     if delete_key in st.session_state and st.session_state[delete_key] is not None:
#         idx = st.session_state[delete_key]
#         if len(carton_rows) > 1 and 0 <= idx < len(carton_rows):
#             carton_rows.pop(idx)
#             st.session_state[rows_key] = carton_rows
#         st.session_state[delete_key] = None
#         st.rerun()

#     n_rows = len(carton_rows)

#     col_spec = (
#         [ACTION_WIDTH, ACTION_WIDTH] +
#         [CELL_WIDTH, CELL_WIDTH] +
#         [COLOR_WIDTH, UPC_WIDTH] +
#         [CELL_WIDTH] * n_sizes +
#         [CELL_WIDTH] * 7
#     )

#     # ---- Order quantity by size table ----
#     o_hdr = "".join(f"<th>{escape(str(s['size']))}</th>" for s in sublines)
#     o_color = "".join(f"<td>{escape(str(s['color']))}</td>" for s in sublines)
#     o_upc = "".join(f"<td>{escape(str(s['upc']))}</td>" for s in sublines)
#     o_qty = "".join(f"<td><b>{s['qty']}</b></td>" for s in sublines)

#     st.markdown("##### Order quantity by size")
#     st.markdown(
#         '<style>'
#         '.order-scroll{overflow-x:auto;-webkit-overflow-scrolling:touch;'
#         'border:1px solid #ccc;border-radius:4px;margin:8px 0 12px 0;max-width:100%;background:#fff;}'
#         'table.order-tbl{border-collapse:collapse;width:max-content;min-width:100%;'
#         'font-family:Arial,sans-serif;font-size:12px;color:#000;white-space:nowrap;}'
#         'table.order-tbl th,table.order-tbl td{border:1px solid #333;padding:4px 8px;'
#         'text-align:center;vertical-align:middle;background:#fff;}'
#         'table.order-tbl th{background:#D9D9D9;font-weight:bold;}'
#         'table.order-tbl .rowlbl{position:sticky;left:0;z-index:1;background:#D9D9D9;'
#         'font-weight:bold;text-align:left;}'
#         '</style>'
#         '<div class="order-scroll"><table class="order-tbl">'
#         f'<tr><th class="rowlbl">SIZE</th>{o_hdr}</tr>'
#         f'<tr><td class="rowlbl">COLOR</td>{o_color}</tr>'
#         f'<tr><td class="rowlbl">UPC</td>{o_upc}</tr>'
#         f'<tr><td class="rowlbl">ORDER QTY</td>{o_qty}</tr>'
#         '</table></div>',
#         unsafe_allow_html=True,
#     )

#     # ---- Packing qty MIN / MAX table ----
#     pdf_pack = int(sublines[0].get("items_per_outer_pack") or 1) if sublines else 1

#     st.markdown("##### Packing qty (per carton)")
#     pk_w = [1.2, 1, 1, 2.2]
#     pk_hdr = st.columns(pk_w)
#     pk_hdr[0].markdown("*PACKING QTY*")
#     pk_hdr[1].markdown("*MIN*")
#     pk_hdr[2].markdown("*MAX*")
#     pk_hdr[3].markdown("*PDF (Items Per Outer Pack)*")
#     pk_row = st.columns(pk_w)
#     pk_row[0].markdown("Input")
#     pack_min = pk_row[1].number_input(
#         "Packing qty MIN", min_value=1, value=None, step=1, placeholder="empty",
#         key=f"pack_min_v4_{style_key}", label_visibility="collapsed",
#     )
#     pack_max = pk_row[2].number_input(
#         "Packing qty MAX", min_value=1, value=None, step=1, placeholder="empty",
#         key=f"pack_max_v4_{style_key}", label_visibility="collapsed",
#     )
#     pk_row[3].markdown(f"*{pdf_pack}* (used when MAX is empty)")

#     # --- ADDED: Reset button ---
#     _rh1, _rh2 = st.columns([1, 6])
#     with _rh1:
#         if st.button("🔄 Reset", key=f"reset_rows_v8_{style_key}",
#                      help="Restore carton rows to PDF defaults (wipes manual edits)"):
#             st.session_state.pop(rows_key, None)
#             st.session_state.pop(pack_sig_key, None)
#             for _k in list(st.session_state.keys()):
#                 if (_k.startswith(f"sz_v8_{style_key}_")
#                         or _k.startswith(f"tctn_v8_{style_key}_")
#                         or _k.startswith(f"add_v8_{style_key}_")
#                         or _k.startswith(f"del_v8_{style_key}_")):
#                     st.session_state.pop(_k, None)
#             st.rerun()
#     # --- END ADDED ---

#     st.markdown("##### Main carton table (editable)")

#     edited_sizes = []
#     edited_ctn_pcs = []
#     edited_total_ctn = []
#     grs_wts, net_wts = [], []
#     row_total_pcs = []
#     ctn_ranges = []
#     row_colors, row_upcs = [], []
#     cum = 0

#     with st.container(border=True, key="carton_grid"):

#         header_cols = st.columns(col_spec, gap="small")

#         with header_cols[0]:
#             st.markdown('<div class="grid-header grid-header-action">+</div>', unsafe_allow_html=True)
#         with header_cols[1]:
#             st.markdown('<div class="grid-header grid-header-action">−</div>', unsafe_allow_html=True)
#         with header_cols[2]:
#             st.markdown('<div class="grid-header">#</div>', unsafe_allow_html=True)
#         with header_cols[3]:
#             st.markdown('<div class="grid-header">Ctn No</div>', unsafe_allow_html=True)
#         with header_cols[4]:
#             st.markdown('<div class="grid-header grid-header-color">Color</div>', unsafe_allow_html=True)
#         with header_cols[5]:
#             st.markdown('<div class="grid-header grid-header-upc">UPC</div>', unsafe_allow_html=True)

#         for s_idx, sz in enumerate(full_sizes):
#             with header_cols[6 + s_idx]:
#                 st.markdown(f'<div class="grid-header">{sz}</div>', unsafe_allow_html=True)

#         extra_headers = ["CTN PCS", "TOTAL CTN", "TOTAL PCS", "Grs/ctn", "Net/ctn", "tot Grs", "tot Net"]
#         for e_idx, h in enumerate(extra_headers):
#             with header_cols[6 + n_sizes + e_idx]:
#                 st.markdown(f'<div class="grid-header">{h}</div>', unsafe_allow_html=True)

#         for row_index, row in enumerate(carton_rows):
#             row_id = row["id"]
#             row_cols = st.columns(col_spec, gap="small")

#             with row_cols[0]:
#                 if st.button("＋", key=f"add_v8_{style_key}_{row_id}",
#                              help="Add row below"):
#                     st.session_state[insert_key] = row_index
#                     st.rerun()

#             with row_cols[1]:
#                 if st.button("❌", key=f"del_v8_{style_key}_{row_id}",
#                              help="Delete this row",
#                              disabled=n_rows <= 1):
#                     st.session_state[delete_key] = row_index
#                     st.rerun()

#             with row_cols[2]:
#                 st.markdown(f'<div class="data-cell">{row_index + 1}</div>', unsafe_allow_html=True)

#             current_sizes = {}
#             for s_idx, sz in enumerate(full_sizes):
#                 with row_cols[6 + s_idx]:
#                     sz_key = f"sz_v8_{style_key}_{row_id}_{sz}"
#                     current_val = int(row["sizes"].get(sz, 0) or 0)
#                     new_val = st.number_input(
#                         f"sz_{row_id}_{sz}",
#                         min_value=0,
#                         value=current_val,
#                         step=1,
#                         key=sz_key,
#                         label_visibility="collapsed",
#                     )
#                     current_sizes[sz] = int(new_val)
#                     carton_rows[row_index]["sizes"][sz] = int(new_val)

#             with row_cols[7 + n_sizes]:
#                 tctn_key = f"tctn_v8_{style_key}_{row_id}"
#                 current_tctn = int(row.get("total_ctn") or 0)
#                 new_tctn = st.number_input(
#                     f"tctn_{row_id}",
#                     min_value=0,
#                     value=current_tctn,
#                     step=1,
#                     key=tctn_key,
#                     label_visibility="collapsed",
#                 )
#                 carton_rows[row_index]["total_ctn"] = int(new_tctn)
#                 total_ctn_val = int(new_tctn)

#             size_sum = sum(current_sizes.values())
#             ctn_start = cum + 1
#             cum += total_ctn_val
#             ctn_end = cum
#             ctn_range_str = f"{ctn_start}-{ctn_end}" if total_ctn_val > 0 else "-"

#             piece_wt = sum(
#                 float(current_sizes.get(sz, 0) or 0) * float(nw_by_size.get(sz, 0.0) or 0.0)
#                 for sz in full_sizes
#             )
#             grs = piece_wt + empty_ctn_wt
#             net = piece_wt
#             tot_pcs = total_ctn_val * size_sum

#             edited_sizes.append(current_sizes)
#             edited_ctn_pcs.append(size_sum)
#             edited_total_ctn.append(total_ctn_val)
#             grs_wts.append(grs)
#             net_wts.append(net)
#             row_total_pcs.append(tot_pcs)
#             ctn_ranges.append(ctn_range_str)
#             row_colors.append(row.get("color") or default_color)
#             row_upcs.append(row.get("upc") or "MIXED")

#             with row_cols[3]:
#                 st.markdown(f'<div class="data-cell">{ctn_range_str}</div>', unsafe_allow_html=True)

#             mixed_cls = " mixed" if row.get("mixed") else ""
#             with row_cols[4]:
#                 st.markdown(
#                     f'<div class="data-cell{mixed_cls}">{row.get("color") or default_color}</div>',
#                     unsafe_allow_html=True
#                 )

#             with row_cols[5]:
#                 st.markdown(
#                     f'<div class="data-cell">{row.get("upc") or "MIXED"}</div>',
#                     unsafe_allow_html=True
#                 )

#             with row_cols[6 + n_sizes]:
#                 st.markdown(f'<div class="data-cell"><b>{size_sum}</b></div>', unsafe_allow_html=True)

#             with row_cols[8 + n_sizes]:
#                 st.markdown(f'<div class="data-cell"><b>{tot_pcs}</b></div>', unsafe_allow_html=True)

#             with row_cols[9 + n_sizes]:
#                 st.markdown(f'<div class="data-cell">{grs:.2f}</div>', unsafe_allow_html=True)

#             with row_cols[10 + n_sizes]:
#                 st.markdown(f'<div class="data-cell">{net:.2f}</div>', unsafe_allow_html=True)

#             with row_cols[11 + n_sizes]:
#                 st.markdown(f'<div class="data-cell">{total_ctn_val * grs:.2f}</div>', unsafe_allow_html=True)

#             with row_cols[12 + n_sizes]:
#                 st.markdown(f'<div class="data-cell">{total_ctn_val * net:.2f}</div>', unsafe_allow_html=True)

#         total_ctn = sum(edited_total_ctn) if edited_total_ctn else 0
#         total_pcs = sum(row_total_pcs) if row_total_pcs else 0
#         total_grs = sum(tc * g for tc, g in zip(edited_total_ctn, grs_wts)) if edited_total_ctn else 0.0
#         total_net = sum(tc * n for tc, n in zip(edited_total_ctn, net_wts)) if edited_total_ctn else 0.0

#         size_packet_totals = {
#             sz: sum(int(edited_sizes[i].get(sz, 0) or 0) * edited_total_ctn[i] for i in range(n_rows))
#             for sz in full_sizes
#         } if n_rows else {sz: 0 for sz in full_sizes}

#         total_cols = st.columns(col_spec, gap="small")

#         with total_cols[0]:
#             st.markdown('<div class="total-action-cell"></div>', unsafe_allow_html=True)
#         with total_cols[1]:
#             st.markdown('<div class="total-action-cell">TOTAL</div>', unsafe_allow_html=True)
#         with total_cols[2]:
#             st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)
#         with total_cols[3]:
#             st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)
#         with total_cols[4]:
#             st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)
#         with total_cols[5]:
#             st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)

#         for s_idx, sz in enumerate(full_sizes):
#             with total_cols[6 + s_idx]:
#                 st.markdown(f'<div class="total-cell">{size_packet_totals.get(sz, 0)}</div>', unsafe_allow_html=True)

#         with total_cols[6 + n_sizes]:
#             st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)
#         with total_cols[7 + n_sizes]:
#             st.markdown(f'<div class="total-cell">{total_ctn}</div>', unsafe_allow_html=True)
#         with total_cols[8 + n_sizes]:
#             st.markdown(f'<div class="total-cell">{total_pcs}</div>', unsafe_allow_html=True)
#         with total_cols[9 + n_sizes]:
#             st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)
#         with total_cols[10 + n_sizes]:
#             st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)
#         with total_cols[11 + n_sizes]:
#             st.markdown(f'<div class="total-cell">{total_grs:.2f}</div>', unsafe_allow_html=True)
#         with total_cols[12 + n_sizes]:
#             st.markdown(f'<div class="total-cell">{total_net:.2f}</div>', unsafe_allow_html=True)

#     st.session_state[rows_key] = carton_rows

#     st.markdown("**Live totals**")
#     m1, m2, m3, m4 = st.columns(4)
#     m1.metric("Total Cartons", total_ctn)
#     m2.metric("Total Pieces", total_pcs)
#     m3.metric("Total Gross Wt", f"{total_grs:.2f}")
#     m4.metric("Total Net Wt", f"{total_net:.2f}")

#     # --- ADDED: Excess / Short warning banner ---
#     _order_total_now = sum(s["qty"] for s in sublines)
#     _ship_total_now = sum(
#         int(edited_sizes[_i].get(_sz, 0) or 0) * edited_total_ctn[_i]
#         for _i in range(n_rows) for _sz in full_sizes
#     )
#     _exs_now = _ship_total_now - _order_total_now
#     if _exs_now > 0:
#         st.success(f"✅ Over-ship: **+{_exs_now}** pieces")
#     elif _exs_now < 0:
#         st.error(f"⚠️ Short-ship: **{_exs_now}** pieces")
#     else:
#         st.info("Ship qty matches order qty exactly.")
#     # --- END ADDED ---

#     with st.expander(f"Edit cut qty / CTN MEAS / CBM for {style_key}", expanded=False):
#         cut_qty = {}
#         cut_cols = st.columns(len(full_sizes)) if full_sizes else []
#         for col, size in zip(cut_cols, full_sizes):
#             with col:
#                 cut_qty[size] = st.number_input(
#                     f"Cut qty — {size}", value=0, key=f"cut_v8_{style_key}_{size}"
#                 )

#         ctn_meas_1 = st.text_input("Ctn meas — line 1", "", key=f"meas1_v8_{style_key}")

#         st.markdown("**Ctn meas — line 2 (inches)**")
#         d1, d2, d3 = st.columns(3)
#         ctn_l = d1.number_input(
#             "Length (L)", min_value=0.0, value=24.0, step=0.5, format="%.2f",
#             key=f"ctn_l_v9_{style_key}",
#         )
#         ctn_w = d2.number_input(
#             "Width (W)", min_value=0.0, value=16.0, step=0.5, format="%.2f",
#             key=f"ctn_w_v9_{style_key}",
#         )
#         ctn_h = d3.number_input(
#             "Height (H)", min_value=0.0, value=10.0, step=0.5, format="%.2f",
#             key=f"ctn_h_v9_{style_key}",
#         )

#         CUIN_PER_CBM = 61000
#         ctn_meas_2 = f"L {ctn_l:g}'' X W {ctn_w:g}'' X H {ctn_h:g}''"
#         cbm_per_ctn = (ctn_l * ctn_w * ctn_h) / CUIN_PER_CBM

#         st.caption(
#             f"CTN MEAS: {ctn_meas_2}  |  CBM per carton: {cbm_per_ctn:.8f}  |  "
#             f"CBM = ({ctn_l:g} × {ctn_w:g} × {ctn_h:g} × total cartons) / {CUIN_PER_CBM}"
#         )

#     ship_qty_by_size = {sz: 0 for sz in full_sizes}
#     for size_vals, tc in zip(edited_sizes, edited_total_ctn):
#         for sz in full_sizes:
#             ship_qty_by_size[sz] += int(size_vals.get(sz, 0) or 0) * tc

#     order_qty_by_size = {s["size"]: s["qty"] for s in sublines}
#     order_total = sum(order_qty_by_size.get(sz, 0) for sz in full_sizes)
#     ship_total = sum(ship_qty_by_size.get(sz, 0) for sz in full_sizes)
#     cut_total = sum(cut_qty.get(sz, 0) for sz in full_sizes)
#     exs_by_size = {sz: ship_qty_by_size.get(sz, 0) - order_qty_by_size.get(sz, 0) for sz in full_sizes}
#     exs_total = ship_total - order_total
#     pct_by_size = {
#         sz: (exs_by_size[sz] / order_qty_by_size[sz]) if order_qty_by_size.get(sz) else None
#         for sz in full_sizes
#     }
#     pct_total = (exs_total / order_total) if order_total else None
#     cbm_total = cbm_per_ctn * total_ctn

#     st.markdown('<div class="plist-wrap">', unsafe_allow_html=True)

#     st.markdown(f"""
# <table class="plist noborder">
# <tr class="title-row">
#   <td style="width:16%"><b>BUYER</b></td><td style="width:14%">VANS</td>
#   <td style="width:14%"></td>
#   <td style="width:16%"><b>ORDER QTY</b></td><td style="width:8%">{line['order_qty']}</td><td style="width:6%">PCS</td>
#   <td style="width:8%"></td>
#   <td style="width:6%"><u><b>CRD</b></u></td><td>{line.get('crd','')}</td>
# </tr>
# <tr class="title-row">
#   <td><b>LOT NO</b></td><td>{line['style']}</td><td></td>
#   <td><b>PACK QTY</b></td><td>{line['order_qty']}</td><td>PCS</td><td></td><td></td><td></td>
# </tr>
# <tr class="title-row">
#   <td><b>P.O NO</b></td><td>{line['po_line_no']}</td><td></td>
#   <td><b>EXCESS/SHORT QTY</b></td><td>{exs_total}</td><td>PCS</td><td></td><td></td><td></td>
# </tr>
# <tr class="title-row">
#   <td colspan="2" class="desc-cell">DESCRIPTION:<br>{line['description']}</td><td></td>
#   <td><b>PERCENTAGE</b></td>
#   <td>{'-' if pct_total is None else f'{pct_total*100:.2f}'}</td><td>%</td><td></td>
#   <td colspan="2">{line.get('destination_country','')}</td>
# </tr>
# </table>
# """, unsafe_allow_html=True)

#     def fmt_pct(v):
#         return "" if v is None else f"{v*100:.2f}%"

#     n_data_rows = 5
#     size_th = "".join(f"<th>{sz}</th>" for sz in full_sizes)

#     summary_body = ""
#     for i, (label, vals, gt, cls) in enumerate([
#         ("CUT QTY", [cut_qty.get(sz, "") for sz in full_sizes], cut_total, "yellow"),
#         ("ORDER QTY", [order_qty_by_size.get(sz, "") for sz in full_sizes], order_total, "gray"),
#         ("SHIP QTY", [ship_qty_by_size.get(sz, "") for sz in full_sizes], ship_total, ""),
#         ("EXS/SHT QTY", [exs_by_size.get(sz, "") for sz in full_sizes], exs_total, ""),
#         ("PERCENTAGE", [fmt_pct(pct_by_size.get(sz)) for sz in full_sizes], fmt_pct(pct_total), ""),
#     ]):
#         cells = "".join(f'<td class="{cls}">{v}</td>' for v in vals)
#         if i == 0:
#             weight_cells = (
#                 f'<td class="bold" rowspan="{n_data_rows}" style="vertical-align:middle;">'
#                 f'{total_grs:.2f}</td>'
#                 f'<td class="bold" rowspan="{n_data_rows}" style="vertical-align:middle;">'
#                 f'{total_net:.2f}</td>'
#             )
#         else:
#             weight_cells = ""
#         summary_body += (
#             f'<tr><td class="{cls} bold">{label}</td>{cells}'
#             f'<td class="bold">{gt}</td>{weight_cells}</tr>'
#         )

#     st.markdown(f"""
# <table class="plist">
# <tr class="hdr">
#   <th>SIZE</th>{size_th}
#   <th>G Total</th>
#   <th>Grand Total gross weight kg</th>
#   <th>Grand Total Net weight kg</th>
# </tr>
# {summary_body}
# </table>
# """, unsafe_allow_html=True)

#     st.markdown(f"""
# <table class="plist noborder">
# <tr class="title-row"><td style="width:15%"><b>GROSS WEIGHT :</b></td><td>{total_grs:.2f}</td></tr>
# <tr class="title-row"><td><b>NET WEIGHT :</b></td><td>{total_net:.2f}</td></tr>
# <tr class="title-row"><td><b>CTN MEAS :</b></td><td>{ctn_meas_1}<br>{ctn_meas_2}</td></tr>
# <tr class="title-row"><td><b>CBM :</b></td><td>{cbm_total:.8f}</td></tr>
# </table>
# <br><br>
# <table class="plist noborder">
# <tr class="title-row"><td style="width:30%; border-top:1px solid #000; text-align:center;">INCHARGE</td>
# <td style="width:30%"></td>
# <td style="width:30%; border-top:1px solid #000; text-align:center;">FM</td></tr>
# </table>
# """, unsafe_allow_html=True)

#     st.markdown('</div>', unsafe_allow_html=True)

#     rows_export = []
#     for i in range(n_rows):
#         rows_export.append({
#             "ctn_no": ctn_ranges[i],
#             "color": row_colors[i],
#             "upc": row_upcs[i],
#             "sizes": edited_sizes[i],
#             "ctn_pcs": edited_ctn_pcs[i],
#             "total_ctn": edited_total_ctn[i],
#             "grs_per_ctn": grs_wts[i],
#             "net_per_ctn": net_wts[i],
#             "total_pcs": row_total_pcs[i],
#             "total_grs": edited_total_ctn[i] * grs_wts[i],
#             "total_net": edited_total_ctn[i] * net_wts[i],
#         })

#     export_state = {
#         "line": line,
#         "full_sizes": full_sizes,
#         "rows": rows_export,
#         "total_ctn": total_ctn,
#         "total_pcs": total_pcs,
#         "total_grs": total_grs,
#         "total_net": total_net,
#         "size_packet_totals": size_packet_totals,
#         "cut_qty": cut_qty,
#         "order_qty_by_size": order_qty_by_size,
#         "ship_qty_by_size": ship_qty_by_size,
#         "exs_by_size": exs_by_size,
#         "pct_by_size": pct_by_size,
#         "cut_total": cut_total,
#         "order_total": order_total,
#         "ship_total": ship_total,
#         "exs_total": exs_total,
#         "pct_total": pct_total,
#         "ctn_meas_1": ctn_meas_1,
#         "ctn_meas_2": ctn_meas_2,
#         "cbm_per_ctn": cbm_per_ctn,
#         "cbm_total": cbm_total,
#         "nw_by_size": nw_by_size,
#         "nnw_by_size": nnw_by_size,
#         "empty_by_size": empty_by_size,
#     }

#     st.markdown("##### Download this packing list (with your edits)")
#     dl_col1, dl_col2 = st.columns(2)

#     with dl_col1:
#         if st.button("Prepare Excel (.xlsx)", key=f"prep_xlsx_v8_{style_key}"):
#             wb = build_sample_style_workbook(export_state)
#             xbuf = io.BytesIO()
#             wb.save(xbuf)
#             xbuf.seek(0)
#             st.session_state[f"xlsx_bytes_v8_{style_key}"] = xbuf.getvalue()
#             st.session_state[f"xlsx_ts_v8_{style_key}"] = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
#         if st.session_state.get(f"xlsx_bytes_v8_{style_key}"):
#             _ts_x = st.session_state.get(f"xlsx_ts_v8_{style_key}", "latest")
#             st.download_button(
#                 label=f"Download {style_key} packing list (.xlsx)",
#                 data=st.session_state[f"xlsx_bytes_v8_{style_key}"],
#                 file_name=f"packing_list_{style_key}_{_ts_x}.xlsx",
#                 mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
#                 key=f"dl_xlsx_v8_{style_key}",
#             )

#     with dl_col2:
#         if st.button("Prepare PDF", key=f"prep_pdf_v8_{style_key}"):
#             pbuf = build_edited_pdf(export_state)
#             st.session_state[f"pdf_bytes_v8_{style_key}"] = pbuf.getvalue()
#             st.session_state[f"pdf_ts_v8_{style_key}"] = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
#         if st.session_state.get(f"pdf_bytes_v8_{style_key}"):
#             _ts_p = st.session_state.get(f"pdf_ts_v8_{style_key}", "latest")
#             st.download_button(
#                 label=f"Download {style_key} packing list (.pdf)",
#                 data=st.session_state[f"pdf_bytes_v8_{style_key}"],
#                 file_name=f"packing_list_{style_key}_{_ts_p}.pdf",
#                 mime="application/pdf",
#                 key=f"dl_pdf_v8_{style_key}",
#             )


# # ============================================================
# # PART 3 - STREAMLIT UI
# # ============================================================

# st.set_page_config(page_title="PO -> Packing List", layout="wide")
# st.title("PO to Packing List Converter")
# st.write(
#     "Upload a VF/Vans-style Purchase Order PDF. This extracts the line items "
#     "and sizes, then builds a packing-list-style Excel workbook (one sheet per "
#     "PO line) with live formulas for cartons, weights, and totals."
# )

# uploaded_file = st.file_uploader("Upload PO PDF", type=["pdf"])

# if uploaded_file is not None:
#     with st.spinner("Reading PDF..."):
#         try:
#             data = extract_data(uploaded_file)
#         except Exception as e:
#             st.error(f"Failed to read the PDF: {e}")
#             st.stop()

#     if not data["lines"]:
#         st.warning(
#             "No PO lines were found. This PDF's layout may differ from the "
#             "expected format - the extraction patterns may need adjusting."
#         )
#         st.stop()

#     st.success(f"Found {len(data['lines'])} PO line(s) and {len(data['sublines'])} size/color subline(s).")

#     # --- ADDED: clear stale session state when a NEW PDF is uploaded ---
#     _upload_sig = f"{uploaded_file.name}::{uploaded_file.size}"
#     if st.session_state.get("_last_upload_sig") != _upload_sig:
#         for _k in list(st.session_state.keys()):
#             if _k.startswith(("carton_rows_v8_", "pack_sig_v8_",
#                               "xlsx_bytes_v8_", "pdf_bytes_v8_",
#                               "xlsx_ts_v8_", "pdf_ts_v8_",
#                               "insert_after_v8_", "delete_row_v8_")):
#                 st.session_state.pop(_k, None)
#         st.session_state["_last_upload_sig"] = _upload_sig
#     # --- END ADDED ---

#     st.subheader("PO Lines")
#     st.caption("Click a row to filter the sublines below and download just that line's sheet.")
#     po_line_rows = [
#         {
#             "PO Line No.": l["po_line_no"],
#             "Style": l["style"],
#             "Description": l["description"],
#             "Order Qty": l["order_qty"],
#             "Unit Price": l["unit_price"],
#             "Line Cost": l["line_cost"],
#             "CRD": l.get("crd"),
#             "Country": l.get("destination_country"),
#         }
#         for l in data["lines"]
#     ]
#     selection = st.dataframe(
#         po_line_rows,
#         use_container_width=True,
#         on_select="rerun",
#         selection_mode="single-row",
#         key="po_lines_table",
#     )

#     selected_idx = None
#     rows_selected = selection.selection.rows if selection is not None else []
#     if rows_selected:
#         selected_idx = rows_selected[0]
#         selected_line = data["lines"][selected_idx]
#         st.info(f"Selected: **{selected_line['style']}** (PO Line {selected_line['po_line_no']})")

#     st.subheader("Sublines (size / color / UPC)")
#     if selected_idx is not None:
#         filtered_sublines = [s for s in data["sublines"] if s["style"] == selected_line["style"]]
#         st.dataframe(filtered_sublines, use_container_width=True)
#     else:
#         st.dataframe(data["sublines"], use_container_width=True)

#     # --- ADDED: Download extracted JSON ---
#     st.download_button(
#         label="⬇ Download extracted JSON",
#         data=json.dumps(data, indent=2).encode("utf-8"),
#         file_name=f"po_extracted_{len(data['lines'])}_lines.json",
#         mime="application/json",
#         key="dl_po_json",
#     )
#     # --- END ADDED ---

#     col1, col2 = st.columns(2)

#     with col1:
#         if st.button("Build full workbook (all PO lines)", type="primary"):
#             with st.spinner("Building Excel workbook..."):
#                 wb = build_workbook(data)
#                 buf = io.BytesIO()
#                 wb.save(buf)
#                 buf.seek(0)

#             st.success(f"Workbook built with {len(wb.sheetnames)} sheet(s): {', '.join(wb.sheetnames)}")
#             st.download_button(
#                 label="Download packing_list_all_PO_lines.xlsx",
#                 data=buf,
#                 file_name="packing_list_all_PO_lines.xlsx",
#                 mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
#             )

#     with col2:
#         if selected_idx is not None:
#             if st.button(f"Build sheet for {selected_line['style']} only"):
#                 single_line_data = {
#                     "lines": [selected_line],
#                     "sublines": [s for s in data["sublines"] if s["style"] == selected_line["style"]],
#                 }
#                 with st.spinner("Building Excel sheet..."):
#                     wb = build_workbook(single_line_data)
#                     buf = io.BytesIO()
#                     wb.save(buf)
#                     buf.seek(0)

#                 st.success(f"Sheet built for {selected_line['style']}")
#                 st.download_button(
#                     label=f"Download {selected_line['style']}.xlsx",
#                     data=buf,
#                     file_name=f"packing_list_{selected_line['style']}.xlsx",
#                     mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
#                 )
#                 st.session_state["preview_style"] = selected_line["style"]
#         else:
#             st.caption("Select a PO line above to download just that sheet.")

#     if selected_idx is not None and st.session_state.get("preview_style") == selected_line["style"]:
#         st.subheader(f"Packing List Preview — {selected_line['style']}")
#         preview_sublines = [s for s in data["sublines"] if s["style"] == selected_line["style"]]
#         render_packing_list_preview(selected_line, preview_sublines)

#     st.info(
#         "Yellow cells in the downloaded file still need manual entry: "
#         "CUT QTY, CTN MEAS, and CBM factor - these come from your factory's "
#         "cutting/packing records, not the PO. "
#         "CARTON NO., SHIP QTY, GROSS/NET WEIGHT are auto-calculated."
#     )
# else:
#     st.info("Upload a PDF to get started.")




"""
PO -> Packing List Streamlit App
=================================
Upload a VF/Vans-style PO PDF, review the extracted data, and download
a packing-list-style Excel workbook (one sheet per PO line) - all in
the browser, no terminal steps needed.

HOW TO RUN:
    pip install -r requirements.txt
    (requirements.txt should include: streamlit, pdfplumber, openpyxl, reportlab)
    streamlit run app.py
Then open the local URL it prints (usually http://localhost:8501).
"""
import io
import json
import re
import uuid
import datetime

import pdfplumber
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors as rl_colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from html import escape

# ============================================================
# PART 1 - PDF EXTRACTION
# ============================================================

LINE_PATTERN = re.compile(
    r"""
    (?P<po_line_no>\d{15})
    \s+
    (?P<style>[A-Z0-9]+)
    \s+
    (?P<description>[A-Z0-9 \-/]+?)
    \s+
    (?P<order_qty>[\d,]+)
    \s+
    (?:PIECE|EACH)
    \s+
    (?P<unit_price>[\d.]+)
    \s+
    (?P<line_cost>[\d,]+\.\d{2})
    """,
    re.IGNORECASE | re.VERBOSE,
)


SUBLINE_START_PATTERN = re.compile(
    r"(?P<subline_no>\d{12,13})\s+(?P<style>VN[A-Z0-9]+)",
    re.IGNORECASE,
)


def extract_sublines(full_text):
    """Extract VF/Vans sublines from the PDF text."""
    text = full_text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\xa0", " ")

    matches = list(SUBLINE_START_PATTERN.finditer(text))
    sublines = []

    for i, start_match in enumerate(matches):
        block_start = start_match.start()
        block_end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        block = text[block_start:block_end]

        if not re.search(r"UnitOfMeasureCode\s+(?:PC|EA)\b", block, re.I):
            continue

        qty_cost = re.search(
            r"(?P<style>VN[A-Z0-9]+)\s+.*?"
            r"(?P<qty>[\d,]+)\s+(?P<unit_cost>\d+(?:\.\d+)?)\s+.*?"
            r"UnitOfMeasureCode",
            block,
            re.IGNORECASE | re.DOTALL,
        )

        color = re.search(
            r"UnitOfMeasureCode\s+(?:PC|EA)\s+Color\s+(?P<color>.*?)\s+Size",
            block,
            re.IGNORECASE | re.DOTALL,
        )

        size = re.search(
            r"Size\s+(?P<size>.*?)\s+Dimension\s+1",
            block,
            re.IGNORECASE | re.DOTALL,
        )

        sku = re.search(
            r"SKU\s+Number\s+(?P<sku>\S+)",
            block,
            re.IGNORECASE,
        )

        upc = re.search(
            r"UPC\s+Number\s+(?P<upc>\d{12,13})",
            block,
            re.IGNORECASE,
        )

        pack = re.search(
            r"Items\s+Per\s+Outer\s+(?:Pack\s+)?(?P<pack>[\d,]+)",
            block,
            re.IGNORECASE | re.DOTALL,
        )

        if not all((qty_cost, color, size, sku, upc, pack)):
            continue

        sublines.append({
            "subline_no": start_match.group("subline_no"),
            "style": start_match.group("style"),
            "qty": int(qty_cost.group("qty").replace(",", "")),
            "unit_cost": float(qty_cost.group("unit_cost").replace(",", "")),
            "color": re.sub(r"\s+", " ", color.group("color")).strip(),
            "size": re.sub(r"\s+", " ", size.group("size")).strip(),
            "sku": sku.group("sku"),
            "upc": upc.group("upc"),
            "items_per_outer_pack": int(pack.group("pack").replace(",", "")),
        })

    return sublines


CRD_PATTERN = re.compile(r"Brand Requested CRD\s+(\d{4}-\d{2}-\d{2})")
COUNTRY_PATTERN = re.compile(r"([A-Z]{3,})\s*\na\s*\na\s*\nBrand Buyer Code")


def extract_data(pdf_file):
    """pdf_file: a file-like object (e.g. Streamlit's UploadedFile)."""
    with pdfplumber.open(pdf_file) as pdf:
        full_text = "\n".join((page.extract_text() or "") for page in pdf.pages)

    lines = []
    for m in LINE_PATTERN.finditer(full_text):
        lines.append({
            "po_line_no": m.group("po_line_no"),
            "style": m.group("style"),
            "description": re.sub(r"\s+", " ", m.group("description")).strip(),
            "order_qty": int(m.group("order_qty").replace(",", "")),
            "unit_price": float(m.group("unit_price").replace(",", "")),
            "line_cost": float(m.group("line_cost").replace(",", "")),
        })

    sublines = extract_sublines(full_text)

    crd_matches = CRD_PATTERN.findall(full_text)
    for i, l in enumerate(lines):
        l["crd"] = crd_matches[i] if i < len(crd_matches) else None

    country_match = COUNTRY_PATTERN.search(full_text)
    destination_country = country_match.group(1) if country_match else None
    for l in lines:
        l["destination_country"] = destination_country

    print(f"[PDF EXTRACTION] PO lines found: {len(lines)}")
    print(f"[PDF EXTRACTION] Sublines found: {len(sublines)}")
    return {"lines": lines, "sublines": sublines}


# ============================================================
# PART 2 - EXCEL BUILDING
# ============================================================

FONT = "Arial"
bold = Font(name=FONT, bold=True, size=10)
normal = Font(name=FONT, size=10)
title_font = Font(name=FONT, bold=True, size=13)
thin = Side(style="thin")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
yellow = PatternFill("solid", fgColor="FFFF00")
grey = PatternFill("solid", fgColor="D9D9D9")


def sheet_name_for(line):
    raw = f"{line['style']}"
    return re.sub(r'[\[\]:*?/\\]', '-', raw)[:31]


def compute_carton_rows(sublines, pack_override=None):
    if not sublines:
        return []
    pack = pack_override if pack_override is not None else (
        sublines[0].get("items_per_outer_pack") or 1
    )

    full_rows = []
    remainder_pool = []
    for s in sublines:
        qty = s["qty"]
        size = s["size"]
        full_ctn = qty // pack
        rem = qty % pack
        if full_ctn > 0:
            full_rows.append({
                "sizes": {size: pack},
                "total_ctn": full_ctn,
                "ctn_pcs": pack,
                "color": s["color"],
                "upc": s["upc"],
                "mixed": False,
            })
        if rem > 0:
            remainder_pool.append([size, rem, s["color"], s["upc"]])

    mixed_rows = []
    current = {}
    current_total = 0
    idx = 0
    while idx < len(remainder_pool):
        size, rem, color, upc = remainder_pool[idx]
        space_left = pack - current_total
        take = min(rem, space_left)
        if take > 0:
            current[size] = current.get(size, 0) + take
            current_total += take
            remainder_pool[idx][1] -= take
        if remainder_pool[idx][1] == 0:
            idx += 1
        if current_total == pack:
            mixed_rows.append({
                "sizes": dict(current), "total_ctn": 1, "ctn_pcs": pack,
                "color": None, "upc": None, "mixed": True,
            })
            current = {}
            current_total = 0
    if current:
        mixed_rows.append({
            "sizes": dict(current), "total_ctn": 1, "ctn_pcs": current_total,
            "color": None, "upc": None, "mixed": True,
        })

    return full_rows + mixed_rows


def build_sheet(ws, line, sublines):
    sizes = [s["size"] for s in sublines]
    n_sizes = len(sizes)

    ws.merge_cells("A1:H1")
    ws["A1"] = "CREATIVE COLLECTIONS LTD. U-1-A"
    ws["A1"].font = title_font
    ws["A1"].alignment = center

    ws.merge_cells("A2:H2")
    ws["A2"] = "PACKING LIST DETAILS"
    ws["A2"].font = Font(name=FONT, bold=True, size=10)
    ws["A2"].alignment = center

    # ---- Size-spec table ----
    spec_header_row = 4
    ws.cell(row=spec_header_row, column=1, value="SIZE").font = bold
    ws.cell(row=spec_header_row, column=1).border = box
    for sc_idx, size_name in enumerate(sizes):
        c = ws.cell(row=spec_header_row, column=2 + sc_idx, value=size_name)
        c.font = bold
        c.alignment = center
        c.border = box
        c.fill = grey

    spec_rows = ["N.W.", "N.N.W.", "EMPTY CTN"]
    for i, label in enumerate(spec_rows):
        rr = spec_header_row + 1 + i
        lc = ws.cell(row=rr, column=1, value=label)
        lc.font = bold
        lc.border = box
        for sc_idx, size_name in enumerate(sizes):
            c = ws.cell(row=rr, column=2 + sc_idx, value=1)
            c.alignment = center
            c.border = box

    r = spec_header_row + len(spec_rows) + 2
    ws.cell(row=r, column=1, value="BUYER").font = bold
    ws.cell(row=r, column=2, value="VANS").font = normal
    r += 1
    ws.cell(row=r, column=1, value="LOT NO").font = bold
    ws.cell(row=r, column=2, value=line["style"]).font = normal
    r += 1
    ws.cell(row=r, column=1, value="P.O NO").font = bold
    ws.cell(row=r, column=2, value=line["po_line_no"]).font = normal
    r += 1
    ws.cell(row=r, column=1, value="ORDER QTY").font = bold
    ws.cell(row=r, column=2, value=line["order_qty"]).font = normal
    ws.cell(row=r, column=3, value="PCS").font = normal
    r += 1
    ws.cell(row=r, column=1, value="PACK QTY").font = bold
    ws.cell(row=r, column=2, value=line["order_qty"]).font = normal
    ws.cell(row=r, column=3, value="PCS").font = normal
    r += 1
    excess_short_row = r
    ws.cell(row=r, column=1, value="EXCESS/SHORT QTY").font = bold
    ws.cell(row=r, column=3, value="PCS").font = normal
    r += 1
    percentage_row = r
    ws.cell(row=r, column=1, value="PERCENTAGE").font = bold
    r += 1
    ws.cell(row=r, column=1, value="CRD").font = bold
    ws.cell(row=r, column=2, value=line.get("crd")).font = normal
    r += 1
    ws.cell(row=r, column=1, value="COUNTRY").font = bold
    ws.cell(row=r, column=2, value=line.get("destination_country")).font = normal
    r += 1
    ws.cell(row=r, column=1, value="DESCRIPTION").font = bold
    ws.cell(row=r, column=2, value=line["description"]).font = normal

    header_row = r + 2
    headers = ["CASE LABEL NO.", "CARTON NO.", "COLOR", "UPC Number"] + sizes + \
              ["CTN PCS", "TOTAL CTN", "TOTAL PCS", "Grs.wt.pr ctn", "Net.wt.pr ctn",
               "total.Grs.wt", "total.net.wt"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = bold
        c.alignment = center
        c.border = box
        c.fill = grey
    n_cols = len(headers)

    data_row_start = header_row + 1
    carton_rows = compute_carton_rows(sublines)
    for idx, row_data in enumerate(carton_rows):
        rr = data_row_start + idx
        ws.cell(row=rr, column=1, value=idx + 1).border = box

        color_val = row_data["color"] if row_data["color"] is not None else "MIXED"
        color_cell = ws.cell(row=rr, column=3, value=color_val)
        color_cell.border = box
        if row_data["mixed"]:
            color_cell.fill = yellow
        ws.cell(row=rr, column=4, value=row_data["upc"]).border = box
        for sc_idx, size_name in enumerate(sizes):
            col = 5 + sc_idx
            val = row_data["sizes"].get(size_name)
            cell = ws.cell(row=rr, column=col, value=val)
            cell.border = box
            cell.alignment = center

        ctn_pcs_col = 5 + n_sizes
        first_size_letter = get_column_letter(5)
        last_size_letter = get_column_letter(4 + n_sizes)
        ctn_cell = ws.cell(
            row=rr, column=ctn_pcs_col,
            value=f"=SUM({first_size_letter}{rr}:{last_size_letter}{rr})"
        )
        ctn_cell.border = box
        ctn_cell.alignment = center
        ctn_cell.font = bold

        total_ctn_col = ctn_pcs_col + 1
        total_pcs_col = ctn_pcs_col + 2
        total_ctn_letter = get_column_letter(total_ctn_col)

        tc_cell = ws.cell(row=rr, column=total_ctn_col, value=row_data["total_ctn"])
        tc_cell.border = box
        tc_cell.alignment = center

        if idx == 0:
            carton_formula = f"={total_ctn_letter}{rr}"
        else:
            carton_formula = f"={get_column_letter(2)}{rr - 1}+{total_ctn_letter}{rr}"
        cB = ws.cell(row=rr, column=2, value=carton_formula)
        cB.border = box
        cB.alignment = center

        ctn_pcs_letter = get_column_letter(ctn_pcs_col)
        tp_cell = ws.cell(
            row=rr, column=total_pcs_col,
            value=f"={total_ctn_letter}{rr}*{ctn_pcs_letter}{rr}"
        )
        tp_cell.border = box
        tp_cell.alignment = center

        grs_wt_col = total_pcs_col + 1
        nw_row = spec_header_row + 1
        empty_row = spec_header_row + 3
        first_spec_letter = get_column_letter(2)
        last_spec_letter = get_column_letter(1 + n_sizes)
        first_main_size_letter = get_column_letter(5)
        last_main_size_letter = get_column_letter(4 + n_sizes)

        grs_formula = (
            f"=SUMPRODUCT({first_main_size_letter}{rr}:{last_main_size_letter}{rr},"
            f"{first_spec_letter}{nw_row}:{last_spec_letter}{nw_row})"
            f"+{first_spec_letter}{empty_row}"
        )
        grs_cell = ws.cell(row=rr, column=grs_wt_col, value=grs_formula)
        grs_cell.border = box
        grs_cell.alignment = center

        net_wt_col = grs_wt_col + 1
        grs_wt_letter_for_net = get_column_letter(grs_wt_col)
        net_cell = ws.cell(
            row=rr, column=net_wt_col,
            value=f"={grs_wt_letter_for_net}{rr}-{first_spec_letter}{empty_row}"
        )
        net_cell.border = box
        net_cell.alignment = center

        total_grs_col = net_wt_col + 1
        grs_wt_letter = get_column_letter(grs_wt_col)
        tg_cell = ws.cell(
            row=rr, column=total_grs_col,
            value=f"={total_ctn_letter}{rr}*{grs_wt_letter}{rr}"
        )
        tg_cell.border = box
        tg_cell.alignment = center

        total_net_col = total_grs_col + 1
        net_wt_letter = get_column_letter(net_wt_col)
        tn_cell = ws.cell(
            row=rr, column=total_net_col,
            value=f"={total_ctn_letter}{rr}*{net_wt_letter}{rr}"
        )
        tn_cell.border = box
        tn_cell.alignment = center

    total_row = data_row_start + len(carton_rows)
    ws.cell(row=total_row, column=3, value="TOTAL").font = bold
    total_ctn_letter_for_sum = get_column_letter(total_ctn_col)
    for sc_idx, size_name in enumerate(sizes):
        col = 5 + sc_idx
        col_letter = get_column_letter(col)
        cell = ws.cell(
            row=total_row, column=col,
            value=f"=SUMPRODUCT({col_letter}{data_row_start}:{col_letter}{total_row-1},"
                  f"{total_ctn_letter_for_sum}{data_row_start}:{total_ctn_letter_for_sum}{total_row-1})"
        )
        cell.font = bold
        cell.border = box

    total_ctn_col = ctn_pcs_col + 1
    total_pcs_col = ctn_pcs_col + 2
    total_ctn_letter = get_column_letter(total_ctn_col)
    total_pcs_letter = get_column_letter(total_pcs_col)

    ws.cell(row=total_row, column=total_ctn_col,
            value=f'=IF(COUNT({total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1})=0,"",'
                  f'SUM({total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1}))').font = bold

    ws.cell(row=total_row, column=total_pcs_col,
            value=f'=IF(COUNT({total_pcs_letter}{data_row_start}:{total_pcs_letter}{total_row-1})=0,"",'
                  f'SUM({total_pcs_letter}{data_row_start}:{total_pcs_letter}{total_row-1}))').font = bold

    grs_wt_col = total_pcs_col + 1
    net_wt_col = total_pcs_col + 2
    grs_wt_letter = get_column_letter(grs_wt_col)
    net_wt_letter = get_column_letter(net_wt_col)

    ws.cell(row=total_row, column=grs_wt_col,
            value=f"=SUMPRODUCT({grs_wt_letter}{data_row_start}:{grs_wt_letter}{total_row-1},"
                  f"{total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1})").font = bold
    ws.cell(row=total_row, column=net_wt_col,
            value=f"=SUMPRODUCT({net_wt_letter}{data_row_start}:{net_wt_letter}{total_row-1},"
                  f"{total_ctn_letter}{data_row_start}:{total_ctn_letter}{total_row-1})").font = bold

    total_grs_col = total_pcs_col + 3
    total_grs_letter = get_column_letter(total_grs_col)
    ws.cell(row=total_row, column=total_grs_col,
            value=f'=IF(COUNT({total_grs_letter}{data_row_start}:{total_grs_letter}{total_row-1})=0,"",'
                  f'SUM({total_grs_letter}{data_row_start}:{total_grs_letter}{total_row-1}))').font = bold

    total_net_col = total_grs_col + 1
    total_net_letter = get_column_letter(total_net_col)
    ws.cell(row=total_row, column=total_net_col,
            value=f'=IF(COUNT({total_net_letter}{data_row_start}:{total_net_letter}{total_row-1})=0,"",'
                  f'SUM({total_net_letter}{data_row_start}:{total_net_letter}{total_row-1}))').font = bold

    # ---- Summary block ----
    sum_header_row = total_row + 3
    ws.cell(row=sum_header_row, column=1, value="SIZE").font = bold
    for sc_idx, size_name in enumerate(sizes):
        ws.cell(row=sum_header_row, column=2 + sc_idx, value=size_name).font = bold
    ws.cell(row=sum_header_row, column=2 + n_sizes, value="G TOTAL").font = bold

    gw_col = 3 + n_sizes
    nw_col = 4 + n_sizes
    ws.merge_cells(start_row=sum_header_row, start_column=gw_col,
                    end_row=sum_header_row, end_column=gw_col)
    ws.cell(row=sum_header_row, column=gw_col, value="Grand Total gross weight (kg)").font = bold
    ws.cell(row=sum_header_row, column=gw_col).alignment = center
    ws.merge_cells(start_row=sum_header_row, start_column=nw_col,
                    end_row=sum_header_row, end_column=nw_col)
    ws.cell(row=sum_header_row, column=nw_col, value="Grand Total net weight (kg)").font = bold
    ws.cell(row=sum_header_row, column=nw_col).alignment = center

    rows_needed = ["CUT QTY", "ORDER QTY", "SHIP QTY", "EXS/SHT QTY", "PERCENTAGE"]
    for i, label in enumerate(rows_needed):
        rr = sum_header_row + 1 + i
        ws.cell(row=rr, column=1, value=label).font = bold
        for sc_idx, size_name in enumerate(sizes):
            col = 2 + sc_idx
            col_letter = get_column_letter(col)
            cell = ws.cell(row=rr, column=col)
            if label == "ORDER QTY":
                qty = next(s["qty"] for s in sublines if s["size"] == size_name)
                cell.value = qty
            elif label == "CUT QTY":
                cell.fill = yellow
            elif label == "SHIP QTY":
                main_size_col = 5 + sc_idx
                main_size_letter = get_column_letter(main_size_col)
                cell.value = f"={main_size_letter}{total_row}"
            elif label == "EXS/SHT QTY":
                ship_row = sum_header_row + 1 + rows_needed.index("SHIP QTY")
                order_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
                cell.value = f'=IF({col_letter}{ship_row}="","",{col_letter}{ship_row}-{col_letter}{order_row})'
            elif label == "PERCENTAGE":
                exs_row = sum_header_row + 1 + rows_needed.index("EXS/SHT QTY")
                order_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
                cell.value = f'=IF({col_letter}{exs_row}="","",IFERROR({col_letter}{exs_row}/{col_letter}{order_row},0))'
                cell.number_format = "0.00%"
            cell.border = box
            cell.alignment = center

        gt_col = 2 + n_sizes
        gt_letter = get_column_letter(gt_col)
        first_letter = get_column_letter(2)
        last_letter = get_column_letter(1 + n_sizes)
        if label in ("CUT QTY", "EXS/SHT QTY"):
            ws.cell(row=rr, column=gt_col,
                    value=f'=IF(COUNT({first_letter}{rr}:{last_letter}{rr})=0,"",SUM({first_letter}{rr}:{last_letter}{rr}))').font = bold
        elif label in ("ORDER QTY", "SHIP QTY"):
            ws.cell(row=rr, column=gt_col,
                    value=f"=SUM({first_letter}{rr}:{last_letter}{rr})").font = bold
        else:
            exs_row = sum_header_row + 1 + rows_needed.index("EXS/SHT QTY")
            order_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
            ws.cell(row=rr, column=gt_col,
                    value=f'=IF({gt_letter}{exs_row}="","",IFERROR({gt_letter}{exs_row}/{gt_letter}{order_row},0))').number_format = "0.00%"
        ws.cell(row=rr, column=gt_col).border = box

    exs_gt_row = sum_header_row + 1 + rows_needed.index("EXS/SHT QTY")
    order_gt_row = sum_header_row + 1 + rows_needed.index("ORDER QTY")
    gt_col_letter = get_column_letter(2 + n_sizes)

    ws.cell(row=excess_short_row, column=2,
            value=f'={gt_col_letter}{exs_gt_row}').font = normal
    ws.cell(row=percentage_row, column=2,
            value=f'=IF({gt_col_letter}{exs_gt_row}="","-",'
                  f'TEXT(IFERROR({gt_col_letter}{exs_gt_row}/{gt_col_letter}{order_gt_row},0),"0.00%"))'
            ).font = normal

    # ---- Weight / CBM block ----
    wt_row = sum_header_row + len(rows_needed) + 2
    total_grs_letter = get_column_letter(total_grs_col)
    total_net_letter = get_column_letter(total_net_col)
    total_ctn_letter_final = get_column_letter(total_ctn_col)

    ws.cell(row=wt_row, column=1, value="GROSS WEIGHT :").font = bold
    c = ws.cell(row=wt_row, column=2, value=f"={total_grs_letter}{total_row}")
    c.font = bold
    c.border = box
    c.alignment = center

    ws.cell(row=wt_row + 1, column=1, value="NET WEIGHT :").font = bold
    c = ws.cell(row=wt_row + 1, column=2, value=f"={total_net_letter}{total_row}")
    c.font = bold
    c.border = box
    c.alignment = center

    ws.cell(row=wt_row + 2, column=1, value="CTN MEAS :").font = bold
    ws.cell(row=wt_row + 2, column=2).fill = yellow
    ws.cell(row=wt_row + 2, column=2).border = box
    ws.cell(row=wt_row + 3, column=2).fill = yellow
    ws.cell(row=wt_row + 3, column=2).border = box

    ws.cell(row=wt_row + 4, column=1, value="CBM :").font = bold
    ws.cell(row=wt_row + 4, column=2).fill = yellow
    ws.cell(row=wt_row + 4, column=2).border = box
    cbm_result = ws.cell(
        row=wt_row + 4, column=3,
        value=f'=IF(B{wt_row + 4}="","",B{wt_row + 4}*{total_ctn_letter_final}{total_row})'
    )
    cbm_result.font = bold
    cbm_result.border = box
    cbm_result.alignment = center

    first_sum_row = sum_header_row + 1
    last_sum_row = sum_header_row + len(rows_needed)

    ws.merge_cells(start_row=first_sum_row, start_column=gw_col,
                    end_row=last_sum_row, end_column=gw_col)
    gw_cell = ws.cell(row=first_sum_row, column=gw_col,
                       value=f'={total_grs_letter}{total_row}')
    gw_cell.alignment = center
    gw_cell.font = bold
    gw_cell.border = box

    ws.merge_cells(start_row=first_sum_row, start_column=nw_col,
                    end_row=last_sum_row, end_column=nw_col)
    nw_cell = ws.cell(row=first_sum_row, column=nw_col,
                       value=f'={total_net_letter}{total_row}')
    nw_cell.alignment = center
    nw_cell.font = bold
    nw_cell.border = box

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    for i in range(n_sizes):
        ws.column_dimensions[get_column_letter(5 + i)].width = 10
    for col in range(5 + n_sizes, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions[get_column_letter(gw_col)].width = 22
    ws.column_dimensions[get_column_letter(nw_col)].width = 22

    legend_row = wt_row + 6
    ws.cell(row=legend_row, column=1, value="Legend:").font = bold
    ws.cell(row=legend_row + 1, column=1,
            value="Yellow cells = data NOT available in the PO PDF (fill in from factory cutting/packing records)"
            ).font = Font(name=FONT, size=9, italic=True)


def build_workbook(data):
    wb = Workbook()
    wb.remove(wb.active)
    used_names = set()
    for line in data["lines"]:
        sublines = [s for s in data["sublines"] if s["style"] == line["style"]]
        if not sublines:
            continue
        name = sheet_name_for(line)
        base_name, i = name, 1
        while name in used_names:
            i += 1
            name = f"{base_name[:28]}_{i}"
        used_names.add(name)
        ws = wb.create_sheet(title=name)
        build_sheet(ws, line, sublines)
    return wb


# ============================================================
# GRID CONSTANTS (Color & UPC columns are wider)
# ============================================================

ACTION_WIDTH = 42
CELL_WIDTH = 72
COLOR_WIDTH = 120          # ← wider
UPC_WIDTH = 130            # ← wider
HEADER_HEIGHT = 36
CELL_HEIGHT = 34
TOTAL_HEIGHT = 36
GRID_HEIGHT = 520


# ============================================================
# HTML PACKING-LIST PREVIEW
# ============================================================

def render_size_spec_table(style_key, full_sizes):
    st.markdown(
        """
        <style>
        .plist-wrap { font-family: Arial, sans-serif; font-size: 13px; color: #000; }
        table.plist { border-collapse: collapse; width: 100%; margin-bottom: 6px; }
        table.plist td, table.plist th { border: 1px solid #000; padding: 3px 8px; text-align: center; }
        .hdr { font-weight: bold; background: #fff; }
        .yellow { background: #FFFF66; }
        .gray { background: #D9D9D9; font-weight: bold; }
        .bold { font-weight: bold; }
        .noborder td { border: none; padding: 1px 4px; }
        .desc-cell { text-align: left; font-weight: bold; }
        .title-row td { border: none; font-weight: bold; padding: 2px 4px; }
        .ctn-total { background:#D9D9D9; color:#000; font-weight:bold; text-align:center; padding:6px 2px; border-top:2px solid #000; border-bottom:2px solid #000; }
        </style>
        """,
        unsafe_allow_html=True,
    )
    st.markdown(
        '<div class="plist-wrap">'
        '<table class="plist noborder">'
        '<tr class="title-row"><td style="text-align:center; font-size:16px;">'
        '<b>CREATIVE COLLECTIONS LTD. U-1-A</b></td></tr>'
        '<tr class="title-row"><td style="text-align:center;">PACKING LIST DETAILS.</td></tr>'
        '</table></div>',
        unsafe_allow_html=True,
    )

    header_cols = st.columns([1] + [1] * len(full_sizes))
    header_cols[0].markdown("**SIZE**")
    for col, size in zip(header_cols[1:], full_sizes):
        col.markdown(f"**{size}**")

    nw_by_size, nnw_by_size, empty_by_size = {}, {}, {}

    row_cols = st.columns([1] + [1] * len(full_sizes))
    row_cols[0].markdown("**N.W.**")
    for col, size in zip(row_cols[1:], full_sizes):
        nw_by_size[size] = col.number_input(
            f"N.W. — {size}", min_value=0.0, value=1.0, step=0.01, format="%.2f",
            key=f"nw_v2_{style_key}_{size}",
            label_visibility="collapsed",
        )

    row_cols = st.columns([1] + [1] * len(full_sizes))
    row_cols[0].markdown("**N.N.W.**")
    for col, size in zip(row_cols[1:], full_sizes):
        nnw_by_size[size] = col.number_input(
            f"N.N.W. — {size}", min_value=0.0, value=1.0, step=0.01, format="%.2f",
            key=f"nnw_v2_{style_key}_{size}",
            label_visibility="collapsed",
        )

    row_cols = st.columns([1] + [1] * len(full_sizes))
    row_cols[0].markdown("**EMPTY CTN**")
    for col, size in zip(row_cols[1:], full_sizes):
        empty_by_size[size] = col.number_input(
            f"EMPTY CTN — {size}", min_value=0.0, value=1.0, step=0.01, format="%.2f",
            key=f"empty_v2_{style_key}_{size}",
            label_visibility="collapsed",
        )

    return nw_by_size, nnw_by_size, empty_by_size


def build_edited_workbook(export_state):
    line = export_state["line"]
    full_sizes = export_state["full_sizes"]
    rows = export_state["rows"]
    n_sizes = len(full_sizes)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name_for(line)

    ws.merge_cells("A1:H1")
    ws["A1"] = "CREATIVE COLLECTIONS LTD. U-1-A"
    ws["A1"].font = title_font
    ws["A1"].alignment = center

    ws.merge_cells("A2:H2")
    ws["A2"] = "PACKING LIST DETAILS (edited)"
    ws["A2"].font = Font(name=FONT, bold=True, size=10)
    ws["A2"].alignment = center

    r = 4
    for label, value in [
        ("BUYER", "VANS"),
        ("LOT NO", line["style"]),
        ("P.O NO", line["po_line_no"]),
        ("ORDER QTY", line["order_qty"]),
        ("EXCESS/SHORT QTY", export_state["exs_total"]),
        ("CRD", line.get("crd")),
        ("COUNTRY", line.get("destination_country")),
        ("DESCRIPTION", line["description"]),
    ]:
        ws.cell(row=r, column=1, value=label).font = bold
        ws.cell(row=r, column=2, value=value).font = normal
        r += 1

    header_row = r + 1
    headers = ["Ctn No", "Color", "UPC"] + full_sizes + [
        "CTN PCS", "TOTAL CTN", "TOTAL PCS", "Grs.wt.pr ctn", "Net.wt.pr ctn",
        "total.Grs.wt", "total.net.wt"
    ]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = bold
        c.alignment = center
        c.border = box
        c.fill = grey

    data_row_start = header_row + 1
    for idx, row_data in enumerate(rows):
        rr = data_row_start + idx
        ws.cell(row=rr, column=1, value=row_data["ctn_no"]).border = box
        ws.cell(row=rr, column=2, value=row_data["color"]).border = box
        ws.cell(row=rr, column=3, value=row_data["upc"]).border = box
        for sc_idx, size_name in enumerate(full_sizes):
            cell = ws.cell(row=rr, column=4 + sc_idx, value=row_data["sizes"].get(size_name))
            cell.border = box
            cell.alignment = center
        base = 4 + n_sizes
        values = [
            row_data["ctn_pcs"], row_data["total_ctn"], row_data["total_pcs"],
            round(row_data["grs_per_ctn"], 2), round(row_data["net_per_ctn"], 2),
            round(row_data["total_grs"], 2), round(row_data["total_net"], 2),
        ]
        for off, val in enumerate(values):
            cell = ws.cell(row=rr, column=base + off, value=val)
            cell.border = box
            cell.alignment = center

    total_row = data_row_start + len(rows)
    ws.cell(row=total_row, column=2, value="TOTAL").font = bold
    for sc_idx, size_name in enumerate(full_sizes):
        cell = ws.cell(row=total_row, column=4 + sc_idx,
                        value=export_state["size_packet_totals"].get(size_name, 0))
        cell.font = bold
        cell.border = box
    base = 4 + n_sizes
    totals_row_values = [
        "", export_state["total_ctn"], export_state["total_pcs"], "", "",
        round(export_state["total_grs"], 2), round(export_state["total_net"], 2),
    ]
    for off, val in enumerate(totals_row_values):
        cell = ws.cell(row=total_row, column=base + off, value=val)
        cell.font = bold
        cell.border = box

    sum_header_row = total_row + 3
    ws.cell(row=sum_header_row, column=1, value="SIZE").font = bold
    for sc_idx, size_name in enumerate(full_sizes):
        ws.cell(row=sum_header_row, column=2 + sc_idx, value=size_name).font = bold
    ws.cell(row=sum_header_row, column=2 + n_sizes, value="G TOTAL").font = bold

    rows_needed = [
        ("CUT QTY", export_state["cut_qty"], export_state["cut_total"], yellow),
        ("ORDER QTY", export_state["order_qty_by_size"], export_state["order_total"], grey),
        ("SHIP QTY", export_state["ship_qty_by_size"], export_state["ship_total"], None),
        ("EXS/SHT QTY", export_state["exs_by_size"], export_state["exs_total"], None),
    ]
    for ridx, (label, by_size, gt, fill) in enumerate(rows_needed):
        r_row = sum_header_row + 1 + ridx
        ws.cell(row=r_row, column=1, value=label).font = bold
        for sc_idx, size_name in enumerate(full_sizes):
            cell = ws.cell(row=r_row, column=2 + sc_idx, value=by_size.get(size_name, ""))
            cell.border = box
            cell.alignment = center
            if fill is not None:
                cell.fill = fill
        cell = ws.cell(row=r_row, column=2 + n_sizes, value=gt)
        cell.font = bold
        cell.border = box

    pct_row = sum_header_row + 1 + len(rows_needed)
    ws.cell(row=pct_row, column=1, value="PERCENTAGE").font = bold
    for sc_idx, size_name in enumerate(full_sizes):
        pct = export_state["pct_by_size"].get(size_name)
        cell = ws.cell(row=pct_row, column=2 + sc_idx,
                        value="" if pct is None else pct)
        if pct is not None:
            cell.number_format = "0.00%"
        cell.border = box
        cell.alignment = center
    pct_total = export_state["pct_total"]
    cell = ws.cell(row=pct_row, column=2 + n_sizes,
                    value="" if pct_total is None else pct_total)
    if pct_total is not None:
        cell.number_format = "0.00%"
    cell.font = bold
    cell.border = box

    wt_row = pct_row + 3
    ws.cell(row=wt_row, column=1, value="GROSS WEIGHT :").font = bold
    ws.cell(row=wt_row, column=2, value=round(export_state["total_grs"], 2)).font = bold
    ws.cell(row=wt_row + 1, column=1, value="NET WEIGHT :").font = bold
    ws.cell(row=wt_row + 1, column=2, value=round(export_state["total_net"], 2)).font = bold
    ws.cell(row=wt_row + 2, column=1, value="CTN MEAS :").font = bold
    ws.cell(row=wt_row + 2, column=2,
            value=f'{export_state["ctn_meas_1"]} {export_state["ctn_meas_2"]}'.strip())
    ws.cell(row=wt_row + 3, column=1, value="CBM :").font = bold
    ws.cell(row=wt_row + 3, column=2, value=round(export_state["cbm_total"], 8))

    sign_row = wt_row + 6
    ws.cell(row=sign_row, column=1, value="INCHARGE").font = bold
    ws.cell(row=sign_row, column=3, value="FM").font = bold

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    for i in range(n_sizes):
        ws.column_dimensions[get_column_letter(4 + i)].width = 10

    return wb


def build_edited_pdf(export_state):
    line = export_state["line"]
    full_sizes = export_state["full_sizes"]
    rows = export_state["rows"]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=5 * mm, rightMargin=5 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
    )
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("CREATIVE COLLECTIONS LTD. U-1-A", styles["Title"]))
    elements.append(Paragraph("PACKING LIST DETAILS (edited)", styles["Heading3"]))
    elements.append(Spacer(1, 6))

    info_data = [
        ["BUYER", "VANS", "LOT NO", line["style"], "P.O NO", line["po_line_no"]],
        ["ORDER QTY", line["order_qty"], "EXCESS/SHORT QTY", export_state["exs_total"],
         "CRD", line.get("crd") or ""],
        ["COUNTRY", line.get("destination_country") or "", "DESCRIPTION",
         line["description"], "", ""],
    ]
    info_table = Table(info_data, colWidths=[65, 90, 90, 140, 65, 90])
    info_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.black),
        ("BACKGROUND", (0, 0), (0, -1), rl_colors.whitesmoke),
        ("BACKGROUND", (2, 0), (2, -1), rl_colors.whitesmoke),
        ("BACKGROUND", (4, 0), (4, -1), rl_colors.whitesmoke),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 8))

    headers = ["Ctn No", "Color", "UPC"] + full_sizes + [
        "CTN PCS", "TOTAL CTN", "TOTAL PCS", "Grs/ctn", "Net/ctn", "tot Grs", "tot Net"
    ]
    table_data = [headers]
    for rdata in rows:
        table_data.append([
            rdata["ctn_no"], rdata["color"], rdata["upc"],
            *[rdata["sizes"].get(sz, "") for sz in full_sizes],
            rdata["ctn_pcs"], rdata["total_ctn"], rdata["total_pcs"],
            f'{rdata["grs_per_ctn"]:.2f}', f'{rdata["net_per_ctn"]:.2f}',
            f'{rdata["total_grs"]:.2f}', f'{rdata["total_net"]:.2f}',
        ])
    total_row = (
        ["", "TOTAL", ""]
        + [export_state["size_packet_totals"].get(sz, 0) for sz in full_sizes]
        + ["", export_state["total_ctn"], export_state["total_pcs"], "", "",
           f'{export_state["total_grs"]:.2f}', f'{export_state["total_net"]:.2f}']
    )
    table_data.append(total_row)

    main_table = Table(table_data, repeatRows=1)
    main_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 6),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), rl_colors.lightgrey),
        ("BACKGROUND", (0, -1), (-1, -1), rl_colors.lightgrey),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ]))
    elements.append(main_table)
    elements.append(Spacer(1, 10))

    summary_headers = ["SIZE"] + full_sizes + ["G TOTAL"]
    summary_data = [summary_headers]

    def fmt_pct(v):
        return "" if v is None else f"{v * 100:.2f}%"

    summary_data.append(
        ["CUT QTY"] + [export_state["cut_qty"].get(sz, "") for sz in full_sizes]
        + [export_state["cut_total"]]
    )
    summary_data.append(
        ["ORDER QTY"] + [export_state["order_qty_by_size"].get(sz, "") for sz in full_sizes]
        + [export_state["order_total"]]
    )
    summary_data.append(
        ["SHIP QTY"] + [export_state["ship_qty_by_size"].get(sz, "") for sz in full_sizes]
        + [export_state["ship_total"]]
    )
    summary_data.append(
        ["EXS/SHT QTY"] + [export_state["exs_by_size"].get(sz, "") for sz in full_sizes]
        + [export_state["exs_total"]]
    )
    summary_data.append(
        ["PERCENTAGE"] + [fmt_pct(export_state["pct_by_size"].get(sz)) for sz in full_sizes]
        + [fmt_pct(export_state["pct_total"])]
    )

    summary_table = Table(summary_data, repeatRows=1)
    summary_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), rl_colors.lightgrey),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 10))

    wt_data = [
        ["GROSS WEIGHT :", f'{export_state["total_grs"]:.2f}'],
        ["NET WEIGHT :", f'{export_state["total_net"]:.2f}'],
        ["CTN MEAS :", f'{export_state["ctn_meas_1"]} {export_state["ctn_meas_2"]}'.strip()],
        ["CBM :", f'{export_state["cbm_total"]:.8f}'],
    ]
    wt_table = Table(wt_data, colWidths=[100, 150])
    wt_table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
    ]))
    elements.append(wt_table)
    elements.append(Spacer(1, 20))

    sign_table = Table([["INCHARGE", "", "FM"]], colWidths=[120, 120, 120])
    sign_table.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (0, 0), 0.75, rl_colors.black),
        ("LINEABOVE", (2, 0), (2, 0), 0.75, rl_colors.black),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    elements.append(sign_table)

    doc.build(elements)
    buf.seek(0)
    return buf


# ============================================================
# SAMPLE-STYLE WORKBOOK (centered)
# ============================================================

def build_sample_style_workbook(export_state):
    """
    Produces the exact sample-style packing list xlsx, fully center-aligned:
      Title box, size-spec table, info block, main carton table
      (with merged COLOR runs + yellow MIXED), TOTAL row, summary block
      with merged Grand Total columns, footer, signature lines.
    """
    line = export_state["line"]
    full_sizes = export_state["full_sizes"]
    rows = export_state["rows"]
    n_sizes = len(full_sizes)

    nw_by_size    = export_state.get("nw_by_size",    {sz: 0.0 for sz in full_sizes})
    nnw_by_size   = export_state.get("nnw_by_size",   {sz: 0.0 for sz in full_sizes})
    empty_by_size = export_state.get("empty_by_size", {sz: 0.0 for sz in full_sizes})

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name_for(line)

    # ---- Column indices ----
    C_CASE_A  = 1
    C_CASE_B  = 2
    C_COLOR   = 3
    C_UPC     = 4
    C_SZ0     = 5
    C_SZ_LAST = C_SZ0 + n_sizes - 1
    C_CTN_PCS = C_SZ0 + n_sizes
    C_TOT_CTN = C_CTN_PCS + 1
    C_TOT_PCS = C_CTN_PCS + 2
    C_GRS_CTN = C_CTN_PCS + 3
    C_NET_CTN = C_CTN_PCS + 4
    C_TOT_GRS = C_CTN_PCS + 5
    C_TOT_NET = C_CTN_PCS + 6
    LAST_COL  = C_TOT_NET

    BORDER   = Border(left=Side(style="thin"), right=Side(style="thin"),
                      top=Side(style="thin"), bottom=Side(style="thin"))
    CENTER   = Alignment(horizontal="center", vertical="center", wrap_text=True)
    HDR_FILL = PatternFill("solid", fgColor="D9D9D9")

    def _put(r, c, v, bold_=False, border_=True,
             fill_=None, numfmt=None, font_size=10):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=FONT, bold=bold_, size=font_size)
        cell.alignment = CENTER
        if border_:
            cell.border = BORDER
        if fill_ is not None:
            cell.fill = fill_
        if numfmt:
            cell.number_format = numfmt
        return cell

    # ============ TITLE ============
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=LAST_COL)
    _put(1, 1, "CREATIVE COLLECTIONS LTD. U-1-A",
         bold_=True, border_=False, font_size=14)

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=LAST_COL)
    _put(2, 1, "PACKING LIST DETAILS.",
         bold_=True, border_=False, font_size=11)

    # ============ SIZE SPEC TABLE ============
    spec_top = 4
    _put(spec_top, 1, "SIZE", bold_=True, fill_=HDR_FILL)
    for i, sz in enumerate(full_sizes):
        _put(spec_top, 2 + i, sz, bold_=True, fill_=HDR_FILL)

    for offset, (label, values) in enumerate([
        ("N.W.",      nw_by_size),
        ("N.N.W.",    nnw_by_size),
        ("EMPTY CTN", empty_by_size),
    ]):
        rr = spec_top + 1 + offset
        _put(rr, 1, label, bold_=True)
        for i, sz in enumerate(full_sizes):
            _put(rr, 2 + i, values.get(sz, 0), numfmt="0.00")

    # ============ INFO BLOCK ============
    info_top = spec_top + 5

    _put(info_top, 1, "BUYER", bold_=True, border_=False)
    _put(info_top, 2, "VANS", border_=False)
    _put(info_top, 4, "ORDER QTY", bold_=True, border_=False)
    _put(info_top, 5, line["order_qty"], bold_=True, border_=False)
    _put(info_top, 6, "PCS", border_=False)
    _put(info_top, 8, "CRD", bold_=True, border_=False)
    _put(info_top, 9, line.get("crd", "") or "", border_=False)

    _put(info_top + 1, 1, "LOT NO", bold_=True, border_=False)
    _put(info_top + 1, 2, line["style"], border_=False)
    _put(info_top + 1, 4, "PACK QTY", bold_=True, border_=False)
    _put(info_top + 1, 5, line["order_qty"], border_=False)
    _put(info_top + 1, 6, "PCS", border_=False)

    _put(info_top + 2, 1, "P.O NO", bold_=True, border_=False)
    _put(info_top + 2, 2, line["po_line_no"], border_=False)
    _put(info_top + 2, 4, "EXCESS/SHORT QTY", bold_=True, border_=False)
    _put(info_top + 2, 5, export_state.get("exs_total", 0), border_=False)
    _put(info_top + 2, 6, "PCS", border_=False)

    ws.merge_cells(start_row=info_top + 3, start_column=1, end_row=info_top + 3, end_column=3)
    _put(info_top + 3, 1, f"DESCRIPTION:  {line['description']}",
         bold_=True, border_=False)
    _put(info_top + 3, 4, "PERCENTAGE", bold_=True, border_=False)
    pct_total = export_state.get("pct_total")
    _put(info_top + 3, 5,
         "-" if pct_total is None else f"{pct_total*100:.2f}",
         border_=False)
    _put(info_top + 3, 6, "%", border_=False)
    _put(info_top + 3, 8, line.get("destination_country") or "",
         bold_=True, border_=False)

    # ============ MAIN TABLE ============
    table_top = info_top + 5

    ws.merge_cells(start_row=table_top, start_column=C_CASE_A,
                   end_row=table_top, end_column=C_CASE_B)
    _put(table_top, C_CASE_A, "CASE LABEL NO.", bold_=True, fill_=HDR_FILL)
    _put(table_top, C_CASE_B, None, fill_=HDR_FILL)

    ws.merge_cells(start_row=table_top, start_column=C_COLOR,
                   end_row=table_top + 1, end_column=C_COLOR)
    _put(table_top, C_COLOR, "COLOR", bold_=True, fill_=HDR_FILL)
    _put(table_top + 1, C_COLOR, None, fill_=HDR_FILL)

    ws.merge_cells(start_row=table_top, start_column=C_UPC,
                   end_row=table_top + 1, end_column=C_UPC)
    _put(table_top, C_UPC, "UPC Number", bold_=True, fill_=HDR_FILL)
    _put(table_top + 1, C_UPC, None, fill_=HDR_FILL)

    ws.merge_cells(start_row=table_top, start_column=C_SZ0,
                   end_row=table_top, end_column=C_SZ_LAST)
    _put(table_top, C_SZ0, "SIZE", bold_=True, fill_=HDR_FILL)
    for c in range(C_SZ0 + 1, C_SZ_LAST + 1):
        _put(table_top, c, None, fill_=HDR_FILL)

    for c, txt in [
        (C_CTN_PCS, "CTN PCS"),
        (C_TOT_CTN, "TOTAL CTN"),
        (C_TOT_PCS, "TOTAL PCS"),
        (C_GRS_CTN, "Grs.wt.pr ctn"),
        (C_NET_CTN, "Net.wt.pr ctn"),
        (C_TOT_GRS, "total.Grs.wt"),
        (C_TOT_NET, "total.net.wt"),
    ]:
        ws.merge_cells(start_row=table_top, start_column=c,
                       end_row=table_top + 1, end_column=c)
        _put(table_top, c, txt, bold_=True, fill_=HDR_FILL)
        _put(table_top + 1, c, None, fill_=HDR_FILL)

    for i, sz in enumerate(full_sizes):
        _put(table_top + 1, C_SZ0 + i, sz, bold_=True, fill_=HDR_FILL)

    # ---- Data rows ----
    data_top = table_top + 2
    cum_ctn = 0
    row_ranges = []
    for row in rows:
        start = cum_ctn + 1
        cum_ctn += int(row["total_ctn"])
        end = cum_ctn
        row_ranges.append((start, end))

    for idx, row in enumerate(rows):
        rr = data_top + idx
        start_ctn, end_ctn = row_ranges[idx]

        _put(rr, C_CASE_A, start_ctn)
        _put(rr, C_CASE_B, end_ctn)

        if row.get("mixed"):
            _put(rr, C_UPC, None)
        else:
            _put(rr, C_UPC, row.get("upc"))

        for s_idx, sz in enumerate(full_sizes):
            _put(rr, C_SZ0 + s_idx, row["sizes"].get(sz))

        _put(rr, C_CTN_PCS, row["ctn_pcs"], bold_=True)
        _put(rr, C_TOT_CTN, row["total_ctn"])
        _put(rr, C_TOT_PCS, row["total_pcs"])
        _put(rr, C_GRS_CTN, round(row["grs_per_ctn"], 2), numfmt="0.00")
        _put(rr, C_NET_CTN, round(row["net_per_ctn"], 2), numfmt="0.00")
        _put(rr, C_TOT_GRS, round(row["total_grs"], 2), numfmt="0.00")
        _put(rr, C_TOT_NET, round(row["total_net"], 2), numfmt="0.00")

    # ---- Merge COLOR cells in same-color runs ----
    i = 0
    while i < len(rows):
        mixed_i = bool(rows[i].get("mixed"))
        color_i = "MIXED" if mixed_i else (rows[i].get("color") or "MIXED")
        j = i
        while j + 1 < len(rows):
            mixed_j = bool(rows[j + 1].get("mixed"))
            color_j = "MIXED" if mixed_j else (rows[j + 1].get("color") or "MIXED")
            if mixed_i != mixed_j or color_i != color_j:
                break
            j += 1
        top_r = data_top + i
        bot_r = data_top + j
        if bot_r > top_r:
            ws.merge_cells(start_row=top_r, start_column=C_COLOR,
                           end_row=bot_r, end_column=C_COLOR)
        cell = ws.cell(row=top_r, column=C_COLOR)
        cell.value = color_i
        cell.font = Font(name=FONT, size=10)
        cell.alignment = CENTER
        cell.border = BORDER
        if mixed_i:
            cell.fill = yellow
        for rrr in range(top_r, bot_r + 1):
            c2 = ws.cell(row=rrr, column=C_COLOR)
            c2.border = BORDER
            if mixed_i:
                c2.fill = yellow
        i = j + 1

    # ---- TOTAL row ----
    total_row = data_top + len(rows)
    ws.merge_cells(start_row=total_row, start_column=C_CASE_A,
                   end_row=total_row, end_column=C_CASE_B)
    _put(total_row, C_CASE_A, "TOTAL", bold_=True)
    _put(total_row, C_CASE_B, None, bold_=True)

    _put(total_row, C_COLOR, None)
    _put(total_row, C_UPC, None)

    spt = export_state.get("size_packet_totals", {})
    for s_idx, sz in enumerate(full_sizes):
        _put(total_row, C_SZ0 + s_idx, spt.get(sz, 0), bold_=True)

    _put(total_row, C_CTN_PCS, None)
    _put(total_row, C_TOT_CTN, export_state.get("total_ctn", 0), bold_=True)
    _put(total_row, C_TOT_PCS, export_state.get("total_pcs", 0), bold_=True)

    sum_grs = sum(r["grs_per_ctn"] for r in rows) if rows else 0
    sum_net = sum(r["net_per_ctn"] for r in rows) if rows else 0
    _put(total_row, C_GRS_CTN, round(sum_grs, 2), bold_=True, numfmt="0.00")
    _put(total_row, C_NET_CTN, round(sum_net, 2), bold_=True, numfmt="0.00")
    _put(total_row, C_TOT_GRS, round(export_state.get("total_grs", 0), 2),
         bold_=True, numfmt="0.00")
    _put(total_row, C_TOT_NET, round(export_state.get("total_net", 0), 2),
         bold_=True, numfmt="0.00")

    # ============ SUMMARY BLOCK ============
    sum_top = total_row + 3

    _put(sum_top, 1, "SIZE", bold_=True, fill_=HDR_FILL)
    for i, sz in enumerate(full_sizes):
        _put(sum_top, 2 + i, sz, bold_=True, fill_=HDR_FILL)

    gt_col = 2 + n_sizes
    gw_col = gt_col + 1
    nw_col = gt_col + 2

    _put(sum_top, gt_col, "G Total", bold_=True, fill_=HDR_FILL)
    _put(sum_top, gw_col, "Grand Total gross weight kg", bold_=True, fill_=HDR_FILL)
    _put(sum_top, nw_col, "Grand Total Net weight kg",  bold_=True, fill_=HDR_FILL)

    def _summary_row(rr, label, by_size, total_val, fill=None, numfmt=None):
        _put(rr, 1, label, bold_=True, fill_=fill)
        for i, sz in enumerate(full_sizes):
            _put(rr, 2 + i, by_size.get(sz, ""), fill_=fill, numfmt=numfmt)
        _put(rr, gt_col, total_val, bold_=True, fill_=fill, numfmt=numfmt)

    cut_qty = export_state.get("cut_qty", {})
    order_by = export_state.get("order_qty_by_size", {})
    ship_by  = export_state.get("ship_qty_by_size", {})
    exs_by   = export_state.get("exs_by_size", {})
    pct_by   = export_state.get("pct_by_size", {})

    _summary_row(sum_top + 1, "CUT QTY", cut_qty, export_state.get("cut_total", 0), fill=yellow)
    _summary_row(sum_top + 2, "ORDER QTY", order_by, export_state.get("order_total", 0), fill=grey)
    _summary_row(sum_top + 3, "SHIP QTY", ship_by, export_state.get("ship_total", 0))
    _summary_row(sum_top + 4, "EXS/SHT QTY", exs_by, export_state.get("exs_total", 0))

    pct_row = sum_top + 5
    _put(pct_row, 1, "PERCENTAGE", bold_=True)
    for i, sz in enumerate(full_sizes):
        v = pct_by.get(sz)
        _put(pct_row, 2 + i, "" if v is None else float(v), numfmt="0.00%")
    v_total = export_state.get("pct_total")
    _put(pct_row, gt_col, "" if v_total is None else float(v_total),
         bold_=True, numfmt="0.00%")

    ws.merge_cells(start_row=sum_top + 1, start_column=gw_col,
                   end_row=pct_row, end_column=gw_col)
    _put(sum_top + 1, gw_col, round(export_state.get("total_grs", 0), 2),
         bold_=True, numfmt="0.00")
    for rrr in range(sum_top + 1, pct_row + 1):
        ws.cell(row=rrr, column=gw_col).border = BORDER

    ws.merge_cells(start_row=sum_top + 1, start_column=nw_col,
                   end_row=pct_row, end_column=nw_col)
    _put(sum_top + 1, nw_col, round(export_state.get("total_net", 0), 2),
         bold_=True, numfmt="0.00")
    for rrr in range(sum_top + 1, pct_row + 1):
        ws.cell(row=rrr, column=nw_col).border = BORDER

    # ============ FOOTER ============
    foot_top = pct_row + 3
    _put(foot_top,     1, "GROSS WEIGHT :", bold_=True, border_=False)
    _put(foot_top,     2, round(export_state.get("total_grs", 0), 2),
         bold_=True, border_=False)
    _put(foot_top + 1, 1, "NET WEIGHT :",   bold_=True, border_=False)
    _put(foot_top + 1, 2, round(export_state.get("total_net", 0), 2),
         bold_=True, border_=False)
    _put(foot_top + 2, 1, "CTN MEAS :",     bold_=True, border_=False)
    _put(foot_top + 2, 2, export_state.get("ctn_meas_1", ""), border_=False)
    _put(foot_top + 3, 2, export_state.get("ctn_meas_2", ""), border_=False)
    _put(foot_top + 4, 1, "CBM :",          bold_=True, border_=False)
    _put(foot_top + 4, 2, round(export_state.get("cbm_total", 0), 8), border_=False)

    # ---- Signature ----
    sign_top = foot_top + 7
    _put(sign_top, 1, "INCHARGE", bold_=True, border_=False)
    _put(sign_top, 3, "FM",       bold_=True, border_=False)
    ws.cell(row=sign_top, column=1).border = Border(top=Side(style="thin"))
    ws.cell(row=sign_top, column=3).border = Border(top=Side(style="thin"))

    # ============ COLUMN WIDTHS ============
    ws.column_dimensions[get_column_letter(C_CASE_A)].width = 8
    ws.column_dimensions[get_column_letter(C_CASE_B)].width = 8
    ws.column_dimensions[get_column_letter(C_COLOR)].width = 16
    ws.column_dimensions[get_column_letter(C_UPC)].width = 14
    for i in range(n_sizes):
        ws.column_dimensions[get_column_letter(C_SZ0 + i)].width = 9
    ws.column_dimensions[get_column_letter(C_CTN_PCS)].width = 9
    ws.column_dimensions[get_column_letter(C_TOT_CTN)].width = 10
    ws.column_dimensions[get_column_letter(C_TOT_PCS)].width = 10
    ws.column_dimensions[get_column_letter(C_GRS_CTN)].width = 12
    ws.column_dimensions[get_column_letter(C_NET_CTN)].width = 12
    ws.column_dimensions[get_column_letter(C_TOT_GRS)].width = 12
    ws.column_dimensions[get_column_letter(C_TOT_NET)].width = 12

    # --- ADDED: freeze header rows + hide gridlines ---
    ws.freeze_panes = ws.cell(row=table_top + 2, column=C_SZ0).coordinate
    ws.sheet_view.showGridLines = False
    # --- END ADDED ---

    return wb


# ============================================================
# UI THEME (purely cosmetic, safe to remove)
# ============================================================

def inject_ui_theme():
    """Purely cosmetic. Safe to remove — the app works identically without it."""
    st.markdown(
        """
        <style>
        /* ===== Font & base ===== */
        html, body, [class*="css"], .stApp {
            font-family: 'Inter', -apple-system, BlinkMacSystemFont,
                         'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
            -webkit-font-smoothing: antialiased;
        }

        /* ===== Background ===== */
        [data-testid="stAppViewContainer"] {
            background: linear-gradient(180deg, #f8fafc 0%, #f1f5f9 100%);
        }
        [data-testid="stHeader"] {
            background: transparent;
        }

        /* ===== Headings ===== */
        h1 {
            font-weight: 800 !important;
            letter-spacing: -0.02em !important;
            color: #0f172a !important;
            padding-bottom: 4px;
        }
        h2, h3, h4, h5 {
            font-weight: 700 !important;
            letter-spacing: -0.01em !important;
            color: #1e293b !important;
        }

        /* ===== Captions ===== */
        [data-testid="stCaptionContainer"] p {
            color: #64748b !important;
            font-size: 0.85rem !important;
        }

        /* ===== Divider ===== */
        hr {
            border-color: #e2e8f0 !important;
            margin: 1.25rem 0 !important;
        }

        /* ===== Buttons (NOT the grid + / x — those override via !important) ===== */
        .stButton > button,
        .stDownloadButton > button {
            border-radius: 10px;
            border: 1px solid #e2e8f0;
            box-shadow: 0 1px 2px rgba(15, 23, 42, 0.04);
            transition: transform 0.15s ease, box-shadow 0.15s ease,
                        border-color 0.15s ease, background 0.15s ease;
            font-weight: 500;
        }
        .stButton > button:hover,
        .stDownloadButton > button:hover {
            transform: translateY(-1px);
            box-shadow: 0 4px 12px rgba(15, 23, 42, 0.08);
            border-color: #cbd5e1;
        }
        .stButton > button:active,
        .stDownloadButton > button:active {
            transform: translateY(0);
        }

        /* Primary action button (the "type=primary" ones) */
        .stButton > button[kind="primary"] {
            background: linear-gradient(135deg, #1e40af 0%, #2563eb 100%);
            border: none;
            box-shadow: 0 4px 14px rgba(37, 99, 235, 0.25);
        }
        .stButton > button[kind="primary"]:hover {
            box-shadow: 0 6px 20px rgba(37, 99, 235, 0.35);
        }

        /* Download buttons */
        .stDownloadButton > button {
            background: #ffffff;
        }
        .stDownloadButton > button:hover {
            background: #f8fafc;
        }

        /* ===== Metrics ===== */
        [data-testid="stMetric"] {
            background: #ffffff;
            border: 1px solid #e2e8f0;
            border-radius: 12px;
            padding: 14px 18px;
            box-shadow: 0 1px 3px rgba(15, 23, 42, 0.04);
            transition: box-shadow 0.15s ease;
        }
        [data-testid="stMetric"]:hover {
            box-shadow: 0 4px 12px rgba(15, 23, 42, 0.08);
        }
        [data-testid="stMetricLabel"] {
            color: #64748b !important;
            font-weight: 500 !important;
            font-size: 0.8rem !important;
            text-transform: uppercase;
            letter-spacing: 0.04em;
        }
        [data-testid="stMetricValue"] {
            color: #0f172a !important;
            font-weight: 700 !important;
        }

        /* ===== Expanders ===== */
        [data-testid="stExpander"] {
            border: 1px solid #e2e8f0 !important;
            border-radius: 12px !important;
            background: #ffffff !important;
            box-shadow: 0 1px 3px rgba(15, 23, 42, 0.03);
        }
        [data-testid="stExpander"] summary {
            font-weight: 600 !important;
            color: #1e293b !important;
        }

        /* ===== File uploader ===== */
        [data-testid="stFileUploader"] {
            background: #ffffff;
            border: 1px solid #e2e8f0;
            border-radius: 12px;
            padding: 12px;
            box-shadow: 0 1px 3px rgba(15, 23, 42, 0.04);
        }
        [data-testid="stFileUploader"] section {
            border: 2px dashed #cbd5e1 !important;
            border-radius: 10px !important;
            background: #f8fafc !important;
            transition: border-color 0.15s ease, background 0.15s ease;
        }
        [data-testid="stFileUploader"] section:hover {
            border-color: #2563eb !important;
            background: #eff6ff !important;
        }

        /* ===== Dataframe ===== */
        [data-testid="stDataFrame"] {
            border: 1px solid #e2e8f0;
            border-radius: 12px;
            overflow: hidden;
            box-shadow: 0 1px 3px rgba(15, 23, 42, 0.04);
        }

        /* ===== Alerts ===== */
        [data-testid="stAlert"] {
            border-radius: 10px !important;
            border-left-width: 4px !important;
            box-shadow: 0 1px 3px rgba(15, 23, 42, 0.04);
        }

        /* ===== Number / text inputs (outside the grid) ===== */
        .stNumberInput input,
        .stTextInput input,
        .stSelectbox div[data-baseweb="select"] > div {
            border-radius: 8px !important;
            border-color: #e2e8f0 !important;
        }
        .stNumberInput input:focus,
        .stTextInput input:focus {
            border-color: #2563eb !important;
            box-shadow: 0 0 0 3px rgba(37, 99, 235, 0.12) !important;
        }

        /* ===== Section headers (##### markdown) ===== */
        [data-testid="stMarkdownContainer"] h5 {
            margin-top: 1.25rem !important;
            margin-bottom: 0.5rem !important;
            padding-bottom: 6px;
            border-bottom: 2px solid #e2e8f0;
        }

        /* ===== Tabs (if you add any later) ===== */
        .stTabs [data-baseweb="tab-list"] {
            gap: 8px;
        }
        .stTabs [data-baseweb="tab"] {
            border-radius: 8px;
            padding: 6px 14px;
            font-weight: 500;
        }
        .stTabs [aria-selected="true"] {
            background: #eff6ff;
            color: #1e40af !important;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def render_packing_list_preview(line, sublines):
    """Interactive packing-list preview with wider Color & UPC columns."""
    style_key = line["style"]
    full_sizes = [s["size"] for s in sublines]
    n_sizes = len(full_sizes)

    default_color = next((s["color"] for s in sublines if s.get("color")), "MIXED")

    _pdf_pack = int(sublines[0].get("items_per_outer_pack") or 1) if sublines else 1
    _saved_pack_max = st.session_state.get(f"pack_max_v4_{style_key}")
    effective_pack = int(_saved_pack_max) if _saved_pack_max is not None else _pdf_pack

    rows_key = f"carton_rows_v8_{style_key}"
    pack_sig_key = f"pack_sig_v8_{style_key}"
    _pack_changed = st.session_state.get(pack_sig_key) != effective_pack
    if rows_key not in st.session_state or _pack_changed:
        base = compute_carton_rows(sublines, pack_override=effective_pack)
        st.session_state[rows_key] = [
            {
                "id": uuid.uuid4().hex,
                "color": (r["color"] if r["color"] is not None else default_color),
                "upc": r["upc"] if not r["mixed"] else "MIXED",
                "sizes": {sz: int(r["sizes"].get(sz, 0) or 0) for sz in full_sizes},
                "total_ctn": int(r["total_ctn"]),
                "mixed": bool(r["mixed"]),
            }
            for r in base
        ]
        st.session_state[pack_sig_key] = effective_pack

    carton_rows = st.session_state[rows_key]

    nw_by_size, nnw_by_size, empty_by_size = render_size_spec_table(style_key, full_sizes)
    empty_ctn_wt = float(empty_by_size[full_sizes[0]]) if full_sizes else 0.0

    # ---- CSS (Color = 5th column, UPC = 6th column are wider) ----
    st.markdown(
        f"""
        <style>
        .st-key-carton_grid {{
            width: 100% !important;
            max-width: 100% !important;

            overflow-x: auto !important;
            box-sizing: border-box !important;
            border: 1px solid #cbd5e1;
            border-radius: 6px;
            background: white;
        }}

        .st-key-carton_grid [data-testid="stHorizontalBlock"] {{
            display: flex !important;
            flex-wrap: nowrap !important;
            gap: 0 !important;
            width: max-content !important;
            min-width: max-content !important;
            max-width: none !important;
            flex-shrink: 0 !important;
        }}

        .st-key-carton_grid [data-testid="stColumn"] {{
            flex: 0 0 {CELL_WIDTH}px !important;
            width: {CELL_WIDTH}px !important;
            min-width: {CELL_WIDTH}px !important;
            max-width: {CELL_WIDTH}px !important;
            box-sizing: border-box !important;
            padding-left: 0 !important;
            padding-right: 0 !important;
            margin: 0 !important;
            flex-shrink: 0 !important;
        }}

        .st-key-carton_grid [data-testid="stHorizontalBlock"] [data-testid="stColumn"]:nth-child(1),
        .st-key-carton_grid [data-testid="stHorizontalBlock"] [data-testid="stColumn"]:nth-child(2) {{
            flex: 0 0 {ACTION_WIDTH}px !important;
            width: {ACTION_WIDTH}px !important;
            min-width: {ACTION_WIDTH}px !important;
            max-width: {ACTION_WIDTH}px !important;
        }}

        .st-key-carton_grid [data-testid="stHorizontalBlock"] [data-testid="stColumn"]:nth-child(5) {{
            flex: 0 0 {COLOR_WIDTH}px !important;
            width: {COLOR_WIDTH}px !important;
            min-width: {COLOR_WIDTH}px !important;
            max-width: {COLOR_WIDTH}px !important;
        }}

        .st-key-carton_grid [data-testid="stHorizontalBlock"] [data-testid="stColumn"]:nth-child(6) {{
            flex: 0 0 {UPC_WIDTH}px !important;
            width: {UPC_WIDTH}px !important;
            min-width: {UPC_WIDTH}px !important;
            max-width: {UPC_WIDTH}px !important;
        }}

        .grid-header {{
            height: 36px !important;
            min-height: 38px !important;
            max-height: 38px !important;
            box-sizing: border-box !important;
            display: flex;
            align-items: center;
            justify-content: center;
            background: #1f2937;
            color: white;
            border: 1px solid #4b5563 !important;
            font-size: 12px;
            font-weight: 600;
            white-space: nowrap;
            overflow: hidden;
        }}

        .grid-header-action {{
            width: {ACTION_WIDTH}px !important;
            min-width: {ACTION_WIDTH}px !important;
            max-width: {ACTION_WIDTH}px !important;
        }}

        .grid-header-color {{
            width: {COLOR_WIDTH}px !important;
            min-width: {COLOR_WIDTH}px !important;
            max-width: {COLOR_WIDTH}px !important;
        }}

        .grid-header-upc {{
            width: {UPC_WIDTH}px !important;
            min-width: {UPC_WIDTH}px !important;
            max-width: {UPC_WIDTH}px !important;
        }}

        .st-key-carton_grid [data-testid="stButton"] {{
            width: {ACTION_WIDTH}px !important;
            min-width: {ACTION_WIDTH}px !important;
            max-width: {ACTION_WIDTH}px !important;
            margin: 0 !important;
            padding: 0 !important;
        }}

        .st-key-carton_grid [data-testid="stButton"] button {{
            width: {ACTION_WIDTH}px !important;
            min-width: {ACTION_WIDTH}px !important;
            max-width: {ACTION_WIDTH}px !important;
            height: {CELL_HEIGHT}px !important;
            min-height: {CELL_HEIGHT}px !important;
            max-height: {CELL_HEIGHT}px !important;
            padding: 0 !important;
            margin: 0 !important;
            border-radius: 0 !important;
            font-size: 16px !important;
        }}

        .st-key-carton_grid [data-testid="stNumberInput"] {{
            width: {CELL_WIDTH}px !important;
            min-width: {CELL_WIDTH}px !important;
            max-width: {CELL_WIDTH}px !important;
            margin: 0 !important;
            padding: 0 !important;
        }}

        .st-key-carton_grid [data-testid="stNumberInput"] input {{
            width: {CELL_WIDTH}px !important;
            min-width: {CELL_WIDTH}px !important;
            max-width: {CELL_WIDTH}px !important;
            height: {CELL_HEIGHT}px !important;
            min-height: {CELL_HEIGHT}px !important;
            max-height: {CELL_HEIGHT}px !important;
            box-sizing: border-box !important;
            border-radius: 0 !important;
            margin: 0 !important;
            padding: 0 4px !important;
            font-size: 12px !important;
            text-align: center !important;
        }}

        .st-key-carton_grid [data-testid="stNumberInput"] label {{
            display: none !important;
        }}

        .data-cell {{
            height: 34px !important;
            min-height: 34px !important;
            max-height: 34px !important;
            box-sizing: border-box !important;
            display: flex;
            align-items: center;
            justify-content: center;
            border: 1px solid #d1d5db !important;
            background: #fff;
            font-size: 12px;
            white-space: nowrap;
            overflow: hidden;
        }}

        .data-cell.mixed {{
            background: #FFFF66;
        }}

        .total-cell {{
            height: 36px !important;
            min-height: 36px !important;
            max-height: 36px !important;
            box-sizing: border-box !important;
            display: flex;
            align-items: center;
            justify-content: center;
            border: 1px solid #9ca3af !important;
            background: #e5e7eb;
            color: #111827;
            font-weight: 700;
            font-size: 12px;
            white-space: nowrap;
            overflow: hidden;
        }}

        .total-action-cell {{
            width: {ACTION_WIDTH}px !important;
            min-width: {ACTION_WIDTH}px !important;
            max-width: {ACTION_WIDTH}px !important;
            height: {TOTAL_HEIGHT}px !important;
            min-height: {TOTAL_HEIGHT}px !important;
            max-height: {TOTAL_HEIGHT}px !important;
            box-sizing: border-box !important;
            display: flex;
            align-items: center;
            justify-content: center;
            background: #e5e7eb;
            border-right: 1px solid #9ca3af;
            border-bottom: 2px solid #6b7280;
            font-weight: 700;
        }}
        </style>
        """,
        unsafe_allow_html=True,
    )

    st.caption(
        "Main carton table is fully editable. Color & UPC columns are wider. "
        "Change size quantities or TOTAL CTN directly in the cells."
    )

    # ---- Handle insert / delete ----
    insert_key = f"insert_after_v8_{style_key}"
    delete_key = f"delete_row_v8_{style_key}"

    if insert_key in st.session_state and st.session_state[insert_key] is not None:
        idx = st.session_state[insert_key]
        new_row = {
            "id": uuid.uuid4().hex,
            "color": default_color,
            "upc": "MIXED",
            "sizes": {sz: 0 for sz in full_sizes},
            "total_ctn": 1,
            "mixed": True,
        }
        carton_rows.insert(idx + 1, new_row)
        st.session_state[rows_key] = carton_rows
        st.session_state[insert_key] = None
        st.rerun()

    if delete_key in st.session_state and st.session_state[delete_key] is not None:
        idx = st.session_state[delete_key]
        if len(carton_rows) > 1 and 0 <= idx < len(carton_rows):
            carton_rows.pop(idx)
            st.session_state[rows_key] = carton_rows
        st.session_state[delete_key] = None
        st.rerun()

    n_rows = len(carton_rows)

    col_spec = (
        [ACTION_WIDTH, ACTION_WIDTH] +
        [CELL_WIDTH, CELL_WIDTH] +
        [COLOR_WIDTH, UPC_WIDTH] +
        [CELL_WIDTH] * n_sizes +
        [CELL_WIDTH] * 7
    )

    # ---- Order quantity by size table ----
    o_hdr = "".join(f"<th>{escape(str(s['size']))}</th>" for s in sublines)
    o_color = "".join(f"<td>{escape(str(s['color']))}</td>" for s in sublines)
    o_upc = "".join(f"<td>{escape(str(s['upc']))}</td>" for s in sublines)
    o_qty = "".join(f"<td><b>{s['qty']}</b></td>" for s in sublines)

    st.markdown("##### Order quantity by size")
    st.markdown(
        '<style>'
        '.order-scroll{overflow-x:auto;-webkit-overflow-scrolling:touch;'
        'border:1px solid #ccc;border-radius:4px;margin:8px 0 12px 0;max-width:100%;background:#fff;}'
        'table.order-tbl{border-collapse:collapse;width:max-content;min-width:100%;'
        'font-family:Arial,sans-serif;font-size:12px;color:#000;white-space:nowrap;}'
        'table.order-tbl th,table.order-tbl td{border:1px solid #333;padding:4px 8px;'
        'text-align:center;vertical-align:middle;background:#fff;}'
        'table.order-tbl th{background:#D9D9D9;font-weight:bold;}'
        'table.order-tbl .rowlbl{position:sticky;left:0;z-index:1;background:#D9D9D9;'
        'font-weight:bold;text-align:left;}'
        '</style>'
        '<div class="order-scroll"><table class="order-tbl">'
        f'<tr><th class="rowlbl">SIZE</th>{o_hdr}</tr>'
        f'<tr><td class="rowlbl">COLOR</td>{o_color}</tr>'
        f'<tr><td class="rowlbl">UPC</td>{o_upc}</tr>'
        f'<tr><td class="rowlbl">ORDER QTY</td>{o_qty}</tr>'
        '</table></div>',
        unsafe_allow_html=True,
    )

    # ---- Packing qty MIN / MAX table ----
    pdf_pack = int(sublines[0].get("items_per_outer_pack") or 1) if sublines else 1

    st.markdown("##### Packing qty (per carton)")
    pk_w = [1.2, 1, 1, 2.2]
    pk_hdr = st.columns(pk_w)
    pk_hdr[0].markdown("*PACKING QTY*")
    pk_hdr[1].markdown("*MIN*")
    pk_hdr[2].markdown("*MAX*")
    pk_hdr[3].markdown("*PDF (Items Per Outer Pack)*")
    pk_row = st.columns(pk_w)
    pk_row[0].markdown("Input")
    pack_min = pk_row[1].number_input(
        "Packing qty MIN", min_value=1, value=None, step=1, placeholder="empty",
        key=f"pack_min_v4_{style_key}", label_visibility="collapsed",
    )
    pack_max = pk_row[2].number_input(
        "Packing qty MAX", min_value=1, value=None, step=1, placeholder="empty",
        key=f"pack_max_v4_{style_key}", label_visibility="collapsed",
    )
    pk_row[3].markdown(f"*{pdf_pack}* (used when MAX is empty)")

    # --- ADDED: Reset button ---
    _rh1, _rh2 = st.columns([1, 6])
    with _rh1:
        if st.button("🔄 Reset", key=f"reset_rows_v8_{style_key}",
                     help="Restore carton rows to PDF defaults (wipes manual edits)"):
            st.session_state.pop(rows_key, None)
            st.session_state.pop(pack_sig_key, None)
            for _k in list(st.session_state.keys()):
                if (_k.startswith(f"sz_v8_{style_key}_")
                        or _k.startswith(f"tctn_v8_{style_key}_")
                        or _k.startswith(f"add_v8_{style_key}_")
                        or _k.startswith(f"del_v8_{style_key}_")):
                    st.session_state.pop(_k, None)
            st.rerun()
    # --- END ADDED ---

    st.markdown("##### Main carton table (editable)")

    edited_sizes = []
    edited_ctn_pcs = []
    edited_total_ctn = []
    grs_wts, net_wts = [], []
    row_total_pcs = []
    ctn_ranges = []
    row_colors, row_upcs = [], []
    cum = 0

    with st.container(border=True, key="carton_grid"):

        header_cols = st.columns(col_spec, gap="small")

        with header_cols[0]:
            st.markdown('<div class="grid-header grid-header-action">+</div>', unsafe_allow_html=True)
        with header_cols[1]:
            st.markdown('<div class="grid-header grid-header-action">−</div>', unsafe_allow_html=True)
        with header_cols[2]:
            st.markdown('<div class="grid-header">#</div>', unsafe_allow_html=True)
        with header_cols[3]:
            st.markdown('<div class="grid-header">Ctn No</div>', unsafe_allow_html=True)
        with header_cols[4]:
            st.markdown('<div class="grid-header grid-header-color">Color</div>', unsafe_allow_html=True)
        with header_cols[5]:
            st.markdown('<div class="grid-header grid-header-upc">UPC</div>', unsafe_allow_html=True)

        for s_idx, sz in enumerate(full_sizes):
            with header_cols[6 + s_idx]:
                st.markdown(f'<div class="grid-header">{sz}</div>', unsafe_allow_html=True)

        extra_headers = ["CTN PCS", "TOTAL CTN", "TOTAL PCS", "Grs/ctn", "Net/ctn", "tot Grs", "tot Net"]
        for e_idx, h in enumerate(extra_headers):
            with header_cols[6 + n_sizes + e_idx]:
                st.markdown(f'<div class="grid-header">{h}</div>', unsafe_allow_html=True)

        for row_index, row in enumerate(carton_rows):
            row_id = row["id"]
            row_cols = st.columns(col_spec, gap="small")

            with row_cols[0]:
                if st.button("＋", key=f"add_v8_{style_key}_{row_id}",
                             help="Add row below"):
                    st.session_state[insert_key] = row_index
                    st.rerun()

            with row_cols[1]:
                if st.button("❌", key=f"del_v8_{style_key}_{row_id}",
                             help="Delete this row",
                             disabled=n_rows <= 1):
                    st.session_state[delete_key] = row_index
                    st.rerun()

            with row_cols[2]:
                st.markdown(f'<div class="data-cell">{row_index + 1}</div>', unsafe_allow_html=True)

            current_sizes = {}
            for s_idx, sz in enumerate(full_sizes):
                with row_cols[6 + s_idx]:
                    sz_key = f"sz_v8_{style_key}_{row_id}_{sz}"
                    current_val = int(row["sizes"].get(sz, 0) or 0)
                    new_val = st.number_input(
                        f"sz_{row_id}_{sz}",
                        min_value=0,
                        value=current_val,
                        step=1,
                        key=sz_key,
                        label_visibility="collapsed",
                    )
                    current_sizes[sz] = int(new_val)
                    carton_rows[row_index]["sizes"][sz] = int(new_val)

            with row_cols[7 + n_sizes]:
                tctn_key = f"tctn_v8_{style_key}_{row_id}"
                current_tctn = int(row.get("total_ctn") or 0)
                new_tctn = st.number_input(
                    f"tctn_{row_id}",
                    min_value=0,
                    value=current_tctn,
                    step=1,
                    key=tctn_key,
                    label_visibility="collapsed",
                )
                carton_rows[row_index]["total_ctn"] = int(new_tctn)
                total_ctn_val = int(new_tctn)

            size_sum = sum(current_sizes.values())
            ctn_start = cum + 1
            cum += total_ctn_val
            ctn_end = cum
            ctn_range_str = f"{ctn_start}-{ctn_end}" if total_ctn_val > 0 else "-"

            piece_wt = sum(
                float(current_sizes.get(sz, 0) or 0) * float(nw_by_size.get(sz, 0.0) or 0.0)
                for sz in full_sizes
            )
            grs = piece_wt + empty_ctn_wt
            net = piece_wt
            tot_pcs = total_ctn_val * size_sum

            edited_sizes.append(current_sizes)
            edited_ctn_pcs.append(size_sum)
            edited_total_ctn.append(total_ctn_val)
            grs_wts.append(grs)
            net_wts.append(net)
            row_total_pcs.append(tot_pcs)
            ctn_ranges.append(ctn_range_str)
            row_colors.append(row.get("color") or default_color)
            row_upcs.append(row.get("upc") or "MIXED")

            with row_cols[3]:
                st.markdown(f'<div class="data-cell">{ctn_range_str}</div>', unsafe_allow_html=True)

            mixed_cls = " mixed" if row.get("mixed") else ""
            with row_cols[4]:
                st.markdown(
                    f'<div class="data-cell{mixed_cls}">{row.get("color") or default_color}</div>',
                    unsafe_allow_html=True
                )

            with row_cols[5]:
                st.markdown(
                    f'<div class="data-cell">{row.get("upc") or "MIXED"}</div>',
                    unsafe_allow_html=True
                )

            with row_cols[6 + n_sizes]:
                st.markdown(f'<div class="data-cell"><b>{size_sum}</b></div>', unsafe_allow_html=True)

            with row_cols[8 + n_sizes]:
                st.markdown(f'<div class="data-cell"><b>{tot_pcs}</b></div>', unsafe_allow_html=True)

            with row_cols[9 + n_sizes]:
                st.markdown(f'<div class="data-cell">{grs:.2f}</div>', unsafe_allow_html=True)

            with row_cols[10 + n_sizes]:
                st.markdown(f'<div class="data-cell">{net:.2f}</div>', unsafe_allow_html=True)

            with row_cols[11 + n_sizes]:
                st.markdown(f'<div class="data-cell">{total_ctn_val * grs:.2f}</div>', unsafe_allow_html=True)

            with row_cols[12 + n_sizes]:
                st.markdown(f'<div class="data-cell">{total_ctn_val * net:.2f}</div>', unsafe_allow_html=True)

        total_ctn = sum(edited_total_ctn) if edited_total_ctn else 0
        total_pcs = sum(row_total_pcs) if row_total_pcs else 0
        total_grs = sum(tc * g for tc, g in zip(edited_total_ctn, grs_wts)) if edited_total_ctn else 0.0
        total_net = sum(tc * n for tc, n in zip(edited_total_ctn, net_wts)) if edited_total_ctn else 0.0

        size_packet_totals = {
            sz: sum(int(edited_sizes[i].get(sz, 0) or 0) * edited_total_ctn[i] for i in range(n_rows))
            for sz in full_sizes
        } if n_rows else {sz: 0 for sz in full_sizes}

        total_cols = st.columns(col_spec, gap="small")

        with total_cols[0]:
            st.markdown('<div class="total-action-cell"></div>', unsafe_allow_html=True)
        with total_cols[1]:
            st.markdown('<div class="total-action-cell">TOTAL</div>', unsafe_allow_html=True)
        with total_cols[2]:
            st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)
        with total_cols[3]:
            st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)
        with total_cols[4]:
            st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)
        with total_cols[5]:
            st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)

        for s_idx, sz in enumerate(full_sizes):
            with total_cols[6 + s_idx]:
                st.markdown(f'<div class="total-cell">{size_packet_totals.get(sz, 0)}</div>', unsafe_allow_html=True)

        with total_cols[6 + n_sizes]:
            st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)
        with total_cols[7 + n_sizes]:
            st.markdown(f'<div class="total-cell">{total_ctn}</div>', unsafe_allow_html=True)
        with total_cols[8 + n_sizes]:
            st.markdown(f'<div class="total-cell">{total_pcs}</div>', unsafe_allow_html=True)
        with total_cols[9 + n_sizes]:
            st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)
        with total_cols[10 + n_sizes]:
            st.markdown('<div class="total-cell"></div>', unsafe_allow_html=True)
        with total_cols[11 + n_sizes]:
            st.markdown(f'<div class="total-cell">{total_grs:.2f}</div>', unsafe_allow_html=True)
        with total_cols[12 + n_sizes]:
            st.markdown(f'<div class="total-cell">{total_net:.2f}</div>', unsafe_allow_html=True)

    st.session_state[rows_key] = carton_rows

    st.markdown("**Live totals**")
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Total Cartons", total_ctn)
    m2.metric("Total Pieces", total_pcs)
    m3.metric("Total Gross Wt", f"{total_grs:.2f}")
    m4.metric("Total Net Wt", f"{total_net:.2f}")

    # --- ADDED: Excess / Short warning banner ---
    _order_total_now = sum(s["qty"] for s in sublines)
    _ship_total_now = sum(
        int(edited_sizes[_i].get(_sz, 0) or 0) * edited_total_ctn[_i]
        for _i in range(n_rows) for _sz in full_sizes
    )
    _exs_now = _ship_total_now - _order_total_now
    if _exs_now > 0:
        st.success(f"✅ Over-ship: **+{_exs_now}** pieces")
    elif _exs_now < 0:
        st.error(f"⚠️ Short-ship: **{_exs_now}** pieces")
    else:
        st.info("Ship qty matches order qty exactly.")
    # --- END ADDED ---

    with st.expander(f"Edit cut qty / CTN MEAS / CBM for {style_key}", expanded=False):
        cut_qty = {}
        cut_cols = st.columns(len(full_sizes)) if full_sizes else []
        for col, size in zip(cut_cols, full_sizes):
            with col:
                cut_qty[size] = st.number_input(
                    f"Cut qty — {size}", value=0, key=f"cut_v8_{style_key}_{size}"
                )

        ctn_meas_1 = st.text_input("Ctn meas — line 1", "", key=f"meas1_v8_{style_key}")

        st.markdown("**Ctn meas — line 2 (inches)**")
        d1, d2, d3 = st.columns(3)
        ctn_l = d1.number_input(
            "Length (L)", min_value=0.0, value=24.0, step=0.5, format="%.2f",
            key=f"ctn_l_v9_{style_key}",
        )
        ctn_w = d2.number_input(
            "Width (W)", min_value=0.0, value=16.0, step=0.5, format="%.2f",
            key=f"ctn_w_v9_{style_key}",
        )
        ctn_h = d3.number_input(
            "Height (H)", min_value=0.0, value=10.0, step=0.5, format="%.2f",
            key=f"ctn_h_v9_{style_key}",
        )

        CUIN_PER_CBM = 61000
        ctn_meas_2 = f"L {ctn_l:g}'' X W {ctn_w:g}'' X H {ctn_h:g}''"
        cbm_per_ctn = (ctn_l * ctn_w * ctn_h) / CUIN_PER_CBM

        st.caption(
            f"CTN MEAS: {ctn_meas_2}  |  CBM per carton: {cbm_per_ctn:.8f}  |  "
            f"CBM = ({ctn_l:g} × {ctn_w:g} × {ctn_h:g} × total cartons) / {CUIN_PER_CBM}"
        )

    ship_qty_by_size = {sz: 0 for sz in full_sizes}
    for size_vals, tc in zip(edited_sizes, edited_total_ctn):
        for sz in full_sizes:
            ship_qty_by_size[sz] += int(size_vals.get(sz, 0) or 0) * tc

    order_qty_by_size = {s["size"]: s["qty"] for s in sublines}
    order_total = sum(order_qty_by_size.get(sz, 0) for sz in full_sizes)
    ship_total = sum(ship_qty_by_size.get(sz, 0) for sz in full_sizes)
    cut_total = sum(cut_qty.get(sz, 0) for sz in full_sizes)
    exs_by_size = {sz: ship_qty_by_size.get(sz, 0) - order_qty_by_size.get(sz, 0) for sz in full_sizes}
    exs_total = ship_total - order_total
    pct_by_size = {
        sz: (exs_by_size[sz] / order_qty_by_size[sz]) if order_qty_by_size.get(sz) else None
        for sz in full_sizes
    }
    pct_total = (exs_total / order_total) if order_total else None
    cbm_total = cbm_per_ctn * total_ctn

    st.markdown('<div class="plist-wrap">', unsafe_allow_html=True)

    st.markdown(f"""
<table class="plist noborder">
<tr class="title-row">
  <td style="width:16%"><b>BUYER</b></td><td style="width:14%">VANS</td>
  <td style="width:14%"></td>
  <td style="width:16%"><b>ORDER QTY</b></td><td style="width:8%">{line['order_qty']}</td><td style="width:6%">PCS</td>
  <td style="width:8%"></td>
  <td style="width:6%"><u><b>CRD</b></u></td><td>{line.get('crd','')}</td>
</tr>
<tr class="title-row">
  <td><b>LOT NO</b></td><td>{line['style']}</td><td></td>
  <td><b>PACK QTY</b></td><td>{line['order_qty']}</td><td>PCS</td><td></td><td></td><td></td>
</tr>
<tr class="title-row">
  <td><b>P.O NO</b></td><td>{line['po_line_no']}</td><td></td>
  <td><b>EXCESS/SHORT QTY</b></td><td>{exs_total}</td><td>PCS</td><td></td><td></td><td></td>
</tr>
<tr class="title-row">
  <td colspan="2" class="desc-cell">DESCRIPTION:<br>{line['description']}</td><td></td>
  <td><b>PERCENTAGE</b></td>
  <td>{'-' if pct_total is None else f'{pct_total*100:.2f}'}</td><td>%</td><td></td>
  <td colspan="2">{line.get('destination_country','')}</td>
</tr>
</table>
""", unsafe_allow_html=True)

    def fmt_pct(v):
        return "" if v is None else f"{v*100:.2f}%"

    n_data_rows = 5
    size_th = "".join(f"<th>{sz}</th>" for sz in full_sizes)

    summary_body = ""
    for i, (label, vals, gt, cls) in enumerate([
        ("CUT QTY", [cut_qty.get(sz, "") for sz in full_sizes], cut_total, "yellow"),
        ("ORDER QTY", [order_qty_by_size.get(sz, "") for sz in full_sizes], order_total, "gray"),
        ("SHIP QTY", [ship_qty_by_size.get(sz, "") for sz in full_sizes], ship_total, ""),
        ("EXS/SHT QTY", [exs_by_size.get(sz, "") for sz in full_sizes], exs_total, ""),
        ("PERCENTAGE", [fmt_pct(pct_by_size.get(sz)) for sz in full_sizes], fmt_pct(pct_total), ""),
    ]):
        cells = "".join(f'<td class="{cls}">{v}</td>' for v in vals)
        if i == 0:
            weight_cells = (
                f'<td class="bold" rowspan="{n_data_rows}" style="vertical-align:middle;">'
                f'{total_grs:.2f}</td>'
                f'<td class="bold" rowspan="{n_data_rows}" style="vertical-align:middle;">'
                f'{total_net:.2f}</td>'
            )
        else:
            weight_cells = ""
        summary_body += (
            f'<tr><td class="{cls} bold">{label}</td>{cells}'
            f'<td class="bold">{gt}</td>{weight_cells}</tr>'
        )

    st.markdown(f"""
<table class="plist">
<tr class="hdr">
  <th>SIZE</th>{size_th}
  <th>G Total</th>
  <th>Grand Total gross weight kg</th>
  <th>Grand Total Net weight kg</th>
</tr>
{summary_body}
</table>
""", unsafe_allow_html=True)

    st.markdown(f"""
<table class="plist noborder">
<tr class="title-row"><td style="width:15%"><b>GROSS WEIGHT :</b></td><td>{total_grs:.2f}</td></tr>
<tr class="title-row"><td><b>NET WEIGHT :</b></td><td>{total_net:.2f}</td></tr>
<tr class="title-row"><td><b>CTN MEAS :</b></td><td>{ctn_meas_1}<br>{ctn_meas_2}</td></tr>
<tr class="title-row"><td><b>CBM :</b></td><td>{cbm_total:.8f}</td></tr>
</table>
<br><br>
<table class="plist noborder">
<tr class="title-row"><td style="width:30%; border-top:1px solid #000; text-align:center;">INCHARGE</td>
<td style="width:30%"></td>
<td style="width:30%; border-top:1px solid #000; text-align:center;">FM</td></tr>
</table>
""", unsafe_allow_html=True)

    st.markdown('</div>', unsafe_allow_html=True)

    rows_export = []
    for i in range(n_rows):
        rows_export.append({
            "ctn_no": ctn_ranges[i],
            "color": row_colors[i],
            "upc": row_upcs[i],
            "sizes": edited_sizes[i],
            "ctn_pcs": edited_ctn_pcs[i],
            "total_ctn": edited_total_ctn[i],
            "grs_per_ctn": grs_wts[i],
            "net_per_ctn": net_wts[i],
            "total_pcs": row_total_pcs[i],
            "total_grs": edited_total_ctn[i] * grs_wts[i],
            "total_net": edited_total_ctn[i] * net_wts[i],
        })

    export_state = {
        "line": line,
        "full_sizes": full_sizes,
        "rows": rows_export,
        "total_ctn": total_ctn,
        "total_pcs": total_pcs,
        "total_grs": total_grs,
        "total_net": total_net,
        "size_packet_totals": size_packet_totals,
        "cut_qty": cut_qty,
        "order_qty_by_size": order_qty_by_size,
        "ship_qty_by_size": ship_qty_by_size,
        "exs_by_size": exs_by_size,
        "pct_by_size": pct_by_size,
        "cut_total": cut_total,
        "order_total": order_total,
        "ship_total": ship_total,
        "exs_total": exs_total,
        "pct_total": pct_total,
        "ctn_meas_1": ctn_meas_1,
        "ctn_meas_2": ctn_meas_2,
        "cbm_per_ctn": cbm_per_ctn,
        "cbm_total": cbm_total,
        "nw_by_size": nw_by_size,
        "nnw_by_size": nnw_by_size,
        "empty_by_size": empty_by_size,
    }

    st.markdown("##### Download this packing list (with your edits)")
    dl_col1, dl_col2 = st.columns(2)

    with dl_col1:
        if st.button("Prepare Excel (.xlsx)", key=f"prep_xlsx_v8_{style_key}"):
            wb = build_sample_style_workbook(export_state)
            xbuf = io.BytesIO()
            wb.save(xbuf)
            xbuf.seek(0)
            st.session_state[f"xlsx_bytes_v8_{style_key}"] = xbuf.getvalue()
            st.session_state[f"xlsx_ts_v8_{style_key}"] = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        if st.session_state.get(f"xlsx_bytes_v8_{style_key}"):
            _ts_x = st.session_state.get(f"xlsx_ts_v8_{style_key}", "latest")
            st.download_button(
                label=f"Download {style_key} packing list (.xlsx)",
                data=st.session_state[f"xlsx_bytes_v8_{style_key}"],
                file_name=f"packing_list_{style_key}_{_ts_x}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key=f"dl_xlsx_v8_{style_key}",
            )

    with dl_col2:
        if st.button("Prepare PDF", key=f"prep_pdf_v8_{style_key}"):
            pbuf = build_edited_pdf(export_state)
            st.session_state[f"pdf_bytes_v8_{style_key}"] = pbuf.getvalue()
            st.session_state[f"pdf_ts_v8_{style_key}"] = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        if st.session_state.get(f"pdf_bytes_v8_{style_key}"):
            _ts_p = st.session_state.get(f"pdf_ts_v8_{style_key}", "latest")
            st.download_button(
                label=f"Download {style_key} packing list (.pdf)",
                data=st.session_state[f"pdf_bytes_v8_{style_key}"],
                file_name=f"packing_list_{style_key}_{_ts_p}.pdf",
                mime="application/pdf",
                key=f"dl_pdf_v8_{style_key}",
            )


# ============================================================
# PART 3 - STREAMLIT UI
# ============================================================

st.set_page_config(page_title="PO -> Packing List", layout="wide")
inject_ui_theme()
st.title("PO to Packing List Converter")
st.write(
    "Upload a VF/Vans-style Purchase Order PDF. This extracts the line items "
    "and sizes, then builds a packing-list-style Excel workbook (one sheet per "
    "PO line) with live formulas for cartons, weights, and totals."
)

uploaded_file = st.file_uploader("Upload PO PDF", type=["pdf"])

if uploaded_file is not None:
    with st.spinner("Reading PDF..."):
        try:
            data = extract_data(uploaded_file)
        except Exception as e:
            st.error(f"Failed to read the PDF: {e}")
            st.stop()

    if not data["lines"]:
        st.warning(
            "No PO lines were found. This PDF's layout may differ from the "
            "expected format - the extraction patterns may need adjusting."
        )
        st.stop()

    st.success(f"Found {len(data['lines'])} PO line(s) and {len(data['sublines'])} size/color subline(s).")

    # --- ADDED: clear stale session state when a NEW PDF is uploaded ---
    _upload_sig = f"{uploaded_file.name}::{uploaded_file.size}"
    if st.session_state.get("_last_upload_sig") != _upload_sig:
        for _k in list(st.session_state.keys()):
            if _k.startswith(("carton_rows_v8_", "pack_sig_v8_",
                              "xlsx_bytes_v8_", "pdf_bytes_v8_",
                              "xlsx_ts_v8_", "pdf_ts_v8_",
                              "insert_after_v8_", "delete_row_v8_")):
                st.session_state.pop(_k, None)
        st.session_state["_last_upload_sig"] = _upload_sig
    # --- END ADDED ---

    st.subheader("PO Lines")
    st.caption("Click a row to filter the sublines below and download just that line's sheet.")
    po_line_rows = [
        {
            "PO Line No.": l["po_line_no"],
            "Style": l["style"],
            "Description": l["description"],
            "Order Qty": l["order_qty"],
            "Unit Price": l["unit_price"],
            "Line Cost": l["line_cost"],
            "CRD": l.get("crd"),
            "Country": l.get("destination_country"),
        }
        for l in data["lines"]
    ]
    selection = st.dataframe(
        po_line_rows,
        use_container_width=True,
        on_select="rerun",
        selection_mode="single-row",
        key="po_lines_table",
    )

    selected_idx = None
    rows_selected = selection.selection.rows if selection is not None else []
    if rows_selected:
        selected_idx = rows_selected[0]
        selected_line = data["lines"][selected_idx]
        st.info(f"Selected: **{selected_line['style']}** (PO Line {selected_line['po_line_no']})")

    st.subheader("Sublines (size / color / UPC)")
    if selected_idx is not None:
        filtered_sublines = [s for s in data["sublines"] if s["style"] == selected_line["style"]]
        st.dataframe(filtered_sublines, use_container_width=True)
    else:
        st.dataframe(data["sublines"], use_container_width=True)

    # --- ADDED: Download extracted JSON ---
    st.download_button(
        label="⬇ Download extracted JSON",
        data=json.dumps(data, indent=2).encode("utf-8"),
        file_name=f"po_extracted_{len(data['lines'])}_lines.json",
        mime="application/json",
        key="dl_po_json",
    )
    # --- END ADDED ---

    col1, col2 = st.columns(2)

    with col1:
        if st.button("Build full workbook (all PO lines)", type="primary"):
            with st.spinner("Building Excel workbook..."):
                wb = build_workbook(data)
                buf = io.BytesIO()
                wb.save(buf)
                buf.seek(0)

            st.success(f"Workbook built with {len(wb.sheetnames)} sheet(s): {', '.join(wb.sheetnames)}")
            st.download_button(
                label="Download packing_list_all_PO_lines.xlsx",
                data=buf,
                file_name="packing_list_all_PO_lines.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

    with col2:
        if selected_idx is not None:
            if st.button(f"Build sheet for {selected_line['style']} only"):
                single_line_data = {
                    "lines": [selected_line],
                    "sublines": [s for s in data["sublines"] if s["style"] == selected_line["style"]],
                }
                with st.spinner("Building Excel sheet..."):
                    wb = build_workbook(single_line_data)
                    buf = io.BytesIO()
                    wb.save(buf)
                    buf.seek(0)

                st.success(f"Sheet built for {selected_line['style']}")
                st.download_button(
                    label=f"Download {selected_line['style']}.xlsx",
                    data=buf,
                    file_name=f"packing_list_{selected_line['style']}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                st.session_state["preview_style"] = selected_line["style"]
        else:
            st.caption("Select a PO line above to download just that sheet.")

    if selected_idx is not None and st.session_state.get("preview_style") == selected_line["style"]:
        st.subheader(f"Packing List Preview — {selected_line['style']}")
        preview_sublines = [s for s in data["sublines"] if s["style"] == selected_line["style"]]
        render_packing_list_preview(selected_line, preview_sublines)

    st.info(
        "Yellow cells in the downloaded file still need manual entry: "
        "CUT QTY, CTN MEAS, and CBM factor - these come from your factory's "
        "cutting/packing records, not the PO. "
        "CARTON NO., SHIP QTY, GROSS/NET WEIGHT are auto-calculated."
    )
else:
    st.info("Upload a PDF to get started.")